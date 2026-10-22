#!/usr/bin/env python3
"""Mooncake: which plane carries decode-to-master RPC."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def box(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    return shape


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Mooncake Decode 到 Master 的 RPC 平面"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.14, 10.2, 0.22, "Mooncake Store    ·    有 Master    ·    MemFabric 无此节点", 12, True, ACCENT)
    textbox(slide, 0.42, 0.36, 11.8, 0.40, "Decode 到 Master 默认是 TCP RPC", 26, True, NAVY)
    textbox(slide, 11.35, 0.16, 1.55, 0.22, "1  /  1", 12, False, MUTED, PP_ALIGN.RIGHT)

    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(0.84), Inches(12.5), Inches(0.40))
    solid(banner, PALE)
    textbox(
        slide,
        0.56,
        0.84,
        12.22,
        0.40,
        "网卡由 Master 的 IP 决定。这个 IP 配在业务面网卡上，RPC 就走业务面。KV 字节不走这跳。",
        15,
        True,
        TEAL,
    )

    # Decode node
    box(slide, 0.42, 1.42, 3.55, 4.55, SOFT, LINE)
    textbox(slide, 0.56, 1.50, 3.27, 0.32, "Decode 节点", 16, True, NAVY, PP_ALIGN.CENTER)
    textbox(slide, 0.56, 1.82, 3.27, 0.28, "推理进程里的 Mooncake Client", 12, False, MUTED, PP_ALIGN.CENTER)
    box(slide, 0.62, 2.24, 3.15, 1.45, WHITE, NAVY)
    textbox(slide, 0.74, 2.32, 2.91, 0.30, "RPC 出口", 15, True, NAVY, PP_ALIGN.CENTER)
    textbox(slide, 0.74, 2.64, 2.91, 0.85, "路由到 Master IP 的网卡\nTCP，默认端口 50051", 13, False, INK, PP_ALIGN.CENTER)
    box(slide, 0.62, 3.86, 3.15, 1.85, WHITE, TEAL)
    textbox(slide, 0.74, 3.96, 2.91, 0.30, "RDMA 网卡", 15, True, TEAL, PP_ALIGN.CENTER)
    textbox(slide, 0.74, 4.30, 2.91, 1.20, "TransferEngine\nRDMA READ / WRITE\nKV 字节从这里出去", 13, False, INK, PP_ALIGN.CENTER)

    # Control arrow and master
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.12), Inches(2.72), Inches(1.15), Inches(0.28))
    solid(arrow, NAVY)
    textbox(slide, 4.05, 2.28, 1.30, 0.40, "TCP RPC", 13, True, NAVY, PP_ALIGN.CENTER)

    box(slide, 5.42, 1.42, 3.55, 2.55, WHITE, NAVY)
    textbox(slide, 5.56, 1.50, 3.27, 0.32, "Master 节点", 16, True, NAVY, PP_ALIGN.CENTER)
    textbox(slide, 5.56, 1.86, 3.27, 1.90, "mooncake_master\n监听 50051，TCP_NODELAY 开\n处理 PutStart、GetReplicaList\n返回偏移、副本和租约\n默认线程数 4", 13, False, INK, PP_ALIGN.CENTER)

    # Data arrow and DDR
    arrow2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.12), Inches(4.55), Inches(4.85), Inches(0.28))
    solid(arrow2, TEAL)
    textbox(slide, 4.55, 4.14, 3.9, 0.36, "RDMA    KV 字节，不经过 Master", 13, True, TEAL, PP_ALIGN.CENTER)

    box(slide, 9.15, 3.55, 3.75, 2.42, PALE, TEAL)
    textbox(slide, 9.29, 3.66, 3.47, 0.32, "DDR 池", 16, True, TEAL, PP_ALIGN.CENTER)
    textbox(slide, 9.29, 4.04, 3.47, 1.70, "其他节点上已注册的 DRAM\n对端网卡按 rkey 做 DMA\n对端 CPU 不读这次请求", 13, False, INK, PP_ALIGN.CENTER)

    foot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(6.12), Inches(12.5), Inches(0.92))
    solid(foot, GOLD_BG)
    textbox(
        slide,
        0.56,
        6.16,
        12.22,
        0.84,
        "Master 默认 bind 0.0.0.0，所有网卡都听。--rpc_interface=eth0 时改用这张网卡的 IPv4。\n"
        "Decode 从路由到该 IP 的口发出。MC_RPC_PROTOCOL 默认 tcp，只有设成 rdma 时这跳才改走 RDMA 网卡。",
        14,
        False,
        GOLD,
    )
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.22,
        12.5,
        0.22,
        "MemFabric 没有 Master    ·    进程间控制是 Host TCP    ·    数据走 Device UB / RoCE 或 Host RoCE",
        12,
        False,
        MUTED,
    )
    slide.notes_slide.notes_text_frame.text = (
        "Mooncake Store 的 Decode 侧 Client 访问 Master 是 RPC。"
        "master.cpp 里 MC_RPC_PROTOCOL 默认 tcp，只有环境变量精确等于 rdma 时才改协议。"
        "默认端口 50051，来自 --port 的默认值；--rpc_port 为 0 时回退到它。"
        "--rpc_enable_tcp_no_delay 默认 true。"
        "--rpc_address 默认 0.0.0.0，即所有网卡监听。"
        "--rpc_interface 非空时用该网卡的 IPv4 覆盖 rpc_address，例如 --rpc_interface=eth0。"
        "客户端地址是 master_server_addr，形如 IP:50051。报文从 Decode 主机路由到这个 IP 的那张网卡出去。"
        "因此业务面网卡只有在 Master IP 配在业务面地址上时才承载这跳 RPC。代码没有单独的控制面网卡类型。"
        "KV 字节走 TransferEngine 的 RDMA READ 或 WRITE，目的是对端已注册内存的地址和 rkey，不进入 Master。"
        "--rpc_thread_num 为 0 时使用已废弃的 --max_threads，其默认值是 4。"
        "MemFabric 的 README 写明没有这类 Master。acc_links 用 Host TCP 传进程间控制命令。"
        "数据面按南向插件走 Device UB、Device RoCE、Host RoCE 等，xcopy 使用全局虚拟地址。"
        "示例里的 store_url 是 tcp://ip:port，用于配置发现，不是 KV 数据面。"
    )
    out = "/workspace/docs/mooncake-master-plane.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
