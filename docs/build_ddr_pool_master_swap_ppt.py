#!/usr/bin/env python3
"""AS-IS vs TO-BE: decode cards, master queue, and the bus to the DDR pool."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def card(slide, x, y, w, h, title, sub, fill, title_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, LINE)
    textbox(slide, x + 0.06, y + 0.04, w - 0.12, 0.28, title, 12, True, title_color, PP_ALIGN.CENTER)
    textbox(slide, x + 0.06, y + 0.30, w - 0.12, h - 0.36, sub, 11, False, MUTED, PP_ALIGN.CENTER)


def right_arrow(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill)
    return shape


def down_arrow(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill)
    return shape


def queue_slots(slide, x, y, total_w, labels, fill, text_color, line=None):
    gap = 0.06
    slot_w = (total_w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        sx = x + i * (slot_w + gap)
        slot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(sx), Inches(y), Inches(slot_w), Inches(0.36))
        solid(slot, fill, line)
        textbox(slide, sx, y, slot_w, 0.36, label, 11, True, text_color, PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "DDR 池直挂总线的 AS-IS 与 TO-BE"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.14, 9.6, 0.22, "拉高 batch    ·    控制面海军蓝    ·    总线平面青色", 12, True, ACCENT)
    textbox(slide, 0.42, 0.36, 11.6, 0.38, "热路径从 Master 排队，改成直走总线", 24, True, NAVY)
    textbox(slide, 11.35, 0.16, 1.55, 0.22, "1  /  1", 12, False, MUTED, PP_ALIGN.RIGHT)

    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(0.82), Inches(12.5), Inches(0.36))
    solid(banner, PALE)
    textbox(
        slide,
        0.56,
        0.82,
        12.22,
        0.36,
        "左为 AS-IS：每次 swap 进 Master 排队，出队后才 DMA。右为 TO-BE：块表已有偏移，热路径只走总线。",
        14,
        True,
        TEAL,
    )

    panels = (
        (0.40, "AS-IS    每次 swap 都进 Master 排队", NAVY),
        (6.86, "TO-BE    热路径只走总线平面", TEAL),
    )
    for x, title, fill in panels:
        head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.28), Inches(6.08), Inches(0.32))
        solid(head, fill)
        textbox(slide, x + 0.12, 1.28, 5.84, 0.32, title, 14, True, WHITE)

    # AS-IS: three cards fan into one master, then the bus.
    cards = (("Decode 卡 1", "换入 / 换出"), ("Decode 卡 2", "换入 / 换出"), ("Decode 卡 N", "换入 / 换出"))
    for i, (title, sub) in enumerate(cards):
        y = 1.72 + i * 0.78
        card(slide, 0.52, y, 2.05, 0.68, title, sub, SOFT, NAVY)
        right_arrow(slide, 2.64, y + 0.20, 0.46, 0.24, NAVY)

    master = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.16), Inches(1.72), Inches(3.18), Inches(2.24))
    solid(master, WHITE, NAVY)
    textbox(slide, 3.28, 1.78, 2.94, 0.28, "Master 节点", 14, True, NAVY, PP_ALIGN.CENTER)
    textbox(slide, 3.28, 2.06, 2.94, 0.22, "控制面队列", 11, True, GOLD, PP_ALIGN.CENTER)
    queue_slots(slide, 3.30, 2.32, 2.90, ["卡1", "卡2", "卡3", "…", "卡N"], GOLD_BG, GOLD)
    textbox(slide, 3.28, 2.74, 2.94, 0.52, "入队：块 id、字节数、换入或换出\n出队：池内偏移、长度、租约", 12, False, INK, PP_ALIGN.CENTER)
    textbox(slide, 3.28, 3.28, 2.94, 0.52, "多卡同时 swap，队列变长\nDMA 约 0.06 ms，往返按 0.2 ms", 12, False, GOLD, PP_ALIGN.CENTER)

    down_arrow(slide, 3.20, 4.04, 0.28, 0.28, TEAL)
    textbox(slide, 3.52, 4.00, 2.70, 0.36, "出队后才走总线平面", 12, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)

    ddr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.52), Inches(4.40), Inches(5.82), Inches(0.78))
    solid(ddr, PALE, TEAL)
    textbox(slide, 0.66, 4.46, 2.3, 0.30, "DDR 池", 14, True, TEAL)
    textbox(slide, 0.66, 4.76, 5.5, 0.32, "直挂总线。KV 按偏移 DMA 进来，字节不进 Master。", 12, False, INK)

    # TO-BE: cards go straight to the pool; master sits off the hot path.
    for i, (title, sub) in enumerate(cards):
        card(slide, 7.00 + i * 1.96, 1.72, 1.86, 0.68, title, "块表里已有偏移", PALE, TEAL)
    down_arrow(slide, 9.68, 2.48, 0.28, 0.26, TEAL)
    textbox(slide, 10.02, 2.44, 2.5, 0.32, "总线平面", 12, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)

    ddr2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.00), Inches(2.82), Inches(5.80), Inches(1.05))
    solid(ddr2, PALE, TEAL)
    textbox(slide, 7.16, 2.92, 5.48, 0.32, "DDR 池", 16, True, TEAL, PP_ALIGN.CENTER)
    textbox(slide, 7.16, 3.26, 5.48, 0.46, "热路径。三张卡按自己块表里的偏移做 DMA。\n不等待 Master。", 13, False, INK, PP_ALIGN.CENTER)

    master2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.00), Inches(4.02), Inches(5.80), Inches(1.16))
    solid(master2, SOFT, LINE)
    textbox(slide, 7.14, 4.08, 2.4, 0.28, "Master 节点", 14, True, MUTED)
    textbox(slide, 9.50, 4.08, 3.14, 0.28, "热路径上没有这条队列", 12, True, TEAL, PP_ALIGN.RIGHT)
    queue_slots(slide, 7.16, 4.42, 2.55, ["新块", "释放", "租约"], WHITE, NAVY, LINE)
    textbox(slide, 9.84, 4.42, 2.80, 0.36, "控制面只收这三类", 12, True, NAVY, PP_ALIGN.RIGHT)
    textbox(slide, 7.16, 4.82, 5.48, 0.28, "再问一次，返回的仍是块表里的同一行偏移和租约。", 12, False, INK)

    foot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.40), Inches(5.36), Inches(12.54), Inches(1.68))
    solid(foot, GOLD_BG)
    textbox(
        slide,
        0.56,
        5.46,
        12.22,
        1.48,
        "AS-IS：每次 swap 经控制面入队。入队是块 id、字节数、换入或换出。出队是池内偏移、长度、租约。\n"
        "总线 DMA 在出队之后。4 MB、64 GB/s 时约 0.06 ms；往返按 0.2 ms，多卡排队后控制面先饱和。\n"
        "TO-BE：偏移和租约留在卡上。热路径只有总线。Master 只在新块进池、块释放、租约到期时出现。",
        14,
        False,
        GOLD,
    )
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.22,
        12.5,
        0.22,
        "跨卡共享的 DDR 池才有左侧 Master    ·    单机进程内的 host DDR 块表走右侧    ·    0.06 ms / 0.2 ms 为估算",
        11,
        False,
        MUTED,
    )
    slide.notes_slide.notes_text_frame.text = (
        "AS-IS 画的是为拉高 batch，把 DDR 池做成跨卡共享，并且每次换入换出都向 Master 要地址。"
        "三张 Decode 卡经控制面把块 id、字节数、方向送入 Master 内部队列。Master 出队时返回池内偏移、长度、租约。"
        "总线平面的 DMA 由 Decode 卡发往直挂的 DDR 池，发生在出队之后。KV 字节不进入 Master。"
        "估算：4 MB / 64 GB/s 约 0.06 ms；不排队的往返按 0.2 ms。一步 20 ms、batch 32、每请求换入换出各一块，"
        "单卡约 3200 次/秒，8 卡约 2.6 万次/秒。单线程按 50 us 服务约 2 万次/秒，队列变长。"
        "TO-BE 中块表保存偏移和租约。热路径是 Decode 卡到 DDR 池的总线 DMA，Master 队列不在这条路径上。"
        "控制面只在新块第一次进池、块释放、租约到期时进入 Master。"
        "当前 vLLM OffloadingConnector 和 SGLang HiCache L2 的本机 host DDR 已经是进程内块表，形态接近右侧，没有左侧这跳。"
    )
    out = "/workspace/docs/ddr-pool-master-swap.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
