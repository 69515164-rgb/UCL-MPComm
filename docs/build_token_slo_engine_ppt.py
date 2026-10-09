#!/usr/bin/env python3
"""Editable deck: prefill Token SLO Engine, next-batch budget and admission."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)
TEAL = (14, 122, 118)
CORAL = (176, 64, 52)
GOLD = (140, 100, 36)
GREEN = (27, 110, 72)
LINE = (214, 220, 226)
SOFT = (244, 247, 250)
COVER_BG = (232, 243, 242)
GREEN_BG = (232, 243, 236)
GOLD_BG = (255, 246, 230)
CORAL_BG = (255, 240, 236)
PILL = (236, 240, 244)


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    strip_shadow(shape)


def write(shape, text, size, bold, color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def box(slide, x, y, w, h, fill, text="", size=12, bold=True, color=WHITE, line=None, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    if text:
        write(shape, text, size, bold, color, align, anchor)
    return shape


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def header(slide, kicker, title, lead):
    box(slide, 0, 0, 13.333, 0.08, ACCENT, "")
    textbox(slide, 0.32, 0.14, 12.6, 0.22, kicker, 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.32, 0.36, 12.7, 0.36, title, 22, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    box(slide, 0.32, 0.80, 0.07, 0.36, ACCENT, "")
    textbox(slide, 0.48, 0.76, 12.5, 0.42, lead, 14, True, INK, anchor=MSO_ANCHOR.MIDDLE)


def input_group(slide, x, y, w, h, rail, title, items):
    box(slide, x, y, w, h, WHITE, line=LINE)
    box(slide, x, y, 0.08, h, rail, "")
    textbox(slide, x + 0.16, y + 0.06, w - 0.26, 0.26, title, 13, True, rail)
    n = len(items)
    top = y + 0.36
    gap = 0.06
    pill_h = min(0.32, (h - 0.44 - gap * (n - 1)) / n)
    for i, item in enumerate(items):
        py = top + i * (pill_h + gap)
        box(slide, x + 0.16, py, w - 0.28, pill_h, PILL, item, 12, True, NAVY)


def step(slide, x, y, w, h, n, title, body):
    box(slide, x, y, w, h, WHITE, line=LINE)
    box(slide, x + 0.08, y + 0.10, 0.32, 0.28, TEAL, str(n), 13, True, WHITE)
    textbox(slide, x + 0.48, y + 0.08, w - 0.58, 0.30, title, 14, True, TEAL)
    textbox(slide, x + 0.12, y + 0.40, w - 0.22, h - 0.46, body, 12, False, INK, anchor=MSO_ANCHOR.TOP)


def set_cell_border(cell, color_hex="D6DCE2"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", "6350")
        sf = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color_hex)
        sf.append(srgb)
        ln.append(sf)
        tcPr.append(ln)


def write_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold, color)
    set_cell_border(cell)


def slide_engine(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "PD 分离  ·  Prefill 闭环    1 / 2",
        "每拍只改下一批的 token 预算，并决定这条请求进不进",
        "已在 GPU 上的批保持不动。Decode 的批大小由另一条环按 TPOT 调整，本引擎只读它的空槽。",
    )

    dx, dy, dh = 0.28, 1.26, 4.02
    lw, cw, rw = 3.22, 5.22, 3.62
    gap = 0.32
    cx = dx + lw + gap
    rx = cx + cw + gap

    gy = dy
    gh = (dh - 0.16) / 3
    groups = (
        (ACCENT, "约束与现状", ("TTFT 目标", "当前批 token 预算", "prefill 实测吞吐")),
        (TEAL, "队列与本请求", ("本请求的 prompt token 数", "队头已经等待的时间", "排在前面的 token 数")),
        (GOLD, "资源与下游", ("激活显存余量", "KV 有效带宽与在途条数", "decode 空槽数")),
    )
    for rail, title, items in groups:
        input_group(slide, dx, gy, lw, gh, rail, title, items)
        gy += gh + 0.08

    box(slide, cx, dy, cw, dh, COVER_BG, "")
    box(slide, cx, dy, cw, 0.38, TEAL, "Token SLO Engine", 16, True, WHITE)
    sy = dy + 0.48
    sh = (dh - 0.58) / 3
    steps = (
        ("1", "预测", "假设本请求装进下一批。\n预测 TTFT = 队列等待 + 本批计算 + 露出的搬运 + 接纳等待"),
        ("2", "归因", "误差 = TTFT 目标 - 预测 TTFT。\n四段里最大的一段，决定动预算还是动准入。"),
        ("3", "限幅", "死区内预算保持。出死区每拍只改一档。\n上界 = min(TTFT 还允许的预算, 显存放得下的预算)"),
    )
    for i, (n, title, body) in enumerate(steps):
        step(slide, cx + 0.10, sy + i * sh, cw - 0.20, sh - 0.08, n, title, body)

    oh = (dh - 0.10) / 2
    box(slide, rx, dy, rw, oh, WHITE, line=LINE)
    box(slide, rx, dy, 0.08, oh, ACCENT, "")
    textbox(slide, rx + 0.18, dy + 0.08, rw - 0.30, 0.24, "输出 1", 12, True, ACCENT)
    textbox(slide, rx + 0.18, dy + 0.30, rw - 0.30, 0.28, "下一批 token 预算", 16, True, NAVY)
    chip_y = dy + 0.68
    chip_w = (rw - 0.36 - 0.12) / 3
    for i, (label, fill, fg) in enumerate((("+1 档", GREEN_BG, GREEN), ("保持", PILL, NAVY), ("-1 档", CORAL_BG, CORAL))):
        box(slide, rx + 0.16 + i * (chip_w + 0.06), chip_y, chip_w, 0.32, fill, label, 12, True, fg)
    textbox(
        slide,
        rx + 0.18,
        chip_y + 0.40,
        rw - 0.32,
        0.70,
        "相对当前预算移动一档。一档是按显存和 prefill 吞吐预先划好的相邻台阶。执行器记的是 token 数。",
        12,
        False,
        INK,
        anchor=MSO_ANCHOR.TOP,
    )

    oy = dy + oh + 0.10
    box(slide, rx, oy, rw, oh, WHITE, line=LINE)
    box(slide, rx, oy, 0.08, oh, GREEN, "")
    textbox(slide, rx + 0.18, oy + 0.08, rw - 0.30, 0.24, "输出 2", 12, True, GREEN)
    textbox(slide, rx + 0.18, oy + 0.30, rw - 0.30, 0.28, "本请求", 16, True, NAVY)
    for i, (label, fill, fg) in enumerate((("准入", GREEN_BG, GREEN), ("推迟", GOLD_BG, GOLD), ("拒绝", CORAL_BG, CORAL))):
        box(slide, rx + 0.16 + i * (chip_w + 0.06), oy + 0.68, chip_w, 0.32, fill, label, 12, True, fg)
    textbox(
        slide,
        rx + 0.18,
        oy + 1.08,
        rw - 0.32,
        0.72,
        "准入进入下一批。推迟留下一拍再判。拒绝只在队列已经把 TTFT 顶过目标时使用。",
        12,
        False,
        INK,
        anchor=MSO_ANCHOR.TOP,
    )

    arrow_y = dy + dh / 2 - 0.12
    for ax in (dx + lw + 0.02, cx + cw + 0.02):
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(arrow_y), Inches(0.28), Inches(0.24))
        solid(shape, TEAL)

    box(
        slide,
        0.28,
        5.40,
        12.77,
        0.34,
        TEAL,
        "闭环回读：下一拍的当前预算 = 本拍输出。刚完成请求的 TTFT 写入预测窗口。",
        13,
        True,
        WHITE,
    )

    notes = (
        (ACCENT, "预算按 token 计", "一条长请求占掉的 token，等于很多条短请求。台阶按 token 数划，长度直接进入预测。"),
        (TEAL, "只改下一拍", "当前预算是加减一档的基准。已经发射的 kernel 跑完为止，闭环从下一次组批生效。"),
        (GOLD, "准入单独给出", "队列是大头时拒绝新请求。decode 没有空槽时推迟。decode 批大小留在 TPOT 环。"),
    )
    card_y, card_h = 5.86, 1.42
    card_gap = 0.12
    card_w = (12.77 - 2 * card_gap) / 3
    for i, (color, title, body) in enumerate(notes):
        x = 0.28 + i * (card_w + card_gap)
        box(slide, x, card_y, card_w, card_h, WHITE, line=LINE)
        box(slide, x, card_y, card_w, 0.08, color, "")
        textbox(slide, x + 0.14, card_y + 0.16, card_w - 0.28, 0.28, title, 14, True, color)
        textbox(slide, x + 0.14, card_y + 0.48, card_w - 0.28, 0.84, body, 13, False, INK, anchor=MSO_ANCHOR.TOP)

    slide.notes_slide.notes_text_frame.text = (
        "Token SLO Engine 服务 PD 分离下的 prefill 池。"
        "输入保留 TTFT 目标、当前批 token 预算、本请求长度，并补上队列、激活显存、KV 带宽与在途条数、decode 空槽。"
        "输出 1 是下一批 token 预算，相对当前预算每拍最多一档。"
        "输出 2 是本请求的准入、推迟或拒绝。"
        "预测 TTFT = 队列等待 + 本批计算 + 露出的搬运 + 接纳等待。"
        "已发射的批不在执行器里。Decode 批大小由 TPOT 环调整。"
    )


def slide_law(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "PD 分离  ·  Prefill 闭环    2 / 2",
        "TTFT 要越线时，先看是哪一段把它顶上去",
        "本批计算或 KV 搬运是大头，才把下一批预算减一档。队列已是大头就保持预算并拒绝新请求。",
    )

    textbox(
        slide,
        0.32,
        1.20,
        12.7,
        1.36,
        "预测 TTFT = 队列等待 + 本批计算 + 露出的搬运 + 接纳等待\n"
        "误差 = TTFT 目标 - 预测 TTFT\n"
        "本批计算 = 下一批 token 数 / prefill 吞吐\n"
        "队列等待 = 队头已等 + 前面各批剩余时间\n"
        "露出的搬运 = max(0, 同时完成条数 × 单条 KV 字节 × 8 / 有效带宽 - 可重叠时间)\n"
        "单条 KV 字节 = 2 × 层数 × KV 头数 × 头维 × token 数 × 每元素字节",
        13,
        False,
        INK,
        anchor=MSO_ANCHOR.TOP,
    )

    rows = (
        ("相对 TTFT 目标", "主导项", "下一批 token 预算", "本请求", NAVY, WHITE, True),
        ("低于目标，超出死区", "显存与带宽仍有余", "加大一档", "准入", GREEN_BG, GREEN, True),
        ("落在死区内", "四段都在目标内", "保持", "装得下就准入，否则推迟", SOFT, NAVY, True),
        ("高于目标，超出死区", "本批计算", "减小一档", "长请求推迟，短请求可准入", CORAL_BG, CORAL, True),
        ("高于目标，超出死区", "KV 搬运露出", "减小一档", "推迟，把完成时刻错开", CORAL_BG, CORAL, True),
        ("高于目标，超出死区", "队列等待", "保持", "拒绝。预算不再下降", (255, 228, 224), CORAL, True),
        ("高于目标，超出死区", "decode 接纳", "保持", "推迟到出现空槽", GOLD_BG, GOLD, True),
    )
    table_shape = slide.shapes.add_table(len(rows), 4, Inches(0.28), Inches(2.62), Inches(12.77), Inches(3.84))
    table = table_shape.table
    widths = (3.15, 2.55, 2.45, 4.62)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    table.rows[0].height = Inches(0.36)
    for r in range(1, len(rows)):
        table.rows[r].height = Inches(0.58)
    for r, row in enumerate(rows):
        texts = row[:4]
        fill, fg, bold = row[4], row[5], row[6]
        for c, text in enumerate(texts):
            cell_fill = NAVY if r == 0 else fill
            cell_fg = WHITE if r == 0 else (fg if c >= 2 else INK)
            cell_bold = True if r == 0 or c >= 2 else False
            write_cell(table.cell(r, c), text, 13, cell_bold, cell_fg, cell_fill)

    textbox(
        slide,
        0.32,
        6.58,
        12.7,
        0.72,
        "死区是目标附近的一段带宽，避免预算来回跳。观测用队头等待和最近窗口的高分位。\n"
        "prefill 吞吐、有效带宽取该池最近窗口的实测值。台阶数按卡的显存和这个吞吐预先划定。",
        13,
        False,
        MUTED,
        anchor=MSO_ANCHOR.TOP,
    )

    slide.notes_slide.notes_text_frame.text = (
        "控制律按预测 TTFT 的主导项分支。"
        "本批计算或露出的 KV 搬运把 TTFT 顶过目标时，下一批 token 预算减一档。"
        "队列等待已经是大头时，再减预算会降低服务率，所以预算保持，新请求拒绝。"
        "decode 没有空槽时，本请求推迟，prefill 预算不动，空槽由 TPOT 环释放。"
        "低于目标且显存、带宽有余时加一档。死区内保持。"
        "每拍最多改一档。观测用队头和窗口高分位。"
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Token SLO Engine：下一批预算与准入"
    slide_engine(prs)
    slide_law(prs)
    out = "/workspace/docs/token-slo-engine.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
