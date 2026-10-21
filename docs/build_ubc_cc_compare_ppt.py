#!/usr/bin/env python3
"""Editable one-page comparison: connection-level CC vs UBC/CTP shared context."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
HEAD = "0E2A47"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    h = hex_color
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def set_cell_fill(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid_fill = OxmlElement("a:solidFill")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", hex_color)
    solid_fill.append(srgb)
    tcPr.append(solid_fill)


def set_cell_borders(cell, hex_color="D6DCE2", width="6350"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for child in list(tcPr):
            if child.tag == qn(f"a:{edge}"):
                tcPr.remove(child)
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", width)
        fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_color)
        fill.append(srgb)
        ln.append(fill)
        tcPr.append(ln)


def write_cell(cell, text, size, bold, color, align):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def add_table(slide, x, y, widths, rows, heights, kinds):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights)))
    table = shape.table
    table.horz_banding = False
    table.vert_banding = False
    tbl = shape._element.graphic.graphicData.tbl
    tbl_pr = tbl.find(qn("a:tblPr"))
    if tbl_pr is None:
        tbl_pr = tbl.tblPr
    for child in list(tbl_pr):
        if child.tag == qn("a:tableStyleId"):
            tbl_pr.remove(child)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    style = {
        "head": (HEAD, WHITE, True, PP_ALIGN.LEFT),
        "body": (WHITE, INK, False, PP_ALIGN.LEFT),
        "key": (PALE, TEAL, True, PP_ALIGN.LEFT),
        "sum": (SOFT, NAVY, True, PP_ALIGN.LEFT),
    }
    for r, (row, h, kind) in enumerate(zip(rows, heights, kinds)):
        table.rows[r].height = Inches(h)
        fill, color, bold, align = style[kind]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            set_cell_fill(cell, fill if c or kind == "head" else SOFT)
            if kind == "head":
                set_cell_fill(cell, HEAD)
            set_cell_borders(cell)
            this_bold = True if c == 0 and kind != "head" else bold
            this_color = NAVY if c == 0 and kind == "body" else color
            if kind in ("key", "sum", "head"):
                this_color = color
                this_bold = bold or c == 0
            write_cell(cell, text, 12, this_bold, this_color, align)
    return shape


def chrome(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.14, 10.4, 0.26, "UBC  ·  CTP 拥塞控制    ·    对照连接级 CC", 12, True, ACCENT)
    textbox(slide, 0.42, 0.38, 12.5, 0.40, "连接各管一份状态，UBC 按目的和虚通道共用一份", 26, True, NAVY)
    textbox(slide, 11.4, 0.16, 1.5, 0.24, "1  /  1", 12, False, MUTED, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.20,
        12.5,
        0.24,
        "UBC 即 UB_CTP    ·    连接级指 TCP、RoCE 的 QP、UB 的 RTP    ·    CTP 上下文键是 {DCNA 低 16 位, VL}",
        11,
        False,
        MUTED,
    )


def banner(slide, y, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(y), Inches(12.5), Inches(0.42))
    solid(shape, PALE)
    textbox(slide, 0.56, y, 12.22, 0.42, text, 14, True, TEAL)


def column(slide, x, y, w, head, sub, head_fill, steps, step_fill, title_color, result, result_fill, result_color):
    head_h = 0.50
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(head_h))
    solid(bar, head_fill)
    textbox(slide, x + 0.16, y + 0.02, w - 0.28, 0.26, head, 15, True, WHITE)
    textbox(slide, x + 0.16, y + 0.26, w - 0.28, 0.20, sub, 11, False, "D6E4F2" if head_fill == NAVY else "D7F0EE")
    step_y = y + head_h + 0.08
    step_h = 0.52
    gap = 0.14
    for i, (title, detail) in enumerate(steps):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(step_y), Inches(w), Inches(step_h))
        solid(card, step_fill, LINE)
        num = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.10), Inches(step_y + 0.12), Inches(0.28), Inches(0.28))
        solid(num, head_fill)
        textbox(slide, x + 0.10, step_y + 0.12, 0.28, 0.28, str(i + 1), 11, True, WHITE, PP_ALIGN.CENTER)
        textbox(slide, x + 0.46, step_y + 0.03, w - 0.58, 0.24, title, 13, True, title_color)
        textbox(slide, x + 0.46, step_y + 0.26, w - 0.58, 0.22, detail, 11, False, MUTED)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW,
                Inches(x + w / 2 - 0.07),
                Inches(step_y + step_h + 0.01),
                Inches(0.14),
                Inches(0.12),
            )
            solid(arrow, head_fill)
        step_y += step_h + gap
    result_y = step_y - gap + 0.06
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(result_y), Inches(w), Inches(0.36))
    solid(chip, result_fill)
    textbox(slide, x + 0.14, result_y, w - 0.24, 0.36, result, 12, True, result_color, PP_ALIGN.CENTER)
    return result_y + 0.36


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "UBC 与连接级拥塞控制"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide)
    banner(slide, 0.86, "同一目的上开很多条流时，连接级 CC 每条流各有一份窗口；UBC 把这些流收成目的加 VL 的一份速率。")

    left_steps = [
        ("建一条连接", "TCP 会话、RoCE 的 QP、RTP 的 TP Channel"),
        ("这条连接自带一份状态", "拥塞窗口，或这条连接自己的发送速率"),
        ("反馈回到这条连接", "ACK、TPACK，或标着这条连接的 CNP"),
        ("只改这一份窗口或速率", "其他连接保持自己的状态"),
    ]
    right_steps = [
        ("报文直接发出", "简易传输，一个包带完一条事务"),
        ("状态按目的和 VL 取", "键是 {DCNA[15:0], VL}"),
        ("反馈来自下层", "交换机打 FECN，接收端回独立 CNP"),
        ("按这份共享速率发送", "逐跳重传和信用留在链路层"),
    ]
    column(
        slide, 0.42, 1.38, 6.05,
        "连接级 CC", "一条连接，一份拥塞状态",
        NAVY, left_steps, SOFT, NAVY,
        "去往同一目的的 N 条连接  =  N 份状态",
        SOFT, NAVY,
    )
    column(
        slide, 6.86, 1.38, 6.05,
        "UBC 的 CTP", "旧名 UBC_CTP，现名 UB_CTP",
        TEAL, right_steps, PALE, TEAL,
        "同一目的、同一 VL 的流  =  1 份速率",
        PALE, TEAL,
    )

    rows = [
        ["", "连接级 CC", "UBC 的 CTP"],
        ["状态挂在哪", "每一条连接", "目的 CNA 的低 16 位，再加 VL"],
        ["反馈回到哪", "回到发出它的那条连接", "独立 CNP，按目的和 VL 归到同一份"],
        ["同一目的上的多条流", "各调各的窗口或速率", "同一 VL 上共用一份速率"],
        ["可靠闭环放在哪", "端到端重传和拥塞窗口在同一条连接上", "逐跳重传和信用在链路层，速率在这份共享状态上"],
    ]
    add_table(
        slide,
        0.42,
        5.28,
        [2.15, 5.15, 5.20],
        rows,
        [0.30, 0.36, 0.36, 0.36, 0.40],
        ["head", "key", "body", "key", "body"],
    )
    slide.notes_slide.notes_text_frame.text = (
        "这页对照的是拥塞状态的粒度。"
        "连接级 CC 包括 TCP、RoCE 按 QP、以及 UB 可靠传输 RTP 按 TP Channel：每条连接一份窗口或速率，确认回到这条连接，只更新这一份。"
        "UBC 是简易传输 CTP。hcomm 里 COMM_PROTOCOL_UBC_CTP 与 COMM_PROTOCOL_UB_CTP 是同一个枚举值 4；可靠传输的旧名 UBG 现为 UB_RTP。"
        "昇腾 950 架构白皮书 4.6.1：CTP 不支持端到端可靠重传，带宽按 9 个 Port 计，简易传输下也可以多路径；RTP 支持端到端重传，可靠带宽按 4 个 Port 计。"
        "术语表：CTP 借助下层协议共同提供可靠和拥塞控制。"
        "CTP 拥塞上下文的键是 {DCNA[15:0], VL}。目的 CNA 为 24 位时，索引用的仍是低 16 位。同一目的、同一 VL 上的 CTP 流共用这一份速率。"
        "CTP 没有 TPACK。LDCP、CAQM 那种靠传输层确认回传窗口的算法在 RTP 上闭合。CTP 的拥塞信号是交换机 FECN 和接收端独立 CNP，发送端调的是这份共享速率。"
        "可靠性由链路层逐跳重传和信用流控承担。"
    )

    out = "/workspace/docs/ubc-ctp-cc-compare.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
