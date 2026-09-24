#!/usr/bin/env python3
"""RSI principle + communication data-flow schematic."""

from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Emu

W, H = 1920, 1080
OUT = Path("/workspace/docs")
DIR = OUT / "rsi-comm-flow"
FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

BG = (246, 244, 239)
INK = (27, 36, 48)
MUTED = (92, 102, 112)
TEAL = (15, 110, 107)
CORAL = (196, 73, 58)
GOLD = (184, 138, 58)
OK = (42, 122, 75)
NAVY = (28, 49, 68)
CARD = (255, 255, 255)
LINE = (220, 214, 204)
SOFT = (237, 233, 224)
PURPLE = (102, 72, 140)


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def canvas():
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    for i in range(0, W, 48):
        d.line([(i, 0), (i, H)], fill=(236, 232, 224), width=1)
    for j in range(0, H, 48):
        d.line([(0, j), (W, j)], fill=(236, 232, 224), width=1)
    return img


def tw(d, t, f):
    return int(d.textlength(t, font=f))


def wrap(d, text, fnt, max_w):
    lines = []
    for para in text.split("\n"):
        buf = ""
        for ch in para:
            if d.textlength(buf + ch, font=fnt) <= max_w:
                buf += ch
            else:
                if buf:
                    lines.append(buf)
                buf = ch
        if buf:
            lines.append(buf)
    return lines


def shadow(base, box, r=16):
    x0, y0, x1, y1 = box
    layer = Image.new("RGBA", base.size, (0, 0, 0, 0))
    ImageDraw.Draw(layer).rounded_rectangle(
        (x0 + 3, y0 + 5, x1 + 3, y1 + 7), radius=r, fill=(20, 24, 30, 26)
    )
    blur = layer.filter(ImageFilter.GaussianBlur(7))
    tmp = base.convert("RGBA")
    tmp.alpha_composite(blur)
    base.paste(tmp.convert("RGB"))


def card(img, box, fill=CARD, r=16, outline=LINE, w=1, drop=True):
    if drop:
        shadow(img, box, r)
    ImageDraw.Draw(img).rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=w)


def pill(d, xy, text, bg, fg=(255, 255, 255), fnt=None):
    fnt = fnt or font(15)
    x, y = xy
    box = (x, y, x + tw(d, text, fnt) + 28, y + 32)
    d.rounded_rectangle(box, radius=14, fill=bg)
    d.text((x + 14, y + 6), text, font=fnt, fill=fg)
    return box


def arrow(d, p1, p2, color, width=4, head=13):
    d.line([p1, p2], fill=color, width=width)
    ang = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    x, y = p2
    for da in (2.55, -2.55):
        d.line(
            [(x, y), (x - head * math.cos(ang + da), y - head * math.sin(ang + da))],
            fill=color,
            width=width,
        )


def numbered_dot(d, xy, n, color):
    x, y = xy
    d.ellipse((x - 14, y - 14, x + 14, y + 14), fill=color)
    d.text((x, y), str(n), font=font(15), fill=(255, 255, 255), anchor="mm")


def header(img, kicker, title, claim):
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 8), fill=TEAL)
    d.text((56, 24), kicker, font=font(17), fill=TEAL)
    d.text((56, 52), title, font=font(36), fill=INK)
    d.text((56, 104), claim, font=font(20), fill=MUTED)
    d.line([(56, 140), (W - 56, 140)], fill=LINE, width=1)


def footer(img, text):
    d = ImageDraw.Draw(img)
    d.line([(56, 1040), (W - 56, 1040)], fill=LINE, width=1)
    d.text((56, 1052), text, font=font(15), fill=MUTED)
    d.text((W - 56, 1052), "RSI 通信数据流", font=font(15), fill=TEAL, anchor="rt")


def slide_flow():
    img = canvas()
    header(
        img,
        "RSI 原理  ·  通信视角",
        "模型自己设计实验、造数据、跑训推、找缺陷，再开下一轮",
        "传统是一条单向流水线。RSI 把训练网、推理网、环境网绑进同一个环。",
    )
    d = ImageDraw.Draw(img)

    # legend
    legs = [
        (GOLD, "① 控制面  作业/配置"),
        (NAVY, "②③ 训练面  样本/梯度"),
        (TEAL, "④ 推理面  Token/KV"),
        (CORAL, "⑤ 环境面  沙箱/工具"),
        (PURPLE, "⑥⑦ 回流  权重/轨迹/评测"),
    ]
    x = 56
    for c, t in legs:
        box = pill(d, (x, 156), t, c)
        x = box[2] + 12

    # controller
    card(img, (610, 208, 1310, 318), fill=TEAL, outline=TEAL)
    d.text((960, 242), "RSI 外环控制器", font=font(26), fill=(255, 255, 255), anchor="mm")
    d.text((960, 284), "设计实验  ·  搭建流程  ·  定位缺陷  ·  决定下一轮", font=font(16), fill=(210, 232, 230), anchor="mm")

    # three stations
    stations = [
        (80, 400, 600, 760, GOLD, "造数据", "样本工厂", [
            "构造训练/评测数据",
            "写轨迹、标对错边界",
            "供给训练和评估",
        ]),
        (660, 400, 1260, 760, NAVY, "跑训练", "训练集群", [
            "集合通信（梯度/专家）",
            "打检查点",
            "产出新权重",
        ]),
        (1320, 400, 1840, 760, CORAL, "跑推理 + 环境", "推理实例 / 沙箱 / 工具", [
            "吐 Token、搬 KV",
            "Agent 调沙箱、EDA、代码",
            "采轨迹、算回报",
        ]),
    ]
    for x0, y0, x1, y1, color, title, sub, bullets in stations:
        card(img, (x0, y0, x1, y1), outline=color, w=2)
        d.rectangle((x0, y0, x1, y0 + 8), fill=color)
        d.text((x0 + 28, y0 + 28), title, font=font(26), fill=INK)
        d.text((x0 + 28, y0 + 70), sub, font=font(16), fill=color)
        y = y0 + 118
        for b in bullets:
            d.ellipse((x0 + 32, y + 8, x0 + 44, y + 20), fill=color)
            d.text((x0 + 56, y), b, font=font(18), fill=INK)
            y += 38

    # mini ④⑤ path inside infer box
    d.rounded_rectangle((1350, 638, 1810, 736), radius=12, fill=SOFT)
    d.rounded_rectangle((1364, 658, 1474, 716), radius=10, fill=TEAL)
    d.text((1419, 687), "推理", font=font(16), fill=(255, 255, 255), anchor="mm")
    d.rounded_rectangle((1524, 658, 1634, 716), radius=10, fill=INK)
    d.text((1579, 687), "Agent", font=font(16), fill=(255, 255, 255), anchor="mm")
    d.rounded_rectangle((1684, 658, 1796, 716), radius=10, fill=CORAL)
    d.text((1740, 687), "沙箱/工具", font=font(16), fill=(255, 255, 255), anchor="mm")
    arrow(d, (1474, 678), (1520, 678), TEAL, 3, 9)
    arrow(d, (1520, 698), (1474, 698), TEAL, 3, 9)
    numbered_dot(d, (1497, 652), 4, TEAL)
    arrow(d, (1634, 678), (1680, 678), CORAL, 3, 9)
    arrow(d, (1680, 698), (1634, 698), CORAL, 3, 9)
    numbered_dot(d, (1657, 652), 5, CORAL)

    # arrows from controller
    arrow(d, (780, 318), (340, 400), GOLD, 4)
    numbered_dot(d, (520, 350), 1, GOLD)
    arrow(d, (960, 318), (960, 400), GOLD, 4)
    numbered_dot(d, (990, 360), 1, GOLD)
    arrow(d, (1140, 318), (1580, 400), GOLD, 4)
    numbered_dot(d, (1400, 350), 1, GOLD)

    # 2 data -> train
    arrow(d, (600, 540), (660, 540), NAVY, 5)
    numbered_dot(d, (630, 512), 2, NAVY)
    d.text((630, 568), "样本", font=font(14), fill=NAVY, anchor="mm")

    # 3 train internal hint
    d.rounded_rectangle((820, 690, 1100, 736), radius=10, fill=SOFT)
    d.text((960, 713), "③ 集群内梯度 / 专家通信", font=font(15), fill=NAVY, anchor="mm")

    # 6 weights train -> infer
    arrow(d, (1260, 500), (1320, 500), PURPLE, 5)
    numbered_dot(d, (1290, 472), 6, PURPLE)
    d.text((1290, 528), "新权重", font=font(14), fill=PURPLE, anchor="mm")

    # eval bar
    card(img, (80, 800, 1840, 920), fill=SOFT, drop=False)
    d.text((120, 828), "评估 / 定位缺陷", font=font(22), fill=PURPLE)
    d.text((120, 872), "⑦ 轨迹、回报、评测分、失败日志 回流控制器  →  开下一轮    （云栖口径：Qwen3.8-Max 无人值守 33 轮 / 月级）", font=font(18), fill=INK)
    arrow(d, (340, 760), (340, 800), PURPLE, 4)
    arrow(d, (960, 760), (960, 800), PURPLE, 4)
    arrow(d, (1580, 760), (1580, 800), PURPLE, 4)
    numbered_dot(d, (200, 860), 7, PURPLE)

    # loop back outside all boxes
    arrow(d, (1720, 860), (1888, 860), PURPLE, 4)
    arrow(d, (1888, 860), (1888, 248), PURPLE, 3)
    arrow(d, (1888, 248), (1310, 248), PURPLE, 3)
    d.text((1864, 540), "下一轮", font=font(14), fill=PURPLE, anchor="rt")

    footer(img, "图1  ·  RSI 原理与通信数据流")
    return img


def slide_challenge():
    img = canvas()
    header(
        img,
        "对网络的判断",
        "单条流大多是旧题；新挑战是三面同时在线、互相等待",
        "RSI 没有发明一种新的包格式。它把训练、推理、环境、控制绑成一条不能断的环。",
    )
    d = ImageDraw.Draw(img)

    # verdict cards
    old_new = [
        (OK, "不是新题", [
            ("训练集合通信", "还是梯度 / 专家 AlltoAll，HPN 那类东西"),
            ("推理 Token / KV", "还是首字、流畅、搬运，TPN 那类东西"),
            ("检查点进存储", "还是大块顺序写，存储网旧题"),
        ]),
        (GOLD, "半新：旧流换了耦合", [
            ("权重热更新", "训练一边产、推理一边吃，要持续推"),
            ("轨迹 / logprobs 回传", "中等大象流，训练等这条回流"),
            ("训推共集群", "两张网的隔离和抢路变成日常"),
        ]),
        (CORAL, "是新题", [
            ("环境 / 沙箱南北向", "创建密集、时延长尾、工具往返"),
            ("控制面进关键路径", "实验编排失败，整轮作废"),
            ("月级不断环", "稳定对象从“一个作业”变成“一条环”"),
        ]),
    ]
    x = 56
    for color, title, rows in old_new:
        card(img, (x, 168, x + 580, 620), outline=color, w=2)
        d.rectangle((x, 168, x + 580, 176), fill=color)
        d.text((x + 28, 200), title, font=font(26), fill=INK)
        y = 270
        for h, b in rows:
            d.ellipse((x + 36, y + 8, x + 48, y + 20), fill=color)
            d.text((x + 64, y), h, font=font(20), fill=INK)
            d.text((x + 64, y + 34), b, font=font(16), fill=MUTED)
            y += 100
        x += 604

    # mapping to Li Feifei
    card(img, (56, 648, 1864, 1008), r=18)
    d.text((92, 676), "所以，对网络规划多了什么要求", font=font(22), fill=TEAL)
    reqs = [
        ("一张环，不是三张孤岛", "训练、推理、环境、控制要能协同，又不能互相打死。隔离和放行要按环上的角色来。"),
        ("忙时先保外环不断", "沙箱长尾、工具超时，不能把训练卡和推理卡一起拖死。该丢环境、不该堵集合通信。"),
        ("控制面也要当生产网", "作业图、评测回流、权重就绪信号，丢了就是一轮实验作废。小消息，高可靠。"),
        ("弹性对象换成沙箱和环", "云栖对 RSI 底座的原话是：超大规模稳定、毫秒级弹性、百万级并发，训推环境统一协同。"),
    ]
    y = 728
    for i, (h, b) in enumerate(reqs):
        col = i % 2
        row = i // 2
        xx = 92 + col * 880
        yy = y + row * 120
        d.rounded_rectangle((xx, yy, xx + 840, yy + 104), radius=12, fill=SOFT)
        d.text((xx + 20, yy + 14), h, font=font(18), fill=TEAL)
        lines = wrap(d, b, font(16), 790)
        ly = yy + 48
        for line in lines:
            d.text((xx + 20, ly), line, font=font(16), fill=INK)
            ly += 24

    footer(img, "图2  ·  RSI 给网络增加了什么")
    return img


def pack(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for p in paths:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(p), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    out = OUT / "rsi-comm-flow.pptx"
    prs.save(out)
    return out


def main():
    DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, fn in (("01-rsi-flow", slide_flow), ("02-net-challenge", slide_challenge)):
        img = fn()
        p = DIR / f"{name}.png"
        img.save(p, "PNG", optimize=True)
        paths.append(p)
        print("wrote", p)
    print("pptx", pack(paths))


if __name__ == "__main__":
    main()
