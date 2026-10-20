#!/usr/bin/env python3
"""Editable deck: 1.6T to 10T, then TTFT 10s / TPOT 10ms / TPOT 30ms."""

from __future__ import annotations

import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
HEAD = "0E2A47"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    h = hex_color
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def set_cell_fill(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid_fill = OxmlElement("a:solidFill")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", hex_color)
    solid_fill.append(srgb)
    tcPr.append(solid_fill)


def set_cell_borders(cell, hex_color="D6DCE2", width="6350"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for child in list(tcPr):
            if child.tag == qn(f"a:{edge}"):
                tcPr.remove(child)
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", width)
        fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_color)
        fill.append(srgb)
        ln.append(fill)
        tcPr.append(ln)


def write_cell(cell, text, size, bold, color, align):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        ln_spc = OxmlElement("a:lnSpc")
        spc = OxmlElement("a:spcPct")
        spc.set("val", "100000")
        ln_spc.append(spc)
        pPr.append(ln_spc)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def add_table(slide, x, y, widths, rows, heights, kinds):
    """kinds: head / body / key / sum, one per row."""
    n_row, n_col = len(rows), len(widths)
    shape = slide.shapes.add_table(n_row, n_col, Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights)))
    table = shape.table
    table.horz_banding = False
    table.vert_banding = False
    tbl = shape._element.graphic.graphicData.tbl
    tbl_pr = tbl.find(qn("a:tblPr"))
    if tbl_pr is None:
        tbl_pr = tbl.tblPr
    for child in list(tbl_pr):
        if child.tag == qn("a:tableStyleId"):
            tbl_pr.remove(child)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    style = {
        "head": (HEAD, WHITE, True, PP_ALIGN.LEFT),
        "body": (WHITE, INK, False, PP_ALIGN.LEFT),
        "key": (PALE, TEAL, True, PP_ALIGN.LEFT),
        "sum": (SOFT, NAVY, True, PP_ALIGN.LEFT),
        "note": (GOLD_BG, INK, False, PP_ALIGN.LEFT),
    }
    for r, (row, h, kind) in enumerate(zip(rows, heights, kinds)):
        table.rows[r].height = Inches(h)
        fill, color, bold, align = style[kind]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            set_cell_fill(cell, fill if c or kind == "head" else SOFT)
            if kind == "head":
                set_cell_fill(cell, HEAD)
            set_cell_borders(cell)
            this_bold = True if c == 0 and kind != "head" else bold
            this_color = NAVY if c == 0 and kind == "body" else color
            if kind in ("key", "sum", "head", "note"):
                this_color = color
                this_bold = bold or c == 0
            write_cell(cell, text, 12 if kind != "head" else 12, this_bold, this_color, align)
    return shape


def chrome(slide, page, title, kicker):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.16, 10.6, 0.28, kicker, 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.42, 0.42, 12.5, 0.42, title, 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 11.4, 0.18, 1.5, 0.26, f"{page}  /  4", 12, False, MUTED, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.20,
        12.5,
        0.24,
        "10T 为外推    ·    单请求一份 128 卡副本    ·    访存与通信取峰值的 80%    ·    计算按 Cube 16 行对齐",
        11,
        False,
        MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def banner(slide, y, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(y), Inches(12.5), Inches(0.48))
    solid(shape, PALE)
    textbox(slide, 0.58, y, 12.18, 0.48, text, 14, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)


def model():
    E, L, h, dff, topk = 2469, 61, 7168, 3072, 6
    n_csa, n_hca = 30, 31
    N, TP = 128, 8
    DP = N // TP
    w_one = 3 * h * dff * L
    rest = 1.6e12 - 384 * w_one
    total = E * w_one + rest
    routed_bytes = E * w_one * 0.5
    shared = w_one  # FP8 bytes
    attn = 19.35e9
    peak = 486e12
    hbm = 4e12 * 0.80
    bus = 400e9 * 0.80
    p2p = 50e9 * 0.80
    param = 50e9 * 0.80
    moe = topk * L * 6 * h * dff + L * 6 * h * dff
    proj = attn * 2
    bpt = n_csa * (576 / 4 + 64 / 4) + n_hca * (576 / 128)
    fixed = L * 128 * 576
    room = N * (96e9 - routed_bytes / N - shared - attn / TP - 6e9)

    def align(M):
        if M <= 0:
            return 1.0
        return M / (math.ceil(M / 16.0) * 16)

    def prefill(S):
        idx = n_csa * 2048 * S * S
        hca = n_hca * 1024 * S * S
        core = n_csa * 262144 * (1024 + 128) * S
        flops = idx + hca + core + proj * S + moe * S
        t_c = flops / (N * peak)
        t_w = (routed_bytes / N + shared + attn / TP) / hbm
        t_ship = S * bpt / (N * param)
        return {
            "t": max(t_c, t_w) + t_ship,
            "idx": idx / (N * peak),
            "hca": hca / (N * peak),
            "core": core / (N * peak),
            "proj": proj * S / (N * peak),
            "moe": moe * S / (N * peak),
            "tc": t_c,
            "tw": t_w,
            "ship": t_ship,
        }

    lo, hi = 1.0, 5e7
    for _ in range(70):
        mid = (lo + hi) / 2
        if prefill(mid)["t"] > 10:
            hi = mid
        else:
            lo = mid
    S10 = lo

    def decode(BS, S):
        frac = 1 - math.exp(-topk * BS / E)
        e = E / N
        std = math.sqrt(max(e * frac * (1 - frac), 0))
        e_hot = min(e, e * frac + 2 * std)
        tok = (BS * topk / E) / frac if frac > 1e-9 else 1
        util_e = align(tok)
        util_a = align(BS / DP)
        score = n_csa * 4096 * S + n_hca * 2048 * S + n_csa * 262144 * (1024 + 128)
        t_w = (e_hot * w_one * 0.5 + shared) / hbm
        kv = n_csa * (S / 4) * 64 + n_hca * (S / 128) * 576 + n_csa * 1024 * 576 + fixed
        t_kv = BS * kv / (N * hbm)
        t_wa = (attn / TP) / hbm
        t_ca = BS * (score + proj) / (N * peak * util_a)
        t_cm = BS * moe / (N * peak * util_e)
        pairs = (e_hot / e) * (BS * topk / N)
        t_ep = pairs * 2 * h * L / bus
        toks = math.ceil(BS / DP)
        ring = 2 * (TP - 1) / TP * (toks * h * 2)
        t_ar = ring * L / p2p
        t_attn = max(t_ca, t_kv + t_wa, t_ar)
        t_moe = max(t_cm, t_w, t_ep)
        return {
            "t": t_attn + t_moe,
            "frac": frac,
            "e_hot": e_hot,
            "tok": tok,
            "ue": util_e,
            "ua": util_a,
            "tw": t_w,
            "tcm": t_cm,
            "tca": t_ca,
            "tkv": t_kv,
            "tep": t_ep,
            "tar": t_ar,
            "t_attn": t_attn,
            "t_moe": t_moe,
        }

    def search(S, limit):
        cap = int(max(room, 0) // (S * bpt + fixed))
        if cap < 1 or decode(1, S)["t"] > limit:
            return 0, cap
        lo_b, hi_b, best = 1, cap, 0
        while lo_b <= hi_b:
            mid = (lo_b + hi_b) // 2
            if decode(mid, S)["t"] <= limit:
                best = mid
                lo_b = mid + 1
            else:
                hi_b = mid - 1
        return best, cap

    return {
        "S10": S10,
        "prefill": prefill(S10),
        "d10_4k": decode(*search(4096, 0.01)[:1], 4096) if False else decode(search(4096, 0.01)[0], 4096),
        "bs10": {S: search(S, 0.01)[0] for S in (4096, 131072, 10**6, int(S10))},
        "bs30": {S: search(S, 0.03)[0] for S in (4096, 131072, 10**6, int(S10))},
        "cap": {S: int(room // (S * bpt + fixed)) for S in (4096, 131072, 10**6, int(S10))},
        "d30_4k": decode(search(4096, 0.03)[0], 4096),
        "d30_1m": decode(search(10**6, 0.03)[0], 10**6),
        "d10_long": decode(search(int(S10), 0.01)[0], int(S10)),
        "room": room,
        "bpt": bpt,
        "fixed": fixed,
        "routed_tb": routed_bytes / 1e12,
        "per_gpu_gb": routed_bytes / N / 1e9,
        "experts_per_gpu": E / N,
        "shared_gb": shared / 1e9,
        "attn_gb": attn / 1e9,
        "attn_gpu_gb": attn / TP / 1e9,
        "full_scan": (routed_bytes / N + shared) / hbm,
        "total": total,
        "moe": moe,
        "proj": proj,
        "E": E,
    }


def ms(v):
    return f"{v * 1e3:.1f} 毫秒" if v < 0.1 else f"{v * 1e3:.0f} 毫秒"


def sec(v):
    if v >= 9.5:
        return f"{v:.1f} 秒"
    return f"{v:.2f} 秒"


def main():
    m = model()
    S = m["S10"]
    p = m["prefill"]
    d10 = m["d10_4k"]
    d30 = m["d30_4k"]
    dlong = m["d10_long"]
    d1m = m["d30_1m"]
    bs10 = m["bs10"]
    bs30 = m["bs30"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "10T DeepSeek：结构外推与 TTFT / TPOT"
    blank = prs.slide_layouts[6]

    # ----- slide 1 -----
    s1 = prs.slides.add_slide(blank)
    chrome(s1, 1, "同一骨架，只加路由专家", "模型结构    ·    1.6T 论文到 10T 外推")
    textbox(
        s1,
        0.42,
        0.90,
        12.5,
        0.32,
        "1.6T 是 V4-Pro 论文原模型。10T 沿用层数、隐藏和中间维，用路由专家个数把总参数对齐到 10^13。",
        14,
        False,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    rows = [
        ["项目", "1.6T  论文 V4-Pro", "10T  外推", "备注"],
        ["总参数", "1.600×10^12", "1.000×10^13", "DeepSeek 未发布 10T，本档是外推"],
        ["路由专家", "384", "2469", "×6.43。增量全部记在专家个数上"],
        ["激活参数", "4.9×10^10", "4.9×10^10", "每 token 仍激活 6 个路由专家 + 1 个共享"],
        ["激活比", "3.06%", "0.49%", "专家表变大，单 token 激活不变"],
        ["FP4 专家表", "0.774 TB", "4.975 TB", "单个专家全层 2.015 GB"],
        ["层 / 隐藏 / 中间维", "61 / 7168 / 3072", "同左", "外推不动这三项"],
        ["注意力", "CSA m=4，topk=1024\nHCA m′=128，滑窗 128", "同左", "30 层 CSA、31 层 HCA。KV 仍是 4939.5 B/token"],
        ["单 token 专家通信", "7.87×10^6 字节", "同左", "6 × 3 × 7168 × 61，与专家总数无关"],
        ["640 GB 卡下限", "2 张", "8 张", "只按 FP4 专家表容量。时延分析改用 96 GB × 128 卡"],
    ]
    kinds = ["head"] + ["body"] * 9
    heights = [0.36] + [0.50] * 9
    add_table(s1, 0.42, 1.28, [2.15, 3.15, 2.55, 4.65], rows, heights, kinds)
    note = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(6.28), Inches(12.5), Inches(0.80))
    solid(note, GOLD_BG)
    textbox(
        s1,
        0.58,
        6.32,
        12.18,
        0.72,
        "外推式：总参数 = 专家数 × 4.030×10^9 + 5.260×10^10。单个专家 = 3 × 7168 × 3072 × 61。\n"
        "依据是 V3 到 V4-Pro 的增量主要进专家表。层数、隐藏、中间维一起放大的另一案约 8.9T，注意力超参定不下来，不收进本表。",
        13,
        False,
        GOLD,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    s1.notes_slide.notes_text_frame.text = (
        "1.6T 来自 V4-Pro 论文 arXiv:2606.19348：61 层，隐藏 7168，384 个路由专家，中间维 3072，每 token 激活 6 个，激活参数 49B。"
        "单个路由专家在全部层上是 3×7168×3072×61 = 4.030×10^9。"
        "384 个专家为 1.547×10^12，总参数 1.600×10^12 减去这一项，其余 5.260×10^10 含共享专家、注意力、mHC、词表和 MTP。"
        "10T 把专家数解成 2469，总参数 = 2469×4.030×10^9 + 5.260×10^10 = 1.000×10^13。"
        "FP4 容量 = 2469×4.030×10^9×0.5 字节 = 4.975 TB。640 GB 卡下限是 ceil(4.975 TB / 640 GB) = 8。"
        "单 token 通信量 6×3×7168×61 = 7.87×10^6 字节，专家总数变了，每 token 仍只派 6 个，所以通信量不变。"
        "后面三页不用 640 GB 这套卡，改用 96 GB、128 卡超节点。"
    )

    # ----- slide 2 -----
    s2 = prs.slides.add_slide(blank)
    chrome(s2, 2, "10 秒对应的最大输入是 2.08×10^6", "Prefill    ·    TTFT ≤ 10 秒")
    banner(
        s2,
        0.92,
        f"一条请求放在一个 128 卡副本上。计算 {sec(p['tc'])}，读权重 {ms(p['tw'])}，送 KV {ms(p['ship'])}。S² 占 {sec(p['idx']+p['hca'])}，线性项占 {sec(p['core']+p['proj']+p['moe'])}。",
    )
    textbox(s2, 0.42, 1.46, 6.2, 0.26, "并行策略", 13, True, ACCENT)
    par_rows = [
        ["", "Attention", "MoE"],
        ["并行", "机内 TP=8\n16 个节点做 DP=16", "超节点 EP=128\n一份模型铺满 128 卡"],
        ["通信域", "8 卡 fullmesh\n有效 40 GB/s", "128 卡 D2D 总线，有效 320 GB/s\nKV 走参数面，有效 40 GB/s"],
        ["每卡权重", "注意力等 19.35 GB / 8\n= 2.42 GB FP8", f"路由专家 {m['experts_per_gpu']:.1f} 个，{m['per_gpu_gb']:.1f} GB FP4\n共享专家 {m['shared_gb']:.2f} GB，每卡一份"],
        ["本阶段角色", "S² 来自 indexer 和 HCA\n决定输入能写多长", "专家 GEMM 随 token 线性\n到这个长度约占 1.89 秒"],
    ]
    add_table(
        s2,
        0.42,
        1.72,
        [1.55, 5.40, 5.55],
        par_rows,
        [0.30, 0.56, 0.62, 0.56, 0.56],
        ["head", "body", "body", "body", "body"],
    )
    textbox(s2, 0.42, 4.42, 12.5, 0.26, "10 秒怎么落到这个输入长度", 13, True, ACCENT)
    calc_rows = [
        ["这一笔", "算法", f"S = {S/1e6:.2f}×10^6 时"],
        ["CSA indexer", "30 层 × 2048 × S²，除以 128 × 486×10^12", sec(p["idx"])],
        ["HCA", "31 层 × 1024 × S²，除以同一算力", sec(p["hca"])],
        ["专家 GEMM", "每 token 5.64×10^10 FLOP × S", sec(p["moe"])],
        ["投影", "每 token 3.87×10^10 FLOP × S", sec(p["proj"])],
        ["稀疏注意力", "每 token 看 1024 条压缩 KV + 128 滑窗", sec(p["core"])],
        ["读权重 + 送 KV", "每卡约 45 GB / 3.2 TB/s；KV 10.3 GB / (128×40 GB/s)", f"{ms(p['tw'])} + {ms(p['ship'])}"],
        ["合计", "计算和读权重取长的一项，再加送 KV。S 再大，前两行先超", "10.0 秒"],
    ]
    add_table(
        s2,
        0.42,
        4.68,
        [1.70, 7.55, 3.25],
        calc_rows,
        [0.28, 0.28, 0.28, 0.28, 0.28, 0.28, 0.28, 0.32],
        ["head", "body", "body", "body", "body", "body", "body", "sum"],
    )
    s2.notes_slide.notes_text_frame.text = (
        "单请求不跨副本。增加 Prefill 节点是多放几份 128 卡副本，提高并发，不缩短这一条的 TTFT。"
        "注意力在机内 8 卡做 TP，16 个节点是 16 路 DP。MoE 把 2469 个专家铺到 128 卡，每卡约 19.3 个。"
        "时间模型：max(计算, 读权重) + 把 KV 送到 Decode。"
        "S 到两百万时，每卡上的矩阵行数远超 16，Cube 对齐效率按 1 计。"
        f"在 S={S:.0f}：indexer {p['idx']:.2f}s，HCA {p['hca']:.2f}s，专家 {p['moe']:.2f}s，投影 {p['proj']:.2f}s，稀疏注意力 {p['core']:.2f}s，合计计算 {p['tc']:.2f}s。"
        f"读权重 {p['tw']*1e3:.1f} ms，KV 搬运 {p['ship']*1e3:.1f} ms。"
        "若只留 S² 两项，10 秒能写到约 2.58×10^6；线性项占掉约 3.5 秒，输入收到 2.08×10^6。"
        "带宽 80%：访存 3.2 TB/s，参数面 40 GB/s。这两笔相对 10 秒可以忽略，所以 80% 不改变最大输入。"
    )

    # ----- slide 3 -----
    s3 = prs.slides.add_slide(blank)
    chrome(s3, 3, "10 毫秒、短输入，全局批 231", "Decode    ·    TPOT ≤ 10 毫秒")
    banner(
        s3,
        0.92,
        "10 毫秒里 9.2 毫秒在读被叫醒的专家，0.8 毫秒在读注意力权重。全部叫醒要 13.4 毫秒，这一档做不到。两路通信都不到 1 毫秒。",
    )
    textbox(s3, 0.42, 1.46, 6.0, 0.26, "并行策略", 13, True, ACCENT)
    textbox(s3, 7.15, 1.46, 5.7, 0.26, "短输入（4K）这一步怎么加到 10 毫秒", 13, True, ACCENT)
    par3 = [
        ["", "Attention", "MoE"],
        ["并行", "TP=8，DP=16\n每卡 token 数 = BS / 16", "EP=128\n2469 个专家，每卡 19.3 个"],
        ["通信", "每层机内 allreduce\n这一步 0.6 毫秒", "最忙卡 BF16 combine\n这一步 0.02 毫秒"],
        ["访存", "注意力权重 2.42 GB\n这一步 0.8 毫秒", "命中专家这一步 9.2 毫秒\n全部叫醒要 13.4 毫秒"],
        ["计算", "Cube 行数 = BS / 16\n不足 16 按 16 行砖", "行数 = 命中专家上的 token 数\n同样按 16 行砖"],
    ]
    add_table(
        s3,
        0.42,
        1.74,
        [1.15, 2.55, 2.70],
        par3,
        [0.30, 0.78, 0.78, 0.82, 0.78],
        ["head", "body", "body", "key", "body"],
    )
    step3 = [
        ["笔", "白话", "时间"],
        ["命中", "1 − exp(−6×231 / 2469) = 43%\n最忙卡按均值 + 2σ，约 12.6 个", "—"],
        ["读专家", "12.6 × 2.015 GB + 共享 4.03 GB\n除以 3.2 TB/s", "9.2 毫秒"],
        ["读注意力", "2.42 GB / 3.2 TB/s\n长于机内 allreduce 的 0.6 毫秒", "0.8 毫秒"],
        ["专家计算", "每个命中专家约 1.3 个 token\n16 行砖，效率 8%", "2.6 毫秒"],
        ["专家通信", "最忙卡送出的 BF16 combine / 320 GB/s", "0.02 毫秒"],
        ["相加", "注意力侧取 0.8 毫秒，专家侧取 9.2 毫秒", "10.0 毫秒"],
    ]
    add_table(
        s3,
        7.05,
        1.74,
        [1.20, 3.50, 1.15],
        step3,
        [0.30, 0.60, 0.60, 0.54, 0.52, 0.46, 0.44],
        ["head", "body", "key", "body", "body", "body", "sum"],
    )
    # result chips
    chips = [
        ("4K – 32K", "231", "每卡 1.8 条"),
        ("128K", "229", "仍是读权重"),
        ("1×10^6", "214", "KV 开始占额度"),
        ("2.08×10^6", "179", f"读权重 {dlong['tw']*1e3:.1f} 毫秒\n注意力计算 {dlong['tca']*1e3:.1f} 毫秒"),
    ]
    for i, (cap, num, sub) in enumerate(chips):
        x = 0.42 + i * 3.20
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.40), Inches(3.05), Inches(1.62))
        solid(card, SOFT, LINE)
        textbox(s3, x + 0.16, 5.48, 2.75, 0.26, cap, 12, True, MUTED)
        textbox(s3, x + 0.16, 5.74, 2.75, 0.50, num, 28, True, TEAL)
        textbox(s3, x + 0.16, 6.28, 2.75, 0.62, sub if "\n" in sub else "全局批\n" + sub, 12, False, INK)
    s3.notes_slide.notes_text_frame.text = (
        "Decode 与 Prefill 用同一份 128 卡副本：注意力 TP=8、DP=16，MoE EP=128。"
        "加 Decode 节点是多副本，单副本的 TPOT 不变。"
        "一步 = 注意力侧最长的一笔 + 专家侧最长的一笔。"
        "注意力侧在计算、读 KV、机内 allreduce 里取最大。专家侧在计算、读权重、EP combine 里取最大。"
        "命中率 = 1 − exp(−6×BS / 2469)。BS=231 时为 43%。"
        "最忙卡专家数 = 均值 + 2 个标准差，约 12.6，封顶 19.3。"
        "读权重 (12.6×2.015 GB + 4.03 GB) / 3.2 TB/s = 9.2 ms。"
        "全部路由专家加共享专家是 13.4 ms，超过 10 ms，所以不能把专家表一次扫完。"
        "专家计算受 Cube 对齐拖累：1.3 token 按 16 行砖，效率 8%，时间 2.6 ms，仍短于读权重。"
        "机内 allreduce：ceil(231/16)=15 条，每条 7168×2 字节，环上每层 2×7/8，61 层除以 40 GB/s = 0.57 ms。"
        "输入加长后，注意力计算和 KV 读取变多，读权重能占用的额度变少，批从 231 降到 179。"
        f"在最大输入上，读权重 {dlong['tw']*1e3:.1f} ms，注意力计算 {dlong['tca']*1e3:.1f} ms，相加仍约 10 ms。"
        "从主机换入一个专家：2.015 GB / (64 GB/s×80%) ≈ 39 ms，放不进 10 ms 这一步。"
    )

    # ----- slide 4 -----
    s4 = prs.slides.add_slide(blank)
    chrome(s4, 4, "30 毫秒、短输入，全局批 6928", "Decode    ·    TPOT ≤ 30 毫秒")
    banner(
        s4,
        0.92,
        "专家可以全部叫醒，读权重停在 13.4 毫秒。剩下的预算被机内 8 卡 allreduce 用掉，16.6 毫秒。专家域通信只有 0.9 毫秒。",
    )
    textbox(s4, 0.42, 1.46, 6.0, 0.26, "并行策略    ·    和 10 毫秒同一份切分", 13, True, ACCENT)
    textbox(s4, 7.15, 1.46, 5.7, 0.26, "短输入、全局批 6928，这一步怎么加", 13, True, ACCENT)
    par4 = [
        ["", "Attention", "MoE"],
        ["并行", "仍是 TP=8，DP=16\n6928 / 16 = 433 条 / 组", "仍是 EP=128\n每卡 19.3 个专家全醒"],
        ["通信", "allreduce 随批线性涨\n这一档是注意力侧最长", "combine 走 320 GB/s\n0.9 毫秒，短于读权重"],
        ["访存", "KV 在短输入上可忽略\n长输入改由容量卡住", "38.9 GB 路由 + 4.03 GB 共享\n/ 3.2 TB/s = 13.4 毫秒"],
        ["计算", "注意力计算 5.6 毫秒\n短于 allreduce", "每专家约 17 token\n按 32 行砖，效率 53%，11.9 毫秒"],
    ]
    add_table(
        s4,
        0.42,
        1.74,
        [1.15, 2.55, 2.70],
        par4,
        [0.30, 0.78, 0.82, 0.78, 0.78],
        ["head", "body", "key", "body", "body"],
    )
    step4 = [
        ["笔", "白话", "时间"],
        ["读权重", "命中率接近 1，19.3 个专家全读\n加共享专家，除以 3.2 TB/s", "13.4 毫秒"],
        ["机内通信", "每组 433 条 × 7168 × 2 字节\n61 层 × 2×7/8 / 40 GB/s", "16.6 毫秒"],
        ["专家计算", "6928 × 6 / 2469 ≈ 17 token / 专家\n对齐效率 53%", "11.9 毫秒"],
        ["专家通信", "最忙卡 combine / 320 GB/s", "0.9 毫秒"],
        ["注意力计算", "短输入的打分和投影", "5.6 毫秒"],
        ["相加", "注意力侧取 16.6，专家侧取 13.4", "30.0 毫秒"],
    ]
    add_table(
        s4,
        7.05,
        1.74,
        [1.25, 3.45, 1.15],
        step4,
        [0.30, 0.60, 0.60, 0.54, 0.46, 0.46, 0.50],
        ["head", "key", "key", "body", "body", "body", "sum"],
    )
    chips4 = [
        ("4K – 128K", "6928", "每卡 54 条\n一步顶满 30 毫秒"),
        ("1×10^6", "1156", "每条 4.94 GB\n这一步 18.2 毫秒"),
        ("2.08×10^6", "555", "每条 10.3 GB\n这一步 18.1 毫秒"),
        ("KV 余量", "5.72 TB", "96−38.9−4.03−2.42−6\n每卡 44.7 GB"),
    ]
    for i, (cap, num, sub) in enumerate(chips4):
        x = 0.42 + i * 3.20
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.40), Inches(3.05), Inches(1.62))
        solid(card, SOFT, LINE)
        textbox(s4, x + 0.16, 5.48, 2.75, 0.26, cap, 12, True, MUTED)
        textbox(s4, x + 0.16, 5.74, 2.75, 0.48, num, 26, True, TEAL)
        textbox(s4, x + 0.16, 6.26, 2.75, 0.64, sub, 13, False, INK)
    s4.notes_slide.notes_text_frame.text = (
        "30 毫秒用的切分与 10 毫秒相同：TP=8，DP=16，EP=128，128 卡一份副本。"
        "BS=6928 时 6×6928/2469 ≈ 16.8，命中率 1−exp(−16.8) 已经是 1，每卡 19.3 个专家全醒。"
        "读权重因此停在整表扫描：(38.9 GB + 4.03 GB) / 3.2 TB/s = 13.4 ms。"
        "机内 allreduce：ceil(6928/16)=433 条，payload=433×7168×2，环上每层 2×7/8，乘 61 层，除以 40 GB/s = 16.6 ms。"
        "两项相加 30.0 ms。专家 combine 只有 0.9 ms，所以 30 毫秒的通信限制在机内 8 卡，不在 128 卡专家域。"
        "专家计算 11.9 ms：每专家约 17 个 token，ceil(17/16)×16=32，效率 17/32=53%。"
        "显存：每卡 96 − 38.9 − 4.03 − 2.42 − 6 = 44.7 GB，128 卡共 5.72 TB。"
        "每条 KV = S×4939.5 B + 4.50 MB。1×10^6 时约 4.94 GB，能放 1156 条，这一步 18.2 ms，30 ms 还没用完。"
        "2.08×10^6 时约 10.3 GB，能放 555 条，这一步 18.1 ms。"
        "1×10^6 的 30 毫秒搜索结果被显存上限截断，批等于 1156，时间 18.2 ms。"
    )

    out = "/workspace/docs/deepseek-10t-slo.pptx"
    prs.save(out)
    print(out)
    print("S10", S)
    print("bs10", bs10)
    print("bs30", bs30)
    print("d10", {k: round(d10[k], 5) if isinstance(d10[k], float) else d10[k] for k in ("t", "tw", "tcm", "tar", "tep", "frac", "e_hot")})
    print("d30", {k: round(d30[k], 5) if isinstance(d30[k], float) else d30[k] for k in ("t", "tw", "tcm", "tar", "tep", "tca", "frac")})
    print("prefill", {k: round(p[k], 4) for k in ("tc", "idx", "hca", "moe", "proj", "core", "tw", "ship")})


if __name__ == "__main__":
    main()
