#!/usr/bin/env python3
"""Editable one-page loop of an agentic coding turn."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)
TEAL = (14, 122, 118)
GOLD = (122, 84, 28)
LINE = (214, 220, 226)
SOFT = (244, 247, 250)


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def write(shape, text, size, bold, color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def box(slide, x, y, w, h, fill, text="", size=14, bold=True, color=WHITE, line=None, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    if text:
        write(shape, text, size, bold, color, align)
    return shape


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def arrow(slide, x, y, w, h, fill, kind):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill)
    return shape


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Agentic Coding 一轮循环"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)

    textbox(slide, 0.42, 0.20, 12.4, 0.26, "Agentic Coding", 14, True, ACCENT)
    textbox(slide, 0.42, 0.46, 12.4, 0.44, "未完成就再转一圈", 28, True, NAVY)

    steps = (
        (TEAL, "1  组装上下文", "代码、对话、记忆"),
        (ACCENT, "2  推理", "决定下一步"),
        (GOLD, "3  工具调用", "读、改、跑"),
        (NAVY, "4  写入记忆", "结果写回"),
    )
    card_w, card_h = 2.55, 1.70
    gap = 0.58
    row_w = 4 * card_w + 3 * gap
    x0 = (13.333 - row_w) / 2
    y0 = 2.35

    for i, (fill, title, body) in enumerate(steps):
        x = x0 + i * (card_w + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y0), Inches(card_w), Inches(card_h)
        )
        solid(card, fill)
        write(card, f"{title}\n{body}", 18, True, WHITE)
        if i < 3:
            arrow(
                slide,
                x + card_w + 0.08,
                y0 + card_h / 2 - 0.14,
                gap - 0.16,
                0.28,
                fill,
                MSO_SHAPE.RIGHT_ARROW,
            )

    # Stop leaves from inference, above the row.
    infer_x = x0 + (card_w + gap)
    done_w, done_h = 2.10, 0.52
    done_x = infer_x + (card_w - done_w) / 2
    done_y = 1.55
    arrow(slide, infer_x + card_w / 2 - 0.14, done_y + done_h, 0.28, 0.24, GOLD, MSO_SHAPE.UP_ARROW)
    done = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(done_x), Inches(done_y), Inches(done_w), Inches(done_h)
    )
    solid(done, GOLD)
    write(done, "完成 · 输出", 16, True, WHITE)

    # Loop: memory down, left along the bottom, up into context.
    loop_y = 4.68
    stub_w, stub_h = 0.34, 0.46
    last_x = x0 + 3 * (card_w + gap)
    stub_h = (loop_y + 0.06) - (y0 + card_h)
    arrow(slide, last_x + card_w / 2 - stub_w / 2, y0 + card_h, stub_w, stub_h, NAVY, MSO_SHAPE.DOWN_ARROW)
    arrow(slide, x0 + card_w / 2 - stub_w / 2, y0 + card_h, stub_w, stub_h, TEAL, MSO_SHAPE.UP_ARROW)
    bar_x = x0 + card_w / 2 + 0.06
    bar_w = (last_x + card_w / 2) - bar_x - 0.06
    loop = slide.shapes.add_shape(
        MSO_SHAPE.LEFT_ARROW, Inches(bar_x), Inches(loop_y), Inches(bar_w), Inches(0.46)
    )
    solid(loop, TEAL)
    textbox(
        slide,
        bar_x + 0.7,
        loop_y,
        bar_w - 1.1,
        0.46,
        "下一轮",
        16,
        True,
        WHITE,
        PP_ALIGN.CENTER,
    )

    textbox(
        slide,
        0.42,
        5.50,
        12.4,
        0.55,
        "记忆写回后，下一轮从组装上下文再读。\n推理判定完成后直接输出。",
        16,
        False,
        MUTED,
        PP_ALIGN.LEFT,
    )

    slide.notes_slide.notes_text_frame.text = (
        "简化的 agentic coding 一轮。"
        "顺序是组装上下文、推理、工具调用、写入记忆。"
        "写入的记忆在下一轮组装上下文时被读回。"
        "推理判定任务完成后直接输出，不再进入工具调用。"
    )

    out = "/workspace/docs/agentic-coding-flow.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
