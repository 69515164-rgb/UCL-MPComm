#!/usr/bin/env python3
"""NVIDIA parameter-plane deck: architecture, standards posture, customer customization.

Slides are drawn with PIL (reliable CJK) and packed full-bleed into a 16:9 pptx.
Published figures are cropped from Khashab et al., arXiv:2605.21187 (20 May 2026).
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

W, H = 1920, 1080
OUT = Path("/workspace/docs/nvidia-param-plane")
FIGS = OUT / "src-figs"
FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

BG = (246, 244, 239)
INK = (27, 36, 48)
MUTED = (92, 102, 112)
TEAL = (15, 110, 107)
CORAL = (196, 73, 58)
NAVY = (28, 49, 68)
GOLD = (176, 128, 40)
CARD = (255, 255, 255)
LINE = (220, 214, 204)
SOFT = (237, 233, 224)
TEAL_BG = (232, 242, 241)
GOLD_BG = (252, 247, 234)
CORAL_BG = (250, 236, 232)
NAVY_BG = (236, 239, 243)
WHITE = (255, 255, 255)

N = 19


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def tw(d, t, f):
    return int(d.textlength(t, font=f))


def wrap(d, text, fnt, max_w):
    lines = []
    for para in str(text).split("\n"):
        if para == "":
            lines.append("")
            continue
        buf = ""
        for ch in para:
            if d.textlength(buf + ch, font=fnt) <= max_w:
                buf += ch
            else:
                if buf:
                    lines.append(buf)
                buf = ch
        if buf:
            lines.append(buf)
    return lines


def canvas():
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    for i in range(0, W, 48):
        d.line([(i, 0), (i, H)], fill=(236, 232, 224), width=1)
    for j in range(0, H, 48):
        d.line([(0, j), (W, j)], fill=(236, 232, 224), width=1)
    return img


def shadow(base, box, r=12):
    x0, y0, x1, y1 = box
    layer = Image.new("RGBA", base.size, (0, 0, 0, 0))
    ImageDraw.Draw(layer).rounded_rectangle(
        (x0 + 2, y0 + 3, x1 + 2, y1 + 5), radius=r, fill=(20, 24, 30, 18)
    )
    blur = layer.filter(ImageFilter.GaussianBlur(5))
    tmp = base.convert("RGBA")
    tmp.alpha_composite(blur)
    base.paste(tmp.convert("RGB"))


def card(img, box, fill=CARD, r=12, outline=LINE, w=1, drop=True):
    if drop:
        shadow(img, box, r)
    ImageDraw.Draw(img).rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=w)


def block(d, x, y, text, fnt, fill, max_w, lh=None, max_lines=30):
    lh = lh or (fnt.size + 7)
    lines = wrap(d, text, fnt, max_w)[:max_lines]
    for i, ln in enumerate(lines):
        d.text((x, y + i * lh), ln, font=fnt, fill=fill)
    return y + max(1, len(lines)) * lh


def pill(d, xy, text, bg, fg=WHITE, fnt=None, pad_x=8, pad_y=3):
    fnt = fnt or font(12)
    x, y = xy
    box = (x, y, x + tw(d, text, fnt) + pad_x * 2, y + fnt.size + pad_y * 2)
    d.rounded_rectangle(box, radius=9, fill=bg)
    d.text((x + pad_x, y + pad_y - 1), text, font=fnt, fill=fg)
    return box


def chrome(img, kicker, title, sub=None):
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 6), fill=TEAL)
    d.text((40, 16), kicker, font=font(13), fill=TEAL)
    d.text((40, 38), title, font=font(28), fill=INK)
    rule_y = 80
    if sub:
        d.text((40, 74), sub, font=font(14), fill=MUTED)
        rule_y = 100
    d.line([(40, rule_y), (1880, rule_y)], fill=LINE, width=1)
    return rule_y + 14


def footer(img, page, source):
    d = ImageDraw.Draw(img)
    d.line([(40, 1032), (1880, 1032)], fill=LINE, width=1)
    d.text((40, 1042), source, font=font(11), fill=MUTED)
    label = f"{page:02d}  /  {N:02d}"
    f = font(12)
    d.text((1880 - tw(d, label, f), 1040), label, font=f, fill=MUTED)


def judgement(img, box, text):
    card(img, box, fill=GOLD_BG, outline=(232, 214, 170), drop=False)
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = box
    d.rectangle((x0, y0, x0 + 6, y1), fill=GOLD)
    pill(d, (x0 + 16, y0 + 8), "判断", GOLD, fnt=font(11))
    block(d, x0 + 16, y0 + 32, text, font(13), INK, x1 - x0 - 32, lh=18, max_lines=3)


def trim_white(im, thr=248, pad=6):
    g = im.convert("L")
    bw = g.point(lambda p: 0 if p > thr else 255)
    bbox = bw.getbbox()
    if not bbox:
        return im
    x0, y0, x1, y1 = bbox
    x0 = max(0, x0 - pad)
    y0 = max(0, y0 - pad)
    x1 = min(im.width, x1 + pad)
    y1 = min(im.height, y1 + pad)
    return im.crop((x0, y0, x1, y1))


def contain(im, mw, mh):
    s = min(mw / im.width, mh / im.height)
    nw, nh = max(1, int(im.width * s)), max(1, int(im.height * s))
    return im.resize((nw, nh), Image.Resampling.LANCZOS)


def figure_card(img, box, path, caption, crop=None):
    card(img, box, fill=WHITE)
    im = Image.open(path).convert("RGB")
    if crop:
        im = im.crop(crop)
    im = trim_white(im)
    x0, y0, x1, y1 = box
    cap_h = 28
    fitted = contain(im, x1 - x0 - 16, y1 - y0 - cap_h - 14)
    px = x0 + 8 + (x1 - x0 - 16 - fitted.width) // 2
    py = y0 + 8 + (y1 - y0 - cap_h - 14 - fitted.height) // 2
    img.paste(fitted, (px, py))
    d = ImageDraw.Draw(img)
    d.text((x0 + 12, y1 - 24), caption, font=font(11), fill=MUTED)


def bullets(d, x, y, items, fnt, fill, max_w, lh, gap=8, mark=TEAL):
    for text in items:
        d.ellipse((x, y + 6, x + 7, y + 13), fill=mark)
        y = block(d, x + 16, y, text, fnt, fill, max_w - 16, lh=lh)
        y += gap
    return y


def simple_table(img, box, headers, rows, col_w, head_fill=NAVY, head_fg=WHITE,
                 body=13, rh=None, aligns=None):
    """rows: list of list of strings. Text wraps inside the cell."""
    x0, y0, x1, y1 = box
    n = len(rows)
    head_h = 32
    body_h = (y1 - y0 - head_h)
    rh = rh or body_h // max(1, n)
    d = ImageDraw.Draw(img)
    # header
    x = x0
    d.rectangle((x0, y0, x1, y0 + head_h), fill=head_fill)
    hf = font(12)
    for i, h in enumerate(headers):
        d.text((x + 8, y0 + 8), h, font=hf, fill=head_fg)
        x += col_w[i]
    bf = font(body)
    for r, row in enumerate(rows):
        yy = y0 + head_h + r * rh
        bg = WHITE if r % 2 == 0 else SOFT
        d.rectangle((x0, yy, x1, yy + rh), fill=bg)
        d.line([(x0, yy + rh), (x1, yy + rh)], fill=LINE, width=1)
        x = x0
        for i, cell in enumerate(row):
            block(d, x + 8, yy + 6, cell, bf, INK, col_w[i] - 16, lh=body + 4, max_lines=4)
            x += col_w[i]
    d.rectangle((x0, y0, x1, y0 + head_h + n * rh), outline=LINE, width=1)


def new():
    return canvas()


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_01():
    img = canvas()
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 780, H), fill=NAVY)
    d.rectangle((0, 0, W, 6), fill=TEAL)
    d.text((52, 56), "研究截止   2026-09-28", font=font(16), fill=(154, 196, 194))
    d.text((52, 150), "NVIDIA", font=font(22), fill=(120, 196, 190))
    d.text((52, 196), "参数面", font=font(68), fill=WHITE)
    d.text((52, 278), "完整解决方案", font=font(52), fill=WHITE)
    d.rectangle((52, 360, 180, 364), fill=TEAL)
    lines = [
        "01–09    参数面架构与公开实测",
        "10–13    UEC、SUE、ESUN、UALink 与标准策略",
        "14–19    OpenAI 等大客户定制与定制策略",
    ]
    y = 390
    for ln in lines:
        d.text((52, y), ln, font=font(18), fill=(214, 220, 226))
        y += 36
    d.text((52, 620), "参数面指承载集合通信的东西向 RDMA 织物。", font=font(16), fill=(186, 198, 208))
    d.text((52, 648), "Spectrum-X 以太与 Quantum InfiniBand 都在里面。", font=font(16), fill=(186, 198, 208))
    d.text((52, 676), "NVLink、BlueField 南北向、上下文存储不在里面。", font=font(16), fill=(186, 198, 208))
    d.text((52, 1000), "数字只来自论文、官方文件与已注明的媒体报道", font=font(14), fill=(140, 156, 168))

    items = [
        (TEAL, "参数面",
         "技术本体是多平面拓扑加上三条硬件控制环：交换机逐包自适应路由、网卡平面负载均衡、发送端拥塞控制。三者信号和时间尺度分开。以太是 2026 参考设计的默认项，InfiniBand 仍是旗舰训练的专家选项。"),
        (GOLD, "标准策略",
         "物理层深参与并担任领导职务；传输层只占座、不做符合性声明；NVLink 协议不捐赠，改用许可和投资换采用。可检验的说法是：买入的层做成标准，卖出的层发放许可。"),
        (CORAL, "定制策略",
         "固定核加可谈判的壳。硅片不按客户加功能；NVLink 域客户可以否决、不能另起；有合同杠杆时东西向织物写成必须是 InfiniBand 或 Spectrum-X。机柜、CPU、交换机整机可以让。"),
    ]
    y = 48
    for color, title, text in items:
        box = (820, y, 1872, y + 300)
        card(img, box)
        d = ImageDraw.Draw(img)
        d.rectangle((820, y, 828, y + 300), fill=color)
        d.text((852, y + 22), title, font=font(26), fill=INK)
        block(d, 852, y + 72, text, font(18), INK, 980, lh=28, max_lines=6)
        y += 324
    footer(img, 1, "本页三条是后文的目录，不是口号。每一条都在对应章节用公开材料展开。")
    return img


def slide_02():
    img = new()
    y0 = chrome(img, "02    定位", "参数面在五支柱里的位置",
                "五支柱是 NVIDIA 2026 年公开表述（Hot Chips，Shainer 向媒体解释 scale-in）。参数面只占其中一柱。")
    pillars = [
        ("Scale-Up", "机柜内", "NVLink 6", "72 GPU 一个域\n每 GPU 3.6 TB/s\n整域 260 TB/s", "不算参数面", SOFT, MUTED),
        ("Scale-Out", "参数面", "Spectrum-X 以太\nQuantum InfiniBand", "机柜之间的东西向\nRDMA 集合通信", "本报告的对象", TEAL_BG, TEAL),
        ("Scale-Across", "跨园区", "Spectrum-XGS", "数百米到数百公里\n距离自适应拥塞控制", "参数面的延伸", GOLD_BG, GOLD),
        ("Context Memory", "推理上下文", "BlueField-4 STX\nCMX", "共享 KV cache\n以太挂载的闪存层", "不算参数面", SOFT, MUTED),
        ("Scale-In", "南北向", "BlueField-4 + DOCA", "安全、存储、控制面\n信任域 ASTRA", "故意不在参数面", CORAL_BG, CORAL),
    ]
    gap = 14
    cw = (1840 - gap * 4) // 5
    x = 40
    for name, role, tech, desc, tag, bg, color in pillars:
        box = (x, y0, x + cw, y0 + 520)
        card(img, box, fill=bg, outline=color if color == TEAL else LINE, w=2 if color == TEAL else 1)
        d = ImageDraw.Draw(img)
        d.rectangle((x, y0, x + cw, y0 + 8), fill=color)
        d.text((x + 16, y0 + 24), name, font=font(18), fill=color)
        d.text((x + 16, y0 + 54), role, font=font(15), fill=MUTED)
        block(d, x + 16, y0 + 100, tech, font(16), INK, cw - 32, lh=22)
        block(d, x + 16, y0 + 220, desc, font(15), INK, cw - 32, lh=22)
        pill(d, (x + 16, y0 + 460), tag, color, fnt=font(13))
        x += cw + gap
    # bandwidth generations
    card(img, (40, y0 + 540, 1880, 1010))
    d = ImageDraw.Draw(img)
    d.text((60, y0 + 556), "参数面端点带宽的三代", font=font(16), fill=INK)
    gens = [
        ("Hopper", "400 GbE / GPU", "xAI Colossus：BlueField-3\n论文测试床：ConnectX-7"),
        ("Blackwell", "800 Gb/s / GPU", "ConnectX-8，2×400G\n论文四平面测试床用这一代"),
        ("Vera Rubin", "1.6 Tb/s / GPU", "ConnectX-9\n端口拆分口径官网不一致，总数一致"),
    ]
    x = 60
    for i, (g, bw, note) in enumerate(gens):
        d.text((x, y0 + 600), g, font=font(14), fill=MUTED)
        d.text((x, y0 + 626), bw, font=font(26), fill=TEAL if i == 2 else INK)
        block(d, x, y0 + 670, note, font(14), MUTED, 520, lh=20)
        if i < 2:
            d.text((x + 500, y0 + 630), "→", font=font(28), fill=LINE)
        x += 600
    footer(img, 2, "端点带宽：NVIDIA Spectrum-X 产品页 FAQ；Vera Rubin 技术博客。五支柱：SDxCentral / Converge Digest 对 Shainer 的转述，2026。")
    return img


def slide_03():
    img = new()
    y0 = chrome(img, "03    分工", "参数面有两条：InfiniBand 与 Spectrum-X 以太",
                "同一套 SuperNIC 家族可以接两条织物。差别在交换机上的网内计算、运维栈，以及 NVIDIA 自己把谁写成默认。")
    cols = [
        (NAVY, "Quantum InfiniBand", [
            "Quantum-X800：144×800 Gb/s，SHARP v4，自适应路由，基于遥测的拥塞控制，UFM 管理口。",
            "官方定位：万亿参数训练。Huang 2025-08 财报电话：对超算和头部模型厂商，Quantum InfiniBand 是明确选择。",
            "公开客户落点：Azure NDv6 GB300 集群内用 Quantum-X800；CoreWeave GB300 用 Quantum-X800 + ConnectX-8，每 GPU 800 Gb/s。",
            "上一代 Quantum-2 仍在 2024 年的 Meta 对照簇和 OCI Supercluster 里。",
        ]),
        (TEAL, "Spectrum-X 以太", [
            "Spectrum 交换机 + ConnectX / BlueField SuperNIC，端到端配对。支持 SONiC。",
            "论文架构：多平面、逐包自适应路由、网卡平面负载均衡、无损加发送端拥塞控制。",
            "2026-03 Vera Rubin DSX 参考设计点名的织物是 Spectrum-X Ethernet，不是 InfiniBand。",
            "Huang 原话：Spectrum Ethernet is not off the shelf。Shainer 的配套说明：抖动问题不能只在交换机或只在网卡上解决，两边要一起做。",
        ]),
    ]
    cw = 890
    for i, (color, title, items) in enumerate(cols):
        x = 40 + i * (cw + 20)
        box = (x, y0, x + cw, y0 + 560)
        card(img, box)
        d = ImageDraw.Draw(img)
        d.rectangle((x, y0, x + cw, y0 + 8), fill=color)
        d.text((x + 20, y0 + 22), title, font=font(22), fill=color)
        bullets(d, x + 20, y0 + 70, items, font(15), INK, cw - 48, lh=22, gap=14, mark=color)
    judgement(img, (40, y0 + 580, 1880, y0 + 680),
              "以太已经是 NVIDIA 自己 2026 旗舰参考设计的默认参数面；InfiniBand 留在要 SHARP、要 UFM、要旗舰训练尾延迟的场景。两条都卖，默认项换了。")
    footer(img, 3, "Quantum-X800 产品页；Huang Q2 FY2026 财报电话；DSX 新闻稿 2026-03-16；Spectrum-X 由交换机与 SuperNIC 配对，见 Shainer theCUBE 与 Huang 本人。")
    return img


def slide_04():
    img = new()
    y0 = chrome(img, "04    物料", "三代参数面物料：只写两边口径一致的数",
                "ConnectX-9 的端口拆分，官网 FAQ 写 4×200G，技术博客表格写 2×800G，算术对不上。片子采用两边一致的每 GPU 1.6 Tb/s。")
    headers = ["", "Hopper 一代", "Blackwell", "Vera Rubin"]
    rows = [
        ["以太交换", "Spectrum-4\n51.2 Tb/s\nSN5600 64×800GbE", "SN5000 系列\n同一代 Spectrum-4", "Spectrum-6\n102.4 Tb/s / 芯片\nSN6600 128×800G\nSN6800 512×800G，5U"],
        ["端点", "论文测试床 CX-7\nxAI 部署 BF-3\n400 GbE / GPU", "ConnectX-8\n800 Gb/s / GPU\n2×400G", "ConnectX-9\n1.6 Tb/s / GPU\n每计算托盘四块板"],
        ["InfiniBand", "Quantum-2\nMeta 2024 对照簇\nOCI Supercluster", "Quantum-X800\n144×800 Gb/s\nSHARP v4", "官方仍把 Quantum-X800\n与 Spectrum-X 并列\n为 Rubin 的 scale-out"],
        ["网内计算", "SHARP 在 IB 交换机\n以太侧公开特性清单无对等项", "SHARP v4 在\nQuantum-X800\nNVLink Switch 另有", "NVLink 6 交换托盘\n14.4 TFLOPS FP8\n厂商口径，在 scale-up"],
        ["论文测试床", "1024 GPU，128 节点\n1 平面，三层胖树，CX-7", "144 GPU 单平面 CX-7\n1152 GPU，4 平面，CX-8\n两层 rail-optimized", "架构章写到 Blackwell Ultra\nVR200 只作为代际点到"],
    ]
    simple_table(img, (40, y0, 1880, y0 + 780), headers, rows,
                 [200, 520, 520, 600], body=14, rh=148)
    footer(img, 4, "论文 Table 2；Spectrum-X FAQ；Vera Rubin 技术博客；Quantum-X800 产品页；xAI 新闻稿 2024-10-28。SN6800 的 409.6 Tb/s 是 4×102.4 的推算，片子不写。")
    return img


def slide_05():
    img = new()
    y0 = chrome(img, "05    拓扑    公开图 Figure 2", "用平面并行代替再加一层脊",
                "Spectrum-X 论文第 3 节。讨论范围停在 Blackwell Ultra 的多平面，南北向、scale-across、NVLink 作者标明超出本文。")
    card(img, (40, y0, 860, 900))
    d = ImageDraw.Draw(img)
    items = [
        "平面在交换机侧断开。每张网卡经无源光 shuffle box 接入全部平面，网卡到网卡仍然全可达。",
        "论文给出的规模：两层多平面到 12.8 万端点，三层到 1600 万。依据是把 800 Gb/s 拆成 4×200 Gb/s，每平面一张两层胖树。",
        "对应用透明。无论几个平面，网卡只暴露一个 RoCE 设备。集合通信库和传输层不用改。",
        "乱序由网卡直接写入目的内存，完成队列在网卡里重排，应用看不到重排。",
        "非满配时用并行链路填满端口，换弹性和端口利用率。论文的测试床也是这么缩的：100 台 10% 占用的脊，收成 10 台满配脊加 10 条并行链路。",
        "调度不要求作业按机柜对齐集合通信模式。作者称为 job-allocation agnostic，并用隔离实验支撑。",
    ]
    bullets(d, 60, y0 + 20, items, font(15), INK, 760, lh=22, gap=12)
    figure_card(img, (880, y0, 1880, 900), FIGS / "fig02.png",
                "Figure 2  四平面两层胖树，rail-optimized。Khashab et al., arXiv:2605.21187")
    footer(img, 5, "图与规模数字均来自论文原文。12.8 万与 1600 万是拓扑容量上界，不是已交付集群规模。")
    return img


def slide_06():
    img = new()
    y0 = chrome(img, "06    控制环    公开图 Figure 3 / 4", "三条硬件环：信号、目标、时间尺度都分开",
                "作者的原话：把三条环关进各自的带宽时延积之内，软件控制路径做不到。硬件加速是结构要求，不是性能优化。")
    loops = [
        ("自适应路由  AR", "交换机，fabric 口", "无状态，逐包", "数百纳秒",
         "量化的 Join-Shortest-Queue。每个包对 ECMP 组里全部出口按队列深度打分，队列以亚微秒采样，转发到最空的口之一。"),
        ("平面负载均衡  PLB", "网卡，平面口", "每平面每目的地", "逐包，数个 RTT",
         "每个目的地四份拥塞控制上下文。先按速率额度掩掉限速或故障平面，再在剩余平面里选本地出口队列最浅的。"),
        ("拥塞控制  CC", "发送端", "每目的地，有状态", "RTT",
         "RTT 探针，加上交换机显式信号。ECN 只在负载均衡容量用尽时才打标记。拥塞控制只对这些信号反应，不追瞬时微突发。"),
    ]
    cw = 600
    for i, (name, where, state, scale, desc) in enumerate(loops):
        x = 40 + i * (cw + 16)
        box = (x, y0, x + cw, y0 + 250)
        card(img, box)
        d = ImageDraw.Draw(img)
        d.rectangle((x, y0, x + 6, y0 + 250), fill=TEAL)
        d.text((x + 18, y0 + 14), name, font=font(16), fill=TEAL)
        d.text((x + 18, y0 + 42), f"{where}    {state}", font=font(13), fill=MUTED)
        pill(d, (x + 18, y0 + 66), scale, GOLD, fnt=font(12))
        block(d, x + 18, y0 + 100, desc, font(13), INK, cw - 36, lh=18, max_lines=6)
    # lower: explanation + figure
    card(img, (40, y0 + 266, 980, 1008))
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 280), "为什么必须拆开", font=font(16), fill=INK)
    bullets(d, 58, y0 + 312, [
        "fabric 口上的微突发会抬高 RTT。拥塞控制若把这读成端点 incast，会在所有平面上一起限速，或让平面均衡误迁健康平面的流量。",
        "平面内不对称通常很小：一条叶上联在几百条里断掉，其余上联负载增加不到 1%。平面之间不对称很大：四平面断一平面，剩下的会被不成比例地打满。",
        "所以平面内无状态、平面间有状态。一个全局控制器跟不上平面之间的不对称。",
        "逐包喷洒使乱序和丢包分不清，因此链路层无损。NVIDIA 称两年生产、最大规模集群里没有观察到 PFC 风暴。",
        "生产遥测显示超过 97% 的流量容忍乱序。要保序的控制流量不进硬件加速，走普通以太 / TCP / BGP。",
    ], font(13), INK, 900, lh=18, gap=6)
    figure_card(img, (996, y0 + 266, 1880, 1008), FIGS / "fig04.png",
                "Figure 4  网卡逐包选平面：掩掉限速面和故障面，再选最浅队列")
    footer(img, 6, "论文 §3.2、§4.1–§4.3。ECN 标记点见 Figure 3：共享缓存上限之下先由 AR 吸收，标记发生在负载均衡容量用尽之后。")
    return img


def slide_07():
    img = new()
    y0 = chrome(img, "07    故障    公开图 Figure 5 / 12", "瞬时故障走硬件，永久不对称走控制面权重",
                "本地链路，自适应路由在约 100 纳秒内排除。远端永久故障的权重由 BGP 控制面计算，算完之后硬件仍按包平衡。")
    # big numbers
    stats = [
        ("< 3 ms", "四平面断一平面\n硬件 PLB 回到 75% 线速", TEAL),
        ("1.08 s", "NCCL 之上的软件负载均衡\n大约慢 400 倍", CORAL),
        ("−11%", "永久少 10% 链路\n带宽从 378 降到 335 Gb/s", NAVY),
        ("+1 µs", "同一实验的 p99 时延\n14.97 → 15.96 µs", GOLD),
    ]
    cw = 450
    for i, (n, t, c) in enumerate(stats):
        x = 40 + i * (cw + 14)
        card(img, (x, y0, x + cw, y0 + 130))
        d = ImageDraw.Draw(img)
        d.text((x + 16, y0 + 12), n, font=font(28), fill=c)
        block(d, x + 16, y0 + 52, t, font(14), INK, cw - 32, lh=20)
    figure_card(img, (40, y0 + 146, 940, y0 + 620), FIGS / "fig05.png",
                "Figure 5  加权自适应路由：远端 D 容量下降，去 D 的权重转向 S1")
    figure_card(img, (956, y0 + 146, 1880, y0 + 400), FIGS / "fig12d.png",
                "Figure 12  硬件 PLB 约 2.7 ms 回到三条平面；软件负载均衡约 1.08 s")
    card(img, (956, y0 + 412, 1880, y0 + 620))
    d = ImageDraw.Draw(img)
    bullets(d, 972, y0 + 424, [
        "Nemotron 3 Ultra，64 节点：主机链路闪断在一个迭代内退到三条平面，恢复后回到基线步时 2.95 秒。叶到脊闪断几乎不动步时。",
        "单平面时主机链路一断，RDMA 连接直接崩溃。",
        "NSX 仿真 25.6 万 GPU：收敛在 10 ms 内，P99 完成时间约 +20%；100 ms 约 +53%，300 ms 约 +260%。闪断率是作者称为偏保守的最坏估计。",
    ], font(13), INK, 880, lh=18, gap=6)
    judgement(img, (40, y0 + 636, 1880, 1010),
              "参数面的弹性指标不是丢包率，是收敛时间。10 毫秒以内，损失约等于少掉的那条平面；再慢，集合通信的长尾会把整步拖垮。这是后文看 UEC 和客户自研喷洒时的同一把尺子。")
    footer(img, 7, "论文 §4.4、§6.4–§6.6，Figure 5、12、13、14。75% 与 3 毫秒来自 Blackwell_Ultra_MP 上的硬件 PLB 对软件 LB。")
    return img


def slide_08():
    img = new()
    y0 = chrome(img, "08    实测    公开图 Figure 8 / 13", "论文里的四个结果，对照物是 DCQCN + ECMP",
                "对照不是 UEC，也不是云厂商自研喷洒。作者写明：ETH 基线是传统 RoCEv2，DCQCN 加普通 ECMP。")
    kpis = [
        ("98%", "p01 带宽达到线速", "64 节点，最差放置，流量全部过脊。p01 = 377 Gb/s。"),
        ("8–9 µs", "75% 负载下的 p99", "以太中位约 13 µs，散到 22 µs。"),
        ("668 ms", "有背景流量的步时", "DeepSeek-V3 代理，16 节点。单独跑 667 ms。以太从 735 ms 涨到 1.18 s。"),
        ("1.13×", "安静 All2All", "49.3 / 49.5 GB/s，约 99.5% 硬件能力。以太顶在 43 GB/s。"),
    ]
    cw = 450
    for i, (n, t, s) in enumerate(kpis):
        x = 40 + i * (cw + 14)
        card(img, (x, y0, x + cw, y0 + 150))
        d = ImageDraw.Draw(img)
        d.text((x + 14, y0 + 10), n, font=font(26), fill=TEAL)
        d.text((x + 14, y0 + 46), t, font=font(15), fill=INK)
        block(d, x + 14, y0 + 72, s, font(13), MUTED, cw - 28, lh=18, max_lines=3)
    figure_card(img, (40, y0 + 166, 960, y0 + 560), FIGS / "fig08.png",
                "Figure 8  高负载：左为带宽分布，右为 300 Gb/s 下的 p99 时延")
    figure_card(img, (976, y0 + 166, 1880, y0 + 560), FIGS / "fig13.png",
                "Figure 13  Nemotron 3 Ultra。灰为主机链路闪断，绿为叶—脊闪断")
    card(img, (40, y0 + 576, 1880, 1010))
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 590), "隔离、运维，以及这些数不能外推的地方", font=font(16), fill=INK)
    bullets(d, 58, y0 + 624, [
        "两份 All2All 并行，16 节点受害、48 节点背景：以太受害带宽掉到 10.9 GB/s 以下，大约掉八成。Spectrum-X 近乎不掉。",
        "不对称实验：两台叶的上联从 16×200G 削到 4×200G。全局一份拥塞控制上下文时，一对多带宽从 94.5 掉到 47.3 GB/s。按平面分开的上下文维持 93.7。4 MB 以下报文还没攒够每平面信号，归一化性能可低到 0.75–0.85。",
        "运维：自适应路由把流量打成对称。叶上联、rail、平面任何一组偏离对称，就是故障或配置错误。高频遥测采样在 100 微秒到 10 毫秒，用来抓住集体通信内部的双峰带宽。",
        "部署陈述（论文，未经第三方审计）：数十个客户集群，合计超过 100 万 GPU，生产超过两年。",
    ], font(14), INK, 1780, lh=20, gap=4)
    footer(img, 8, "论文 §5、§6，Table 1，Figure 8、9、10、13、15。DeepSeek 步时是 16 节点 NVL8 代理，不是全量预训练。")
    return img


def slide_09():
    img = new()
    y0 = chrome(img, "09    边界", "和参数面一起卖、但不在同一层的三件事",
                "SHARP、跨园区、共封装光，经常被写进同一份 Spectrum 材料。它们的位置不同。")
    cols = [
        (NAVY, "SHARP  网内归约", "在 scale-up 和 InfiniBand 上",
         [
             "Quantum-X800 官方特性：硬件网内计算 SHARP v4。",
             "NVLink 6 交换托盘：14.4 TFLOPS FP8。厂商称 AllReduce 通信量最多减少 50%，张量并行最多加快 20%。",
             "Spectrum-X 以太平台的公开特性清单是自适应路由、可编程拥塞控制、多平面、高频遥测、scale-across、共封装光。没有交换机侧 SHARP。",
             "云伙伴规范 NET-2 的写法是 NCCL collectives（SHARP where supported）。支持与否按织物区分。",
         ]),
        (GOLD, "Spectrum-XGS  跨园区", "参数面向外的一跳",
         [
             "2025-08 Hot Chips 发布。距离从 500 米以上到数百公里。",
             "距离自适应拥塞控制、精度时延管理、感知距离的自适应路由。",
             "厂商口径：10 公里、大消息，NCCL AllReduce 带宽相对现成以太 1.9 倍。NVIDIA 内部实验，只覆盖大消息。",
             "具名首个采用者是 CoreWeave。微软 Fairwater 的跨站公开名称是 MRC，不是 XGS。",
         ]),
        (TEAL, "共封装光", "同一台以太交换机的光形态",
         [
             "Vera Rubin 平台的 Spectrum-X Ethernet Photonics，基于 Spectrum-6。",
             "厂商口径，相对可插拔：网络能效 5 倍、平均故障间隔 10 倍、正常运行时间 5 倍。",
             "GTC 2026 称已投产，与台积电合作。独立功耗拆解未见公开第三方测量。",
             "它优化的是功耗和故障半径，不改变三条控制环。",
         ]),
    ]
    cw = 600
    for i, (color, title, role, items) in enumerate(cols):
        x = 40 + i * (cw + 16)
        card(img, (x, y0, x + cw, 900))
        d = ImageDraw.Draw(img)
        d.rectangle((x, y0, x + cw, y0 + 8), fill=color)
        d.text((x + 16, y0 + 20), title, font=font(18), fill=color)
        d.text((x + 16, y0 + 50), role, font=font(13), fill=MUTED)
        bullets(d, x + 16, y0 + 84, items, font(14), INK, cw - 40, lh=20, gap=10, mark=color)
    judgement(img, (40, 916, 1880, 1010),
              "参数面的以太性能来自交换机和 SuperNIC 的配对控制环。SHARP 是另一条织物和 scale-up 的加速器。把三者写成一个协议栈，会对不上 NVIDIA 自己的产品页。")
    footer(img, 9, "SHARP：Quantum-X800 产品页与 Vera Rubin 技术博客。XGS：2025-08-22 新闻稿。CPO：Spectrum-X 产品页与 GTC 2026 报道。NET-2：docs.nvidia.com DGX Cloud 伙伴规范。")
    return img


def slide_10():
    img = new()
    y0 = chrome(img, "10    标准    版图", "标准组织里 NVIDIA 的位置，按层而不是按名单",
                "会员身份截至 2026-09-28。贡献深度凡是组织不公布名册的，都标成查无，而不是写成零贡献。")
    headers = ["层", "组织", "NVIDIA 的公开位置", "和参数面的关系"]
    rows = [
        ["物理层", "IEEE 802.3dj\nOIF CEI", "深。802.3dj 2026-03 出席至少 6 名工程师并署名评论决议。Karl Bois 任 OIF 技术委员会副主席（2026）。", "买入的层。200G/lane 和 1.6T 光电气，谁都要用。"],
        ["机柜机械", "OCP", "2024-10-15 贡献 GB200 NVL72 机柜、托盘、线缆仓体积。Spectrum-X 支持 SAI 与 SONiC。", "标准化机柜扩大自己的供应链。协议文本不在捐赠里。"],
        ["Scale-out 传输", "UEC", "2024-06 向媒体确认会员，2024-09 公开列入一般会员。无指导委员会席位。查无工作组主席或编辑。", "威胁的是以太 AI 网这个市场。占座，不做符合性声明。"],
        ["RoCE 血统", "IBTA", "Shainer：RoCE 在 IBTA 标准化，NVIDIA 在其中继续增强规范。", "自己参与制定的传输。UEC 的 UET 绕开这条血统。"],
        ["Scale-up 以太", "OCP ESUN\nSUE-T", "ESUN 12 家创始之一（2025-10-13），含 OpenAI。Broadcom 把 NVIDIA 列入 SUE-T 支持者。", "以太网如果做成 scale-up，NVIDIA 的交换机仍可能卖进去。"],
        ["Scale-up 协议", "UALink", "不在任何公开名册。2024-05 发起时未被邀请，TechCrunch 称 NVIDIA 拒绝对此置评。", "对手协议，NVIDIA 的交换机卖不进去。两年半未加入。"],
        ["中国 scale-up", "ODCC\nGSE / ETH-X / ALS", "查无 NVIDIA。阿里、百度、字节、华为、腾讯在 2023-11 加入的是 UEC。", "国内客户在 UEC 里，不在 NVIDIA 缺席的国内规范里等它。"],
    ]
    simple_table(img, (40, y0, 1880, y0 + 860), headers, rows,
                 [180, 220, 860, 580], body=13, rh=114)
    footer(img, 10, "UEC 会员：The Next Platform 2024-06-26；UEC 2024-09 会员帖。802.3dj 会议纪要。OIF 2026-01-14 官员公告。ESUN：OCP 与 Meta 工程博客。")
    return img


def slide_11():
    img = new()
    y0 = chrome(img, "11    UEC", "UEC 在标准化 NVIDIA 已经用专有配对实现的那些功能",
                "1.0 于 2025-06-11 发布。1.1 截至 2026-09-28 未发布。联盟 2025-11 给 IEEE 的联络函写过 CSIG 计划在 2026 Q1 随 1.1 公布。")
    # left timeline
    card(img, (40, y0, 620, y0 + 430))
    d = ImageDraw.Draw(img)
    d.text((56, y0 + 12), "NVIDIA 在 UEC 的足迹", font=font(16), fill=INK)
    bullets(d, 56, y0 + 44, [
        "2023-07-19 成立。九家创始加 Oracle 进指导委员会。NVIDIA 不在。",
        "2024-06-26 对 The Next Platform：会员身份是为了支持对客户有用的规范。未来也许会在 Spectrum-X 之外提供一个 UEC 版本的以太。",
        "2024-09 列入一般会员。指导委员会名册持续没有 NVIDIA。",
        "ConnectX-8 手册的符合性清单是 IEEE 802.3 与 IBTA 1.7，没有 UEC。",
        "2026-03 GTC 的 ConnectX-9 与 Spectrum-6 发布稿没有 UEC。",
        "OFC 2026 首次公开的 LLR/CBFC 800GE 互通是 Keysight 加 Broadcom。NVIDIA 不在场。",
        "对照：Broadcom Thor Ultra 宣称 fully feature compliant with UEC。",
    ], font(13), INK, 540, lh=18, gap=4)
    # mapping
    headers = ["功能", "UEC", "Spectrum-X 对应物"]
    rows = [
        ["多路径", "Packet spray，端与交换机协同", "交换机自适应路由，SuperNIC 重排"],
        ["拥塞控制", "NSCC / RCCC；1.1 计划 PCM，算法可跨网卡", "SuperNIC 可编程拥塞控制"],
        ["遥测", "CSIG，4 或 8 字节，沿途比较替换", "端到端高频遥测，线格式不公开"],
        ["链路可靠", "LLR + CBFC，LLDP 协商", "单链路故障限制在该链路；无公开互操作规范"],
        ["网内集合", "INC，可选实现", "SHARP，在 NVLink 与 IB，不在以太交换机"],
        ["传输", "新传输 UET，无握手的短连接", "RoCEv2 加 NVIDIA 扩展"],
        ["Scale-up", "1.1 计划：UFH、RODL、单向 <1 µs", "NVLink。这是对护城河的规格"],
    ]
    simple_table(img, (636, y0, 1880, y0 + 430), headers, rows,
                 [140, 520, 584], body=12, rh=54)
    judgement(img, (40, y0 + 446, 1880, y0 + 560),
              "映射是分析，不是 NVIDIA 的对照表。功能重叠成立；实现位置相反：UEC 把重排和拥塞算法放进标准传输，NVIDIA 放进交换机与 SuperNIC 的配对。PCM 的公开卖点就是同一算法能跑在任何支持它的网卡上。")
    card(img, (40, y0 + 576, 1880, 1010), fill=SOFT, drop=False)
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 590), "1.0 已经写进规范的，和 1.1 还停在计划里的", font=font(15), fill=INK)
    left = "1.0（已发布）：UET 四子层 SES / PDS / CMS / TSS；逐包多路径；无握手短连接；选择性重传；可选的包修剪、LLR、CBFC；发送端 NSCC 与接收端 RCCC；可选网内集合；AI Base / AI Full / HPC 三档。物理层和 IP 层故意不动。"
    right = "1.1（未发布，ITU-T 2026-07-11 的计划页）：scale-up、统一转发头、RODL、单向时延从 <10 µs 收到 <1 µs、Load/Store/Atomic、AI Local 档、CSIG、可编程拥塞管理 PCM。1.0.3（2026-07-16）只补了 200 Gb/s 每车道，没有把 1.1 带出来。"
    block(d, 58, y0 + 622, left, font(14), INK, 860, lh=20)
    block(d, 960, y0 + 622, right, font(14), INK, 880, lh=20)
    footer(img, 11, "UEC 规范史页面；1.0.3 发布说明；Paul Congdon ITU-T 2026-07-11；ConnectX-8 用户手册；Broadcom Thor Ultra 新闻稿 2025-10-14。")
    return img


def slide_12():
    img = new()
    y0 = chrome(img, "12    Scale-up 标准", "四条 scale-up 路线，NVIDIA 的席位各不相同",
                "ESUN 管织物需求，SUE-T 管端点传输，UALink 是另一套协议，NVLink Fusion 是许可。四者不是同一个标准的四个名字。")
    cols = [
        ("SUE / SUE-Lite", CORAL, [
            "作者是 Broadcom。2025-04 贡献给 OCP，修订史上唯一作者。",
            "SUE 含端到端可靠传输和固定窗口拥塞控制。",
            "SUE-Lite，2025-07：砍掉端到端可靠、拥塞控制和分区，只留逐跳 LLR 与 CBFC。规范写明 IP 面积最多减少 50%。",
            "机柜内时延短，逐跳重传就够。这是用面积和功耗打 NVLink 的算法。",
        ]),
        ("ESUN 1.0", TEAL, [
            "2025-10-13，OCP。12 家创始含 NVIDIA、OpenAI、Meta、微软、Oracle。",
            "2026-02-12 指导委员会批准。参与公司超过 175 家。",
            "文本是网络运营商需求基线，不是线协议。大部分需求引用 UEC 1.0，压缩头用自写的 ESUN Header。",
            "NVIDIA 在创始名单里。参与深度超出名单本身，公开材料不够。",
        ]),
        ("UALink", NAVY, [
            "2024-05-30 发起。NVIDIA 不在，也未被邀请。",
            "1.0，2025-04-08：200G 每车道，每 pod 1024 加速器。",
            "2.0，2026-04-07：把网内计算写进 Common 2.0。联盟自己的新闻标题写着 2.0 规范早于 1.0 硅片出货。",
            "3.0 指向 2027。首批硅片的公开日期仍是分析师估计。",
        ]),
        ("NVLink Fusion", GOLD, [
            "2025-05-18 Computex。许可 NVLink-C2C、融合芯粒、交换机芯片和 MGX 机柜。",
            "协议文本没有捐给任何标准组织。",
            "投资换采用：Intel 50 亿美元（2025-12 交割），Marvell 20 亿美元（2026-03），MediaTek 35 亿美元可转债（2026-08）。",
            "已公开的接入方含 AWS Trainium4、Fujitsu、Qualcomm、SiFive。",
            "每套部署必须含一件 NVIDIA 产品：只有媒体报道，没有 NVIDIA 文件。不写成事实。",
        ]),
    ]
    cw = 450
    for i, (title, color, items) in enumerate(cols):
        x = 40 + i * (cw + 14)
        card(img, (x, y0, x + cw, 900))
        d = ImageDraw.Draw(img)
        d.rectangle((x, y0, x + cw, y0 + 8), fill=color)
        block(d, x + 14, y0 + 18, title, font(16), color, cw - 28, lh=20)
        bullets(d, x + 14, y0 + 70, items, font(13), INK, cw - 32, lh=18, gap=8, mark=color)
    judgement(img, (40, 916, 1880, 1010),
              "NVIDIA 加入 ESUN，是因为那是以太网，交换机还卖得出去。NVIDIA 不加入 UALink，是因为那是一套替代协议。NVLink 本身保持许可，不进入上述任何一个文本。")
    footer(img, 12, "SUE 规范修订史；Broadcom 2025-07-15 Tomahawk Ultra 新闻稿；OCP ESUN 博客；UALink 1.0 / 2.0 新闻稿；NVLink Fusion 新闻稿 2025-05-18 及后续投资公告。")
    return img


def slide_13():
    img = new()
    y0 = chrome(img, "13    标准策略", "买入的层做成标准，卖出的层发放许可",
                "这是从第 10–12 页归纳的分析判断，不是 NVIDIA 的自我描述。NVIDIA 的自我描述写在下面的原话里。")
    headers = ["层", "威胁的是什么", "公开行为"]
    rows = [
        ["物理层、光", "都不威胁。速率上去，买卖双方都要", "点名工程师、评论决议、OIF 技术委员会副主席"],
        ["机柜机械", "不威胁。标准件扩大可买到的供应链", "捐赠 NVL72 机械件；协议不捐"],
        ["Scale-out 传输", "威胁市场。以太 AI 网 NVIDIA 是强供应商，不是唯一", "一般会员，无指导席，无符合性声明，缺席首次互通"],
        ["Scale-up 以太织物", "威胁护城河，但插座仍可能插 NVIDIA 交换机", "ESUN 创始成员，SUE-T 支持者名单"],
        ["Scale-up 对方协议", "威胁护城河，且卖不进自己的交换机", "UALink 名册上缺席"],
        ["NVLink 本身", "护城河", "不捐赠。选择性许可，并用投资换三家关键设计公司的采用"],
    ]
    simple_table(img, (40, y0, 1880, y0 + 430), headers, rows,
                 [280, 760, 800], body=14, rh=64)
    # quotes
    quotes = [
        ("2024-06-26  对 The Next Platform",
         "我们加入 UEC，是因为策略是支持对客户有用的网络规范。未来也许会在 Spectrum-X 之外，再提供一个 UEC 版本的以太。"),
        ("Huang  Q2 FY2026",
         "Spectrum Ethernet is not off the shelf。它有一整组为低时延、低抖动和拥塞控制新做的技术。"),
        ("Shainer  theCUBE  2026-07-16",
         "我们在 UEC，在 ESUN，在许多联盟里，并且实际有贡献。同时我们必须非常快，因为每年都有新的一代。"),
    ]
    cw = 600
    yy = y0 + 446
    for i, (who, text) in enumerate(quotes):
        x = 40 + i * (cw + 16)
        card(img, (x, yy, x + cw, yy + 168))
        d = ImageDraw.Draw(img)
        d.text((x + 14, yy + 10), who, font=font(12), fill=TEAL)
        block(d, x + 14, yy + 34, text, font(13), INK, cw - 28, lh=18, max_lines=6)
    judgement(img, (40, yy + 184, 1880, 1010),
              "原话把开放落在接口能连通。性能层仍要求交换机和 SuperNIC 配对，Huang 自己说不是现成以太。年度一代被 Shainer 用来解释为什么不能等共识流程。弱项：不能说 NVIDIA 不做标准。它做物理层标准，避开传输层标准。")
    footer(img, 13, "原话出处见引语抬头。每吉瓦收入机会 180→250→400 亿美元（Q2 FY2027 电话，含 CPU、GPU、NVLink、IB 或以太、Groq），是这条策略的收入背景，不是某一层的报价。")
    return img


def slide_14():
    img = new()
    y0 = chrome(img, "14    定制边界", "固定的核，可谈判的壳",
                "边界写在云伙伴规范的 MUST 里，也写在客户实际买到的东西里。两边对照，避免把新闻稿的标准化写成架构让步。")
    card(img, (40, y0, 930, y0 + 430))
    d = ImageDraw.Draw(img)
    d.rectangle((40, y0, 46, y0 + 430), fill=CORAL)
    d.text((64, y0 + 14), "固定", font=font(20), fill=CORAL)
    bullets(d, 64, y0 + 52, [
        "硅片。公开记录里没有按客户加功能的芯片。A800、H800、H20、L20、L2 是按司法辖区做的减法。",
        "Scale-up 域。NVL72 是单位。NVIDIA 自己提过的 NVL72×2 被云厂商否决。客户能否决，公开记录里没有客户另设计一个 NVIDIA scale-up 域。",
        "有合同杠杆时的东西向织物。NET-2：多节点服务的东西向 RDMA 必须是 InfiniBand 或 Spectrum-X。没有第三项。",
        "CUDA 和单一架构。Huang 反复公开的立场。",
        "信任与控制面：BlueField 南北向、TPM 2.0、安全启动、签名固件、关闭 IPMI。",
    ], font(14), INK, 840, lh=19, gap=6, mark=CORAL)
    card(img, (950, y0, 1880, y0 + 430))
    d = ImageDraw.Draw(img)
    d.rectangle((950, y0, 956, y0 + 430), fill=TEAL)
    d.text((974, y0 + 14), "可谈判", font=font(20), fill=TEAL)
    bullets(d, 974, y0 + 52, [
        "机柜机械、液冷、供电。Azure 把 Boost、自研 HSM、DC-SCM 放进 NVL72，冷却兼容机房水和风冷。",
        "CPU 插座。MGX 主机模块；NVLink Fusion 接 Fujitsu、Qualcomm。",
        "交换机整机，芯片可以留下。Meta Minipack3N 是 Spectrum-4 芯片、Meta 机箱、Accton 代工、FBOSS。",
        "客户自己做得动的 scale-out。AWS 的 EFA/SRD、Google Jupiter、Oracle Acceleron 网卡，NVIDIA 照样卖 GPU。",
        "设施到电网。DSX 把参考设计从机柜伸到电力互联，并用数字孪生消化差异。",
    ], font(14), INK, 870, lh=19, gap=6, mark=TEAL)
    card(img, (40, y0 + 448, 1880, y0 + 640), fill=NAVY, drop=False)
    d = ImageDraw.Draw(img)
    d.text((60, y0 + 462), "Huang，Computex 2025，谈 NVLink Fusion 的买法", font=font(14), fill=(150, 196, 194))
    block(d, 60, y0 + 492,
          "更可能的图景是：他们买一块 NVLink 芯粒，买 NVLink 交换机和脊，买 Spectrum-X 交换机，以及配套的全部软件。一种架构，一种 NVLink，一种网络。CPU 有时是三家的，Fujitsu 或 Qualcomm。NVIDIA 的整个生态和他们的融合在一起。",
          font(16), WHITE, 1780, lh=24, max_lines=4)
    judgement(img, (40, y0 + 656, 1880, 1010),
              "允许变化的是计算插座。收费的是互连、机柜和交换机。MGX 是一份认证菜单，不是按客户重做设计。云伙伴可以换裸金属或虚拟机，不能换 NET-2 的织物。")
    footer(img, 14, "NET-2 与伙伴规范：docs.nvidia.com/dsx/ncp。MGX：2023-05-28 新闻稿。Huang 引语来自 Computex 2025 问答，转录存于第三方站点，措辞按该转录。")
    return img


def slide_15():
    img = new()
    y0 = chrome(img, "15    客户矩阵", "九个客户，四层网络各自是谁的",
                "规律：有成熟自研网络组织的客户留下交换层。没有的，拿走整张织物。Scale-up 是 NVIDIA 目前公开记录里赢面最大的一层。")
    headers = ["客户", "Scale-up", "网卡", "交换", "跨域 / 读法"]
    rows = [
        ["OpenAI 租用", "NVLink", "NVIDIA", "NVIDIA\n经微软 / OCI / CoreWeave", "MRC，三方共研"],
        ["OpenAI 自研", "Broadcom 以太机柜", "Broadcom", "Broadcom", "10 GW ASIC\n不替换现有 NVIDIA 集群"],
        ["xAI", "NVLink", "BlueField-3", "SN5600\nSpectrum-X", "参考部署，不是定制"],
        ["微软", "NVLink", "NVIDIA", "集群内 Quantum-X800\n跨站 Spectrum-X", "MRC，自建 AI WAN"],
        ["Oracle", "NVLink", "Acceleron 自研\n也有 ConnectX", "IB + Spectrum-X", "交换机买 NVIDIA\n网卡自己做"],
        ["CoreWeave", "NVLink", "ConnectX-8", "Quantum-X800", "XGS 首批\nDSX Air"],
        ["Meta", "NVL72；MTIA 自有域", "多厂商", "Spectrum-4 芯片\n进自研机箱", "买 ASIC\n不买平台"],
        ["AWS", "NVLink 6 + MGX\nTrainium4，2025-12", "Nitro / EFA / SRD", "自研", "Scale-out 仍拒绝"],
        ["Google", "NVL72；TPU 用 ICI", "ConnectX-7\nA3 Ultra", "Jupiter + 光电路交换", "交换层不买"],
    ]
    simple_table(img, (40, y0, 1880, y0 + 820), headers, rows,
                 [230, 340, 280, 420, 570], body=13, rh=86)
    footer(img, 15, "每格的出处在第 16–18 页。Nscale 报道使用 UEC 交换机、仍采用 DSX 设施参考设计，证据等级是媒体，不进这张表。")
    return img


def slide_16():
    img = new()
    y0 = chrome(img, "16    OpenAI 与 xAI", "一个客户在改合同结构，一个客户在当参考部署",
                "OpenAI 的定制不在机柜图纸上，在资产负债表上。xAI 的选择是 NVIDIA 已经做好的那一项。")
    card(img, (40, y0, 1120, 900))
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 12), "OpenAI", font=font(20), fill=INK)
    rows = [
        ("2025-09-22", "意向书：至少 10 GW。NVIDIA 有意随每吉瓦部署投资最多 1000 亿美元。意向书写明不是合同。"),
        ("2026-02-27", "OpenAI 确认 1100 亿美元融资：软银 300 亿、NVIDIA 300 亿、亚马逊 500 亿。NVIDIA 这笔改为不绑定部署里程碑的股权。协作改为 Vera Rubin 上 3 GW 推理加 2 GW 训练。"),
        ("2026-08-17", "8-K：为俄亥俄 PORTS-Pike 约 4.25 GW IT 负载做残值担保，支付上限 1050 亿美元，预计 2028 年起。NVIDIA 是该园区独家 AI 算力基础设施提供方。OpenAI 是 8 IT-GW 的 20 年租户。"),
        ("口径差", "同一天，8-K 写可再支持约 3.8 GW，新闻稿写剩余 3.75 IT-GW。两份都是 NVIDIA 文件。"),
        ("反向定制", "2025-10-13，OpenAI 与 Broadcom：10 GW 自研加速器，机柜全部用 Broadcom 以太网。Jalapeño 为推理 ASIC。OpenAI 硬件负责人称 2026 年底小批量，不预期替换 NVIDIA 集群。"),
        ("协议共研", "微软 AI WAN 上的 Multi-Path Reliable Connected，微软称与 OpenAI、NVIDIA 联合开发。这是公开记录里最清楚的三方传输共研。"),
    ]
    yy = y0 + 48
    for k, v in rows:
        pill(d, (58, yy), k, NAVY, fnt=font(12))
        yy = block(d, 200, yy, v, font(13), INK, 880, lh=18)
        yy += 8
    card(img, (1140, y0, 1880, 900))
    d = ImageDraw.Draw(img)
    d.text((1160, y0 + 12), "xAI Colossus", font=font(20), fill=INK)
    bullets(d, 1160, y0 + 52, [
        "2024-10-28 NVIDIA 新闻稿：孟菲斯 10 万 Hopper，RDMA 用 Spectrum-X 以太，不用 InfiniBand。",
        "122 天建成。第一台机柜上架到开始训练 19 天。",
        "SN5600（Spectrum-4，51.2 Tb/s，64×800GbE）加 BlueField-3，每 GPU 400 GbE。",
        "95% 数据吞吐、全网无流碰撞丢包，是 NVIDIA 的数字，没有独立测量。",
        "论文用它说明 Time-to-AI：四个月满容。和 122 天是同一量级的两个出处。",
        "Colossus 2 的 GPU 数量来自 Musk 的帖子，织物代际没有 NVIDIA 或 xAI 的一手说明。片子不推测。",
        "读法：参考客户。差异在部署速度和自备电力，不在网络设计。选以太而不是 IB，选的是 NVIDIA 已经上市的产品。",
    ], font(14), INK, 690, lh=20, gap=8)
    judgement(img, (40, 916, 1880, 1010),
              "2026 年出现的新机制是：用担保换独家，而不是用功能裁剪换订单。俄亥俄条款买的是客户不去定制离开 NVIDIA。OpenAI 同时用 Broadcom 机柜保留离开的能力。")
    footer(img, 16, "OpenAI 2025-09-22 与 2026-02-27 博文；NVIDIA 8-K 2026-08-17 与同日新闻稿；OpenAI–Broadcom 2025-10-13；xAI 新闻稿 2024-10-28。")
    return img


def slide_17():
    img = new()
    y0 = chrome(img, "17    Meta、微软、Oracle", "三个有自研网络的客户，让步停在不同的层",
                "Meta 让到芯片。微软两张织物都用，并把自己的硅放进机柜。Oracle 买交换机、做网卡。")
    cols = [
        ("Meta    买的是 ASIC", [
            "SIGCOMM 2024：两套 24576 卡 H100，一套 400G RoCE（Arista），一套 400G NDR IB（Quantum-2），用来对照。",
            "论文写明选 RoCE 是因为专有互连限制部署灵活性。拥塞控制放在集合通信库的接收端准入，不交给 DCQCN。",
            "因此 Spectrum-X 不是 Meta 放弃 IB 的原因。RoCE 的决定早于这款平台。",
            "2025-10-13 Minipack3N：51.2 Tb/s，Spectrum-4 芯片，Meta 设计，Accton 制造，SAI + FBOSS。旁边的 Minipack3 仍是 Tomahawk。",
            "NVIDIA 标题写 standardizing on Spectrum-X。Meta 引言写的是 Spectrum Ethernet 进 FBOSS。两句都真，买的东西是芯片。",
            "2026-02 多年协议：百万颗级 Blackwell 与 Rubin，Grace CPU 进数据中心通用计算，Spectrum-X 覆盖其基础设施足迹。金额未披露。",
        ]),
        ("微软    两张织物，外加机柜内自研硅", [
            "集群内：NDv6 GB300 是首个规模化量产的 GB300 NVL72，超过 4600 颗 Blackwell Ultra，连接用 Quantum-X800 InfiniBand。",
            "跨站：Ignite 2025，Fairwater 部署下一代 Spectrum-X；全球超过 10 万颗 Blackwell Ultra 用于推理。",
            "广域：MRC，与 OpenAI、NVIDIA 共研，在微软自建的 AI WAN 上选路。光纤一年内增加超过 25%，到约 12 万英里。",
            "机柜内共研：Azure Boost、Azure Integrated HSM、DC-SCM。约 136 kW 每柜。闭环水乙二醇，机房水和风冷都能接。",
            "Fairwater 威斯康星是 GB200 NVL72。亚特兰大 2025-10 投运，两层建筑，依赖电网可靠性，现场不设发电和 UPS。",
        ]),
        ("Oracle    交换机与网卡拆开", [
            "公开集群从 H100 16384 卡、H200 65536 卡，到 Blackwell 最多 131072 卡。织物写的是 RoCEv2（ConnectX-7/8）或 Quantum-2，GB200 实例带 SHARP。",
            "Zettascale10，2025-10：最多 80 万 GPU，Acceleron RoCE 加 InfiniBand。",
            "Acceleron 是 Oracle 自研 RoCE 网卡，网卡内带四口以太网交换，做硬件多平面，用来越过单张 ASIC 的无损 RDMA 规模。",
            "同一时期，NVIDIA 2025-10-13 新闻稿把 Oracle 写成 Spectrum-X 交换机的标准化客户，并用于吉瓦级工厂。",
            "两件同时成立：交换芯片可以是 NVIDIA 的，端点网卡是 Oracle 的。这是数据集里最干净的一层拆分。",
        ]),
    ]
    cw = 600
    for i, (title, items) in enumerate(cols):
        x = 40 + i * (cw + 16)
        card(img, (x, y0, x + cw, 1010))
        d = ImageDraw.Draw(img)
        block(d, x + 16, y0 + 12, title, font(16), TEAL if i else CORAL, cw - 32, lh=20)
        bullets(d, x + 16, y0 + 64, items, font(13), INK, cw - 36, lh=18, gap=7, mark=NAVY)
    footer(img, 17, "Meta SIGCOMM 2024 论文与 2025-10-13 工程博客；NVIDIA/Meta 2026-02-17 新闻稿；Azure GB300 技术博文；Oracle Zettascale 博文。50 亿美元的 Meta 交易额是单一分析师估计，不采用。")
    return img


def slide_18():
    img = new()
    y0 = chrome(img, "18    AWS、Google、中国", "拒绝交换层的客户，以及唯一一批不同的硅",
                "中国 SKU 是按司法辖区做的减法，从来不是按客户加功能。2026 年的约束变成双边许可：美国限制卖，中国限制买。")
    card(img, (40, y0, 930, y0 + 390))
    d = ImageDraw.Draw(img)
    d.text((56, y0 + 10), "AWS    2025-12-02 的反转只发生在 scale-up", font=font(15), fill=INK)
    bullets(d, 56, y0 + 40, [
        "Scale-out 仍是 SRD，做在 Nitro 的 EFA 上。一次把一个块喷到最多 64 条路径，放松有序交付。AWS 称尾延迟大约降一个数量级，幅度是它自己的测量。",
        "理由写在 AWS 自己的 HPC 博客：以太上的投资提供的控制深度，他们不想交出去。",
        "re:Invent：Trainium4 集成 NVLink 6 和 MGX。NVIDIA 技术博客、路透、Q4 FY2026 财报电话三处一致。",
        "反转之后：scale-up 和机柜用 NVIDIA，网卡、交换、传输仍是 AWS。不采用媒体上的 72 颗、3.6 TB/s 等未出现在双方材料里的数字。",
    ], font(13), INK, 850, lh=18, gap=4)
    card(img, (950, y0, 1880, y0 + 390))
    d = ImageDraw.Draw(img)
    d.text((966, y0 + 10), "Google    交换层的成本论证是公开的", font=font(15), fill=INK)
    bullets(d, 966, y0 + 40, [
        "NSDI 2024：TPUv4 的光电路交换加光纤，资本开支不到 pod 的 5%，运行功耗不到 3%。作者直接对比 NVLink 上的两层 NVSwitch 胖树。",
        "GPU 虚机仍买 NVIDIA 的端点：A3 Ultra 用 ConnectX-7，每服务器 3.2 Tb/s 无阻塞 RoCE，并预告 GB200 NVL72。",
        "交换留在 Jupiter。买网卡和 scale-up 域，不买参数面交换机。",
        "Tesla 是失败的垂直整合回到 NVIDIA：Dojo 团队 2025-08 解散，训练回到 GPU。Dojo3 的细节没有可引用的规格。",
    ], font(13), INK, 880, lh=18, gap=4)
    headers = ["部件", "时间", "砍掉的", "留下的"]
    rows = [
        ["A800", "2022-11", "NVLink 从 A100 的 600 降到 400 GB/s", "为了卡在当时的互连门槛下"],
        ["H800", "2023-03", "NVLink 从 900 降到约 400 GB/s", "算力不动。Epoch 分析：400 GB/s 仍够把通信藏在计算后面"],
        ["H20", "2023-11", "BF16 算力砍到 148 TFLOPS", "NVLink 保持 900 GB/s，显存升到 96 GB。规则改考算力之后，砍法反过来"],
        ["许可", "2025–2026", "2025-08 白宫确认 15% 中国芯片销售收入上缴以换许可", "2026-05 约 10 家获准买 H200，当时零交付。2026-08 NVIDIA 确认首批运抵，不到当季数据中心收入的 1%"],
    ]
    simple_table(img, (40, y0 + 406, 1880, y0 + 760), headers, rows,
                 [140, 200, 700, 800], body=13, rh=80)
    footer(img, 18, "AWS 博客与 2025-12-02 路透；Google NSDI 2024；H800/H20：路透与 Epoch AI；15%：白宫 2025-08-11；H200 首批：NVIDIA 2026-08-27。字节、腾讯各约 1 万片是 FT 报道，不写进表。")
    return img


def slide_19():
    img = new()
    y0 = chrome(img, "19    策略合上", "定制策略和标准策略是同一条边界",
                "性能层保持端到端配对。接口层保持能够连通的外观。客户越大，NVIDIA 越愿意把壳让出去，换硅片和 scale-up 留在屋里。")
    shifts = [
        ("01", "默认项换成以太",
         "Vera Rubin DSX 参考设计写的是 Spectrum-X Ethernet。Meta、Oracle 在 2025-10 采用，微软跨站采用，xAI 在 2024 年就已经用它跑满 10 万卡。InfiniBand 留在 Azure、CoreWeave 这类要点名 SHARP 的集群内。"),
        ("02", "从卖系统到担保地产",
         "OpenAI：1000 亿美元意向书，收成 300 亿美元股权，再加上 1050 亿美元租赁残值担保，换俄亥俄园区的独家算力提供方。资本背书的是客户的资产负债表。"),
        ("03", "NVLink 进竞争对手的机柜",
         "拒绝 NVIDIA 网络最久的 AWS，在 2025-12 接受 NVLink 6 和 MGX 来做 Trainium4。互连开始出现在用来替代 NVIDIA GPU 的芯片旁边。Scale-out 仍然是对方的。"),
    ]
    cw = 600
    for i, (n, t, body) in enumerate(shifts):
        x = 40 + i * (cw + 16)
        card(img, (x, y0, x + cw, y0 + 280))
        d = ImageDraw.Draw(img)
        d.text((x + 16, y0 + 14), n, font=font(14), fill=GOLD)
        d.text((x + 56, y0 + 12), t, font=font(18), fill=INK)
        block(d, x + 16, y0 + 50, body, font(14), INK, cw - 32, lh=20, max_lines=8)
    card(img, (40, y0 + 296, 1880, y0 + 470))
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 310), "收入背景，用来衡量这条边界值多少", font=font(15), fill=INK)
    block(d, 58, y0 + 340,
          "Q4 FY2026 网络业务单季 110 亿美元，同比增长超过 3.5 倍；全年超过 310 亿美元。Huang：Ethernet has been a home run。Q2 FY2027 网络环比再增 18%，Spectrum-X 以太同比 2.6 倍。每吉瓦收入机会从 Hopper 约 180 亿美元、Blackwell 250 亿，到 Vera Rubin 400 亿美元。他的效率论证是：利用率从大约 65% 提到 85–90%，相对一座 500 亿美元的工厂，网络等于免费。这是厂商论证。",
          font(14), INK, 1780, lh=20, max_lines=5)
    card(img, (40, y0 + 486, 1880, 1010), fill=SOFT, drop=False)
    d = ImageDraw.Draw(img)
    d.text((58, y0 + 500), "公开信息不足，本套片子不采用", font=font(15), fill=CORAL)
    skipped = [
        "NVLink Fusion 每套部署必须含至少一件 NVIDIA 产品。多家媒体一致，没有 NVIDIA 文件。",
        "Kyber NVL144 延期到 2028。单一分析师来源。NVIDIA 称路线图不变。",
        "Colossus 2 使用哪一代织物。没有一手来源。",
        "B30A / B40 的规格、名字和价格。来源互相矛盾，NVIDIA 未确认。",
        "Trainium4 的 72 颗、每芯片 3.6 TB/s。不在 NVIDIA 或 AWS 的材料里。",
        "Spectrum-X 是 AI 工厂神经系统。只追到一篇分析博客，没有演讲原文。",
        "Meta 交易 500 亿美元、Abilene 600 MW 取消。前者是分析师估计，后者 Oracle 公开否认。",
    ]
    # two columns of skipped
    mid = 4
    bullets(d, 58, y0 + 530, skipped[:4], font(13), INK, 860, lh=18, gap=3, mark=CORAL)
    bullets(d, 980, y0 + 530, skipped[4:], font(13), INK, 860, lh=18, gap=3, mark=CORAL)
    footer(img, 19, "财务数字：Q4 FY2026 与 Q2 FY2027 财报电话。不采用清单与调研笔记 docs/nvidia-param-plane.md 第 8 节对应。")
    return img


SLIDES = [
    slide_01, slide_02, slide_03, slide_04, slide_05,
    slide_06, slide_07, slide_08, slide_09, slide_10,
    slide_11, slide_12, slide_13, slide_14, slide_15,
    slide_16, slide_17, slide_18, slide_19,
]


def pack(pngs, path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for p in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(p), Emu(0), Emu(0), width=prs.slide_width, height=prs.slide_height)
    prs.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    pngs = []
    for i, fn in enumerate(SLIDES, 1):
        img = fn()
        path = OUT / f"{i:02d}.png"
        img.save(path, "PNG")
        pngs.append(path)
        print("wrote", path.name, img.size)
    pptx = Path("/workspace/docs/nvidia-param-plane.pptx")
    pack(pngs, pptx)
    print("pptx", pptx)


if __name__ == "__main__":
    main()
