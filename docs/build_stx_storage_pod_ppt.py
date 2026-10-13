#!/usr/bin/env python3
"""Editable diagram: NVMe, Vera CPU, and ConnectX-9 on BlueField-4 STX."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)
TEAL = (14, 122, 118)
GOLD = (122, 84, 28)
GREEN = (27, 110, 72)
LINE = (214, 220, 226)
SOFT = (244, 247, 250)
RACK = (232, 240, 246)
OPTIC = (88, 62, 130)


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    strip_shadow(shape)


def write(shape, text, size, bold, color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def box(slide, x, y, w, h, fill, text="", size=12, bold=True, color=WHITE, line=None, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    if text:
        write(shape, text, size, bold, color, align)
    return shape


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def link(slide, x1, y1, x2, y2, color):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(2.25)
    ln = shape.line._get_or_add_ln()
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "sm")
    ln.append(tail)
    return shape


def chip(slide, x, y, w, h, fill, text, fg=NAVY):
    box(slide, x, y, w, h, fill, text, 11, True, fg, line=fill)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "盘、Vera CPU 与 ConnectX-9 的连接"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box(slide, 0, 0, 13.333, 0.08, ACCENT, "")
    textbox(slide, 0.40, 0.20, 12.5, 0.24, "BlueField-4 STX", 14, True, ACCENT)
    textbox(slide, 0.40, 0.48, 12.5, 0.42, "盘用 PCIe 接到这颗处理器，CPU 和 ConnectX-9 在片内相连", 26, True, NAVY)
    textbox(
        slide,
        0.40,
        1.00,
        12.5,
        0.36,
        "没有另挂一颗主机 CPU。网口只在 ConnectX-9 上。",
        16,
        False,
        INK,
    )

    # Three nodes on one line. The frame only groups the two functions of one chip.
    disk_x, disk_w, disk_h = 0.40, 2.45, 2.20
    frame_x, frame_w, frame_h = 4.50, 8.45, 3.15
    cpu_x, cpu_w, cpu_h = 4.85, 3.15, 2.20
    cx_x, cx_w = 9.55, 3.05
    node_y = 2.55
    frame_y = 2.05
    mid_y = node_y + cpu_h / 2

    box(slide, frame_x, frame_y, frame_w, frame_h, RACK, line=ACCENT)
    textbox(slide, frame_x + 0.22, frame_y + 0.08, frame_w - 0.44, 0.32, "同一颗 BlueField-4 STX", 16, True, ACCENT)

    box(slide, disk_x, node_y, disk_w, disk_h, GOLD, "")
    textbox(slide, disk_x + 0.14, node_y + 0.48, disk_w - 0.28, 0.46, "NVMe 盘", 26, True, WHITE)
    textbox(slide, disk_x + 0.14, node_y + 1.10, disk_w - 0.28, 0.80, "KV 所在的闪存", 16, False, WHITE, anchor=MSO_ANCHOR.TOP)

    box(slide, cpu_x, node_y, cpu_w, cpu_h, TEAL, "")
    textbox(slide, cpu_x + 0.14, node_y + 0.48, cpu_w - 0.28, 0.46, "Vera CPU", 26, True, WHITE)
    textbox(slide, cpu_x + 0.14, node_y + 1.10, cpu_w - 0.28, 0.80, "存储处理\n元数据、加密、压缩", 16, False, WHITE, anchor=MSO_ANCHOR.TOP)

    box(slide, cx_x, node_y, cx_w, cpu_h, ACCENT, "")
    textbox(slide, cx_x + 0.14, node_y + 0.48, cx_w - 0.28, 0.46, "ConnectX-9", 24, True, WHITE)
    textbox(slide, cx_x + 0.14, node_y + 1.10, cx_w - 0.28, 0.80, "网口\n每端口最高 800 Gb/s", 16, False, WHITE, anchor=MSO_ANCHOR.TOP)

    # Disk lands on the CPU side of this processor. The label stays in the open gap.
    link(slide, disk_x + disk_w, mid_y, cpu_x, mid_y, GOLD)
    chip(slide, 3.08, mid_y - 0.20, 1.18, 0.40, WHITE, "PCIe Gen6", GOLD)

    link(slide, cpu_x + cpu_w, mid_y, cx_x, mid_y, TEAL)
    chip(slide, 8.22, mid_y - 0.18, 1.10, 0.36, WHITE, "片内", TEAL)

    textbox(
        slide,
        0.40,
        5.55,
        12.5,
        1.45,
        "盘接到这颗处理器的 PCIe，规格为 Gen6、96 通道。\n"
        "Vera CPU 与 ConnectX-9 做在同一颗上，中间是片内连接，不再拉线缆。\n"
        "网口只在 ConnectX-9 上。数据离开这颗处理器，只从这只网口出去。",
        18,
        False,
        INK,
        anchor=MSO_ANCHOR.TOP,
    )

    slide.notes_slide.notes_text_frame.text = (
        "依据 BlueField-4 STX 存储处理器产品说明，以及 NVIDIA Vera Rubin POD 技术博客（2026-03-19）。"
        "BlueField-4 STX 是自托管存储处理器：Vera CPU 与 ConnectX-9 做在同一颗上，数据路径上没有另挂一颗主机 CPU。"
        "NVMe 盘接到这颗处理器的 PCIe，规格为 Gen6、96 通道。图中这根线停在处理器边界，不表示每个字节都要由 CPU 核再拷贝一次。"
        "Vera CPU 与 ConnectX-9 之间是片内连接，中间不再拉一根线缆。CPU 侧做存储处理，包括元数据、加密、压缩、完整性与恢复。"
        "网口只在 ConnectX-9 上，面向 Spectrum-X 以太网，每端口最高 800 Gb/s。数据离开这颗处理器，只从这只网口出去。"
    )

    out = "/workspace/docs/stx-storage-pod.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
