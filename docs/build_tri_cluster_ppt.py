#!/usr/bin/env python3
"""One-page 3-column slide: ByteDance / Alibaba / Tencent AI-cluster generations."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Emu

W, H = 1920, 1080
OUT = Path("/workspace/docs")
DIR = OUT / "tri-ai-cluster-gens"
FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

BG = (246, 244, 239)
INK = (27, 36, 48)
MUTED = (92, 102, 112)
TEAL = (15, 110, 107)
CORAL = (196, 73, 58)
GOLD = (184, 138, 58)
NAVY = (28, 49, 68)
CARD = (255, 255, 255)
LINE = (220, 214, 204)
SOFT = (237, 233, 224)
OK = (42, 122, 75)


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def tw(d, t, f):
    return int(d.textlength(t, font=f))


def wrap(d, text, fnt, max_w):
    lines = []
    for para in text.split("\n"):
        buf = ""
        for ch in para:
            if d.textlength(buf + ch, font=fnt) <= max_w:
                buf += ch
            else:
                if buf:
                    lines.append(buf)
                buf = ch
        if buf:
            lines.append(buf)
    return lines


def canvas():
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    for i in range(0, W, 48):
        d.line([(i, 0), (i, H)], fill=(236, 232, 224), width=1)
    for j in range(0, H, 48):
        d.line([(0, j), (W, j)], fill=(236, 232, 224), width=1)
    return img


def shadow(base, box, r=14):
    x0, y0, x1, y1 = box
    layer = Image.new("RGBA", base.size, (0, 0, 0, 0))
    ImageDraw.Draw(layer).rounded_rectangle(
        (x0 + 3, y0 + 4, x1 + 3, y1 + 6), radius=r, fill=(20, 24, 30, 24)
    )
    blur = layer.filter(ImageFilter.GaussianBlur(6))
    tmp = base.convert("RGBA")
    tmp.alpha_composite(blur)
    base.paste(tmp.convert("RGB"))


def card(img, box, fill=CARD, r=14, outline=LINE, w=1, drop=True):
    if drop:
        shadow(img, box, r)
    ImageDraw.Draw(img).rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=w)


def pill(d, xy, text, bg, fg=(255, 255, 255), fnt=None, pad_x=10, pad_y=4):
    fnt = fnt or font(13)
    x, y = xy
    box = (x, y, x + tw(d, text, fnt) + pad_x * 2, y + 18 + pad_y * 2)
    d.rounded_rectangle(box, radius=11, fill=bg)
    d.text((x + pad_x, y + pad_y + 1), text, font=fnt, fill=fg)
    return box


def sw(d, xy, color, r=7):
    x, y = xy
    d.rounded_rectangle((x - r, y - 5, x + r, y + 5), radius=3, fill=color)


def header(img):
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 7), fill=TEAL)
    d.text((28, 16), "三家三代  ·  公开口径对照", font=font(15), fill=TEAL)
    d.text((28, 38), "字节 / 阿里 / 腾讯：近三代 AI 集群怎么建、往哪走", font=font(30), fill=INK)
    d.text(
        (28, 80),
        "场景从训练专网走到训推一体与 PD；组网从三层多轨走向少层 / 超节点；协议从 RoCE 集合通信扩到 Scale-up 与 Token 面。",
        font=font(16),
        fill=MUTED,
    )
    d.line([(28, 110), (W - 28, 110)], fill=LINE, width=1)


def footer(img):
    d = ImageDraw.Draw(img)
    d.line([(28, 1044), (W - 28, 1044)], fill=LINE, width=1)
    d.text((28, 1050), "仅公开论文 / 官网 / 大会口径，不发明未披露数字。", font=font(12), fill=MUTED)
    d.text(
        (28, 1066),
        "HPN8 发布会另有「10万卡 / GPU互联 6.4T / 存储 800G」；WAIC 为「单集群最高13万卡、可扩百万」。星脉1.0 按开发者文 1.6T / 2K / 32K，不与 2023 HCC「3.2T / 10万卡」混用。",
        font=font(12),
        fill=MUTED,
    )


def roadmap(d, x0, y0, x1, color, gens):
    """Three-node horizontal roadmap."""
    n = 3
    y = y0 + 28
    xs = [x0 + 70 + i * ((x1 - x0 - 140) / (n - 1)) for i in range(n)]
    d.line([(xs[0], y), (xs[-1], y)], fill=LINE, width=3)
    for i, (lab, sub, year) in enumerate(gens):
        c = color if i == 2 else (GOLD if i == 1 else MUTED)
        d.ellipse((xs[i] - 8, y - 8, xs[i] + 8, y + 8), fill=c)
        d.text((xs[i], y0 + 2), year, font=font(12), fill=c, anchor="mm")
        d.text((xs[i], y + 16), lab, font=font(14), fill=INK, anchor="mm")
        d.text((xs[i], y + 34), sub, font=font(11), fill=MUTED, anchor="mm")


def draw_byte_arch(img, box):
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = box
    d.rounded_rectangle(box, radius=10, fill=SOFT)
    d.text((x0 + 12, y0 + 8), "最新架构  ·  Rack 3.0 + EthLink + HPN 6.0", font=font(13), fill=TEAL)

    # scale-up racks
    rx, ry = x0 + 18, y0 + 40
    d.text((rx, ry), "Scale-up", font=font(12), fill=TEAL)
    for i, lab in enumerate(("计算柜×8", "交换柜×2")):
        bx = rx + i * 88
        fill = TEAL if i == 0 else NAVY
        d.rounded_rectangle((bx, ry + 20, bx + 78, ry + 88), radius=6, fill=fill)
        d.text((bx + 39, ry + 42), lab, font=font(12), fill=(255, 255, 255), anchor="mm")
        d.text((bx + 39, ry + 62), "NPO", font=font(11), fill=(210, 232, 230), anchor="mm")
    d.text((rx + 78, ry + 108), "1024 XPU 超节点", font=font(12), fill=INK, anchor="mm")
    d.text((rx + 78, ry + 126), "铜互联双柜 576", font=font(11), fill=MUTED, anchor="mm")
    pill(d, (rx, ry + 142), "EthLink  Ld/St+RDMA", TEAL, fnt=font(11))

    # 3-tier clos — sparse links, readable hierarchy
    cx = x0 + 230
    d.text((cx, ry), "Scale-out  三层 Clos", font=font(12), fill=TEAL)
    cores = [cx + 50, cx + 110, cx + 170]
    aggs = [cx + 30, cx + 90, cx + 150, cx + 190]
    tors = [cx + 18 + i * 32 for i in range(6)]
    for x in cores:
        sw(d, (x, ry + 28), NAVY, 10)
    for i, x in enumerate(aggs):
        sw(d, (x, ry + 70), TEAL, 9)
        d.line([(x, ry + 65), (cores[min(i, 2)], ry + 33)], fill=(150, 175, 174), width=2)
        if i > 0:
            d.line([(x, ry + 65), (cores[max(i - 1, 0)], ry + 33)], fill=(190, 208, 206), width=1)
    for i, x in enumerate(tors):
        sw(d, (x, ry + 112), GOLD, 7)
        d.line([(x, ry + 107), (aggs[min(i // 2 + i % 2, 3)], ry + 75)], fill=(205, 190, 150), width=1)
    d.text((cx + 110, ry + 46), "Core", font=font(11), fill=MUTED, anchor="mm")
    d.text((cx + 110, ry + 86), "Agg", font=font(11), fill=MUTED, anchor="mm")
    d.text((cx + 110, ry + 128), "ToR  800G 混速", font=font(11), fill=MUTED, anchor="mm")
    d.text((cx + 110, ry + 152), "POD 65k  →  集群百万级", font=font(13), fill=INK, anchor="mm")
    d.text((cx + 110, ry + 170), "102.4T / 128×800G", font=font(12), fill=TEAL, anchor="mm")


def draw_ali_arch(img, box):
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = box
    d.rounded_rectangle(box, radius=10, fill=SOFT)
    d.text((x0 + 12, y0 + 8), "最新架构  ·  HPN 8.0 + UPN512 + TPN", font=font(13), fill=CORAL)

    layers = [
        (CORAL, "TPN  Token 面", "两层融合网", "接入 +2.5×   规模 +10×   时延 −1/3"),
        (NAVY, "HPN 8.0  Scale-out", "多平面 CLOS", "单集群最高 13万卡混布  →  百万卡"),
        (GOLD, "UPN512  Scale-up", "单层光 CLOS", "512 xPU   LPO/NPO   可用 +3×  成本 −30%"),
    ]
    y = y0 + 36
    for color, title, topo, metric in layers:
        d.rounded_rectangle((x0 + 14, y, x1 - 14, y + 46), radius=8, fill=CARD, outline=color, width=2)
        d.rectangle((x0 + 14, y, x0 + 20, y + 46), fill=color)
        d.text((x0 + 32, y + 6), title, font=font(14), fill=INK)
        d.text((x0 + 210, y + 8), topo, font=font(12), fill=color)
        d.text((x0 + 32, y + 26), metric, font=font(12), fill=MUTED)
        # mini switch row
        for i in range(5):
            sw(d, (x1 - 120 + i * 18, y + 23), color, 6)
        y += 52
    d.text((x0 + (x1 - x0) // 2, y1 - 16), "训推一体  ·  PD 分离  ·  Token / KV 与集合通信分面", font=font(12), fill=MUTED, anchor="mm")


def draw_tx_arch(img, box):
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = box
    d.rounded_rectangle(box, radius=10, fill=SOFT)
    d.text((x0 + 12, y0 + 8), "最新架构  ·  星脉 3.0  光 Shuffle 二级单轨", font=font(13), fill=NAVY)

    # GPU row
    gy = y0 + 36
    d.text((x0 + 16, gy + 10), "GPU", font=font(12), fill=NAVY)
    gxs = [x0 + 68 + i * 50 for i in range(6)]
    for x in gxs:
        d.rounded_rectangle((x - 14, gy + 8, x + 14, gy + 30), radius=4, fill=NAVY)

    # optical shuffle as a woven band
    sy0, sy1 = gy + 36, gy + 68
    d.rounded_rectangle((gxs[0] - 18, sy0, gxs[-1] + 18, sy1), radius=6, fill=(252, 246, 230), outline=GOLD, width=1)
    d.text((x0 + 16, (sy0 + sy1) / 2), "光\nShuffle", font=font(11), fill=GOLD, anchor="lm")
    for i, gx in enumerate(gxs):
        d.line([(gx, sy0 + 3), (gxs[(i + 2) % 6], sy1 - 3)], fill=GOLD, width=2)
        d.line([(gx, sy0 + 3), (gxs[(i + 3) % 6], sy1 - 3)], fill=(210, 180, 110), width=1)

    ly = gy + 86
    for x in gxs:
        sw(d, (x, ly), GOLD, 8)
        d.line([(x, sy1), (x, ly - 8)], fill=GOLD, width=2)
    d.text((x0 + 16, ly), "Leaf", font=font(12), fill=GOLD, anchor="lm")

    spy = gy + 120
    spines = (gxs[1], gxs[2] + 24, gxs[4])
    for x in spines:
        sw(d, (x, spy), NAVY, 10)
    for i, lx in enumerate(gxs):
        d.line([(lx, ly + 8), (spines[i % 3], spy - 6)], fill=(170, 180, 196), width=2)
    d.text((x0 + 16, spy), "Spine", font=font(12), fill=NAVY, anchor="lm")

    pill(d, (x0 + 16, y1 - 32), "Prefill/训练 高带宽A2A", CORAL, fnt=font(11))
    pill(d, (x0 + 200, y1 - 32), "Decode 低时延A2A", TEAL, fnt=font(11))
    pill(d, (x0 + 360, y1 - 32), "TRMT", NAVY, fnt=font(11))


def table(img, box, color, rows):
    """rows: list of (label, g1, g2, g3)."""
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = box
    n = len(rows)
    lab_w = 62
    inner = x1 - x0 - lab_w
    cw = inner / 3
    rh = (y1 - y0) / n
    heads = ("一代", "二代", "三代")
    # header strip is first visual of each cell via bold gen in data; draw col titles above
    d.rectangle((x0, y0, x1, y0 + 22), fill=color)
    d.text((x0 + 8, y0 + 11), "维度", font=font(12), fill=(255, 255, 255), anchor="lm")
    for i, h in enumerate(heads):
        d.text((x0 + lab_w + cw * i + cw / 2, y0 + 11), h, font=font(12), fill=(255, 255, 255), anchor="mm")
    y = y0 + 22
    rh = (y1 - y0 - 22) / n
    for i, (lab, a, b, c) in enumerate(rows):
        bg = CARD if i % 2 == 0 else SOFT
        d.rectangle((x0, y, x1, y + rh), fill=bg)
        d.line([(x0, y), (x1, y)], fill=LINE, width=1)
        d.rectangle((x0, y, x0 + lab_w, y + rh), fill=(color[0], color[1], color[2],))
        # mute the label bar
        d.rectangle((x0, y, x0 + lab_w, y + rh), fill=(245, 243, 238))
        d.text((x0 + lab_w / 2, y + rh / 2), lab, font=font(13), fill=color, anchor="mm")
        cells = (a, b, c)
        fnt = font(12)
        for j, text in enumerate(cells):
            cx0 = x0 + lab_w + cw * j
            lines = wrap(d, text, fnt, cw - 12)
            total = len(lines) * 15
            ty = y + (rh - total) / 2
            ink = color if j == 2 else INK
            for line in lines:
                d.text((cx0 + cw / 2, ty), line, font=fnt, fill=ink, anchor="mm")
                ty += 15
        y += rh
    d.line([(x0, y1), (x1, y1)], fill=LINE, width=1)
    d.line([(x0, y0), (x0, y1)], fill=LINE, width=1)
    d.line([(x1, y0), (x1, y1)], fill=LINE, width=1)
    for i in range(1, 3):
        xx = x0 + lab_w + cw * i
        d.line([(xx, y0 + 22), (xx, y1)], fill=LINE, width=1)
    d.line([(x0 + lab_w, y0), (x0 + lab_w, y1)], fill=LINE, width=1)


def vendor_col(img, box, color, name, tag, gens, rows, arch_fn):
    x0, y0, x1, y1 = box
    card(img, box, r=16, outline=color, w=2)
    d = ImageDraw.Draw(img)
    d.rectangle((x0, y0, x1, y0 + 8), fill=color)
    d.text((x0 + 16, y0 + 20), name, font=font(22), fill=INK)
    pill(d, (x0 + 16, y0 + 50), tag, color, fnt=font(12))
    roadmap(d, x0 + 8, y0 + 80, x1 - 8, color, gens)
    arch = (x0 + 12, y0 + 150, x1 - 12, y0 + 368)
    arch_fn(img, arch)
    table(img, (x0 + 12, y0 + 380, x1 - 12, y1 - 12), color, rows)


def slide():
    img = canvas()
    header(img)

    cols = [
        (
            TEAL,
            "字节跳动",
            "最新：Rack 3.0 + HPN 6.0",
            [
                ("MegaScale", "三层 Clos 1:1", "NSDI'24"),
                ("Rack 2.0", "256 XPU + PD", "2025"),
                ("Rack 3.0", "超节点 + 65k", "2026"),
            ],
            [
                ("场景", "预训练万卡", "训练 + xLLM PD", "训推一体 / 混速多代"),
                ("组网", "三层 Clos 1:1\n8×200G 多轨", "双柜 256 XPU\nRDMA scale-out", "576 铜 / NPO 1024\n三层 Clos  POD 65k"),
                ("协议", "RoCEv2  400G AOC", "RDMA + NVLink 域", "EthLink Ld/St+RDMA\n200/400/800G 混速"),
                ("关键\n技术", "TH4 25.6T\n12288 卡 MFU 55.2%", "PD 吞吐最高 5×\n超节点 256 XPU", "102.4T / SyncMesh μs\n算子+任务 QoS"),
            ],
            draw_byte_arch,
        ),
        (
            CORAL,
            "阿里云",
            "最新：HPN8 + UPN + TPN",
            [
                ("灵骏早期", "万卡 ETH RDMA", "2022–23"),
                ("HPN 7.0", "1k / 15k 两层", "SIGCOMM'24"),
                ("HPN8+UPN+TPN", "13万 → 百万", "WAIC/云栖"),
            ],
            [
                ("场景", "训练万卡", "训练专网  存算分离", "训推一体 / PD / Token"),
                ("组网", "ETH RDMA 双平面\n前后端分离", "双 ToR 双平面\n1k 一跳 / ~15k 两层", "多平面 CLOS 13万\nUPN 单层 512 xPU"),
                ("协议", "RoCE + HPCC", "400G RoCE\n自研 51.2T", "800G / IPv6 Native\nTPN 两层 Token 面"),
                ("关键\n技术", "存算流量分网", "AllReduce +59.3%\n排队 −91.8%  8+月", "TPN +2.5×/+10×/−1/3\n分钟自愈  可用 99.7%"),
            ],
            draw_ali_arch,
        ),
        (
            NAVY,
            "腾讯云",
            "最新：星脉 3.0 规划",
            [
                ("星脉 1.0", "Fat-Tree 1.6T", "2023"),
                ("星脉 2.0", "3.2T 同轨三层", "2024–25"),
                ("星脉 3.0", "光 Shuffle", "2026"),
            ],
            [
                ("场景", "混元训练", "训练扩到十万卡", "MoE 训推一体 / PD 核"),
                ("组网", "Fat-Tree 多轨 1.6T\n典型 2K / 最大 32K", "同轨聚合 三层等带宽\nPod 64k / 集群 ~512k", "光 Shuffle\n扁平二级单轨"),
                ("协议", "ETH RDMA\nTiTa + TCCL", "TiTa2 网卡主动 CC\nTCCL2 NVLink+NET", "TRMT + RoCEv2\nPrefill/Decode 分核"),
                ("关键\n技术", "GPU 利用 +40%\n时延 −40%", "训练 +20% / 通信 +60%\nMoE 8K 效率损 0.6%", "RoCE A2A +100%\nIB +30%  双引擎规划"),
            ],
            draw_tx_arch,
        ),
    ]

    gap = 14
    x = 24
    cw = (W - 48 - gap * 2) / 3
    for color, name, tag, gens, rows, arch in cols:
        vendor_col(img, (x, 122, x + cw, 1040), color, name, tag, gens, rows, arch)
        x += cw + gap

    footer(img)
    return img


def pack(path):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    out = OUT / "tri-ai-cluster-gens.pptx"
    prs.save(out)
    return out


def main():
    DIR.mkdir(parents=True, exist_ok=True)
    img = slide()
    png = DIR / "01-tri-cluster.png"
    img.save(png, "PNG", optimize=True)
    print("wrote", png)
    print("pptx", pack(png))


if __name__ == "__main__":
    main()
