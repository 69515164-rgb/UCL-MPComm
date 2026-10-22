#!/usr/bin/env python3
"""What the KV control plane exchanges, and whether that wait belongs on a DPU."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def chrome(slide, kicker, title, page):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.14, 10.2, 0.22, kicker, 12, True, ACCENT)
    textbox(slide, 0.42, 0.36, 11.5, 0.42, title, 24, True, NAVY)
    textbox(slide, 11.2, 0.16, 1.7, 0.22, page, 12, False, MUTED, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.22,
        12.5,
        0.22,
        "0.2 ms 与 100 GB/s 是估算用的数量级    ·    186 MB / 14 ms 来自 vLLM-Omni 在 H800 + mlx5 上的公开墙钟",
        11,
        False,
        MUTED,
    )


def banner(slide, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(0.86), Inches(12.5), Inches(0.40))
    solid(shape, PALE)
    textbox(slide, 0.56, 0.86, 12.22, 0.40, text, 14, True, TEAL)


def add_table(slide, x, y, widths, rows, font=12, head_fill=NAVY, head_h=0.40, row_h=0.58):
    h = head_h + row_h * (len(rows) - 1)
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(h))
    table = shape.table
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(head_h if r == 0 else row_h)
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = value
            if r == 0:
                set_run_font(run, font, True, WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(head_fill)
            else:
                set_run_font(run, font, c == 0, INK if c == 0 else MUTED)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(WHITE if r % 2 else SOFT)
            for side in ("lnL", "lnR", "lnT", "lnB"):
                ln = OxmlElement(f"a:{side}")
                sf = OxmlElement("a:solidFill")
                srgb = OxmlElement("a:srgbClr")
                srgb.set("val", LINE)
                sf.append(srgb)
                ln.append(sf)
                cell._tc.get_or_add_tcPr().append(ln)
    return shape


def slide_info(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "Decode KV    ·    控制面载荷", "每次挡住数据面的，是这条 KV 的目标地址", "1  /  2")
    banner(slide, "内存注册、rkey 和 QP 只挡第一次。预热之后，post_send 前面只剩这次传输的目的地址。")
    add_table(
        slide,
        0.42,
        1.42,
        [2.55, 6.15, 3.80],
        [
            ["控制消息", "交换的信息", "何时挡在 post_send 前"],
            ["段描述 mooncake/ram", "addr、length、每张网卡一个 rkey，以及 lid、gid", "第一次 openSegment"],
            ["QP 握手", "nic path、qp_num、lid、gid、ready_ack", "该对端的第一条连接"],
            ["PD 的 ZMQ", "hostname、rpc_port、目的 kv base、block id、长度", "每一次读前缀"],
            ["GetReplicaList / PutStart", "key、endpoint、buffer_address、租约或 PROCESSING", "每一次池化读或写"],
            ["CQ / TRANS_DONE / PutEnd", "完成状态，或副本改为 COMPLETE", "在数据面之后"],
        ],
        font=13,
        head_h=0.46,
        row_h=0.72,
    )
    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(5.58), Inches(12.5), Inches(1.46))
    solid(note, GOLD_BG)
    textbox(
        slide,
        0.58,
        5.58,
        12.18,
        1.46,
        "TransferRequest 不出网。它留在发起方，字段是 opcode、source、segment、offset、length。\n"
        "RDMA 工作请求里的远端地址和 rkey 是数据面描述符，随 post_send 交给本端网卡，不再做一次软件往返。imm_data 为 0。\n"
        "ZMQ 在 WRITE 之后回的 TRANS_DONE，以及 Master 上的 PutEnd，挡住的是下一步 attention 或后续读者，不挡住这一次发送。",
        14,
        False,
        GOLD,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    slide.notes_slide.notes_text_frame.text = (
        "段描述来自 TransferEngine 元数据 mooncake/ram/[segment]，缓冲项含 addr、length、rkey 数组，设备项含 name、lid、gid。"
        "rpc_meta 另有 ip_or_host_name 和 rpc_port，握手前用来找到对端。"
        "HandShakeDesc 含 local_nic_path、local_lid、local_gid、peer_nic_path、qp_num、ready_ack、reply_msg。"
        "MC_HANDSHAKE_CONNECT_TIMEOUT 默认 5 秒，是连不上时的上限。MC_HANDSHAKE_MAX_LENGTH 默认 1 MB，限制的是建连报文，不是每条 KV。"
        "段描述会缓存在本端。MC_TE_METADATA_REFRESH_INTERVAL_SECONDS 默认 0，不后台刷新。"
        "vLLM MooncakeConnector 的 receive_kv 先把 MooncakeAgentMetadata 经 ZMQ 发给 Prefill，字段有 remote_hostname、remote_port、"
        "kv_caches_base_addr、block_ids、request_ids。Prefill 算出目的地址后才 batch_transfer_sync_write。"
        "Mooncake Store 的 GetReplicaList 返回 replicas、lease_ttl_ms。MEMORY 副本带 transport_endpoint 和 buffer_address。"
        "PutStart 分配缓冲并把副本标为 PROCESSING，PutEnd 才标 COMPLETE。"
        "因此稳态下挡在数据面前面的软件往返，是 PD 的地址消息，或池化路径的 Master RPC。"
    )


def card(slide, x, y, w, h, title, body, fill, title_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, LINE)
    textbox(slide, x + 0.16, y + 0.08, w - 0.32, 0.30, title, 14, True, title_color)
    textbox(slide, x + 0.16, y + 0.40, w - 0.32, h - 0.50, body, 13, False, INK, anchor=MSO_ANCHOR.TOP)


def slide_judge(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "Decode KV    ·    这段等待与 DPU", "长前缀的时间在数据面，短块才被地址等待盖过", "2  /  2")
    banner(slide, "控制消息是 KB 级。每次传输前都要先完成这一次往返，这个来回才会成为瓶颈。")
    card(
        slide,
        0.42,
        1.42,
        6.16,
        2.55,
        "判据",
        "T_data = KV 字节 / 有效带宽\n"
        "地址 RPC 的时间记为 T_ctrl。\n"
        "T_ctrl 大于 T_data 时，这段等待盖过传输。\n"
        "估算：T_ctrl 取 0.2 ms，带宽取 100 GB/s，交叉点约 20 MB。\n"
        "公开墙钟：约 186 MB 的 KV，TransferEngine 路径约 14 ms。",
        SOFT,
        NAVY,
    )
    card(
        slide,
        6.76,
        1.42,
        6.16,
        2.55,
        "谁会真的挡住",
        "长前缀，几百 MB 到数 GB，时间在网卡搬字节。\n"
        "单层、少量 block，传输落到 20 MB 以下，一次地址往返就占大头。\n"
        "冷握手不在稳态路径上。连不上时的上限是 5 秒。\n"
        "Master 排队随 QPS 变长，和 KV 大小无关。",
        PALE,
        TEAL,
    )
    card(
        slide,
        0.42,
        4.12,
        6.16,
        2.00,
        "DPU 上值得放的",
        "段描述和 QP 的缓存，主机不再做冷握手。\n"
        "ZMQ 与 Master RPC 由 DPU 代收，再向同一张网卡 post WR。\n"
        "SSD 暂存放 DPU：本地 NVMe 读进 DPU 内存，RDMA 从这里发。",
        SOFT,
        NAVY,
    )
    card(
        slide,
        6.76,
        4.12,
        6.16,
        2.00,
        "地址到达之前，DPU 也要等",
        "目的地址还没到，就不能 post_send。DPU 只能缓存，或与上一笔重叠。\n"
        "KV 已在 GPU HBM。先搬进 DPU 再上网，会多一次拷贝。\n"
        "PutStart 的分配和租约仍由 Master 裁决，DPU 只做代发。",
        GOLD_BG,
        GOLD,
    )
    slide.notes_slide.notes_text_frame.text = (
        "交叉点按十进制计算：0.2 ms 乘 100 GB/s 等于 2e7 字节，约 20 MB。这是判断用的数量级，不是 Mooncake 的实测控制面时延。"
        "vLLM-Omni 文档记载，H800 加 mlx5_0、约 186 MB KV，MooncakeTransferEngineConnector 的墙钟约 14 ms，"
        "MooncakeStoreConnector 约 810 ms。14 ms 含控制面。若地址 RPC 在亚毫秒，它在这 186 MB 里只占个位数百分比；"
        "若 Master 排到接近 10 ms，就和传输同量级。"
        "数据面已经是 RNIC 对已注册内存的 DMA。DPU 与这张网卡是同一块芯片时，ARM 核可以持有描述符缓存并代发 post_send，DMA 仍走原网卡。"
        "DPU 作为独立一跳放在 GPU 与 RNIC 之间时，HBM 上的 KV 会多一次搬运。"
        "LOCAL_DISK 路径今天由持有节点把 SSD 读入 ClientBuffer，再被对端 RDMA READ。这块暂存和盘读适合放到带本地 NVMe 的 DPU。"
        "Mooncake 当前代码没有把这条控制面卸到 DPU。上面说的是可落地的切分，不是已有功能。"
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Decode KV 控制面载荷、等待与 DPU"
    slide_info(prs)
    slide_judge(prs)
    out = "/workspace/docs/kv-ctrl-dpu.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
