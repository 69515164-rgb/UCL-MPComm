#!/usr/bin/env python3
"""Decode KV read/write interaction, including the control messages."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"


def rgb(hex_color):
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def chrome(slide, title, page):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.14, 10.4, 0.24, "Decode KV    ·    控制消息与单边 RDMA", 12, True, ACCENT)
    textbox(slide, 0.42, 0.38, 11.4, 0.40, title, 24, True, NAVY)
    textbox(slide, 11.2, 0.16, 1.7, 0.24, page, 12, False, MUTED, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide,
        0.42,
        7.20,
        12.5,
        0.24,
        "注册只做一次：ram 元数据含 addr、length、rkey；rpc_meta 含 ip 和 rpc_port；首次 QP 交换 qp_num、lid、gid",
        11,
        False,
        MUTED,
    )


def banner(slide, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(0.86), Inches(12.5), Inches(0.38))
    solid(shape, PALE)
    textbox(slide, 0.56, 0.86, 12.22, 0.38, text, 14, True, TEAL)


def column_head(slide, x, y, w, text, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    solid(shape, fill)
    textbox(slide, x + 0.14, y, w - 0.28, 0.34, text, 14, True, WHITE)


def step(slide, x, y, w, h, kind, headline, body):
    accent = NAVY if kind == "控制" else TEAL
    fill = SOFT if kind == "控制" else PALE
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(card, fill, LINE)
    mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.10), Inches(0.07), Inches(h - 0.20))
    solid(mark, accent)
    textbox(slide, x + 0.18, y + 0.06, 0.62, 0.26, kind, 12, True, accent)
    textbox(slide, x + 0.82, y + 0.05, w - 0.96, 0.28, headline, 13, True, INK)
    textbox(slide, x + 0.18, y + 0.36, w - 0.32, h - 0.42, body, 12, False, INK, anchor=MSO_ANCHOR.TOP)


def gold(slide, y, h, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(y), Inches(12.5), Inches(h))
    solid(shape, GOLD_BG)
    textbox(slide, 0.56, y, 12.22, h, text, 13, False, GOLD)


def slide_read(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "Decode 读前缀时，控制消息只交换地址", "1  /  2")
    banner(slide, "PD 直连由 Prefill 发起 WRITE。KV 池由 Decode 发起 READ。字节进入本端 HBM 后才开始 attention。")
    col_w = 6.16
    left = 0.42
    right = 0.42 + col_w + 0.18
    column_head(slide, left, 1.36, col_w, "PD 直连    Prefill 把 KV 写入 Decode", NAVY)
    column_head(slide, right, 1.36, col_w, "KV 池    Decode 自己把 KV 读回", TEAL)
    steps_l = [
        ("控制", "调度  →  Decode", "带上 do_remote_prefill、remote_host、remote_port，并分配本地 block。"),
        ("控制", "Decode  →  Prefill，ZMQ", "送出 hostname、rpc_port、各层 kv base、block id。这是控制消息，不是 KV 本体。"),
        ("数据", "Prefill  →  Decode HBM", "RDMA WRITE。源是 Prefill 的 block，目的地址是 Decode base 加 block 乘 block_len。"),
        ("控制", "Prefill  →  Decode，ZMQ", "返回 TRANS_DONE。Decode 收到后才 forward。失败则是 TRANS_ERROR。"),
    ]
    steps_r = [
        ("控制", "Decode  →  Master", "GetReplicaList(block key)。Master 给出可读副本，并发放读租约。"),
        ("控制", "Master  →  Decode", "MEMORY 返回 endpoint 和 buffer_address。LOCAL_DISK 返回盘所在节点的 rpc。"),
        ("数据", "Decode  →  远端内存", "内存副本直接 RDMA READ。盘副本先 RPC，对端把对象读进暂存，再 READ 这块暂存。"),
        ("控制", "本端 CQ  →  Decode", "全部 slice 为 COMPLETED 后，前缀 KV 已经在本地 HBM，Decode 开始 attention。"),
    ]
    y = 1.78
    h = 1.16
    gap = 0.08
    for i, (kind, head, body) in enumerate(steps_l):
        step(slide, left, y + i * (h + gap), col_w, h, kind, head, body)
    for i, (kind, head, body) in enumerate(steps_r):
        step(slide, right, y + i * (h + gap), col_w, h, kind, head, body)
    slide.notes_slide.notes_text_frame.text = (
        "这一页说明 Decode 读取已有前缀 KV 的两条控制路径，都不是时序轴。"
        "PD 直连来自 vLLM MooncakeConnector。调度在 kv_transfer_params 里放 do_remote_prefill、remote_host、remote_port。"
        "Decode worker 的 receive_kv 经 ZMQ 把 MooncakeAgentMetadata 发给 Prefill：remote_hostname、remote_port、"
        "kv_caches_base_addr、block_ids、request_ids。Prefill 的 send_kv_to_decode 计算 "
        "src = 本端 layer base + local block * block_len，dst = Decode layer base + remote block * block_len，"
        "调用 batch_transfer_sync_write。数据面是 Prefill 发起的 RDMA WRITE，Decode 网卡把字节 DMA 进已注册 HBM。"
        "ZMQ 回 TRANS_DONE 后 Decode 才把该请求记入 finished_recving。save_kv_layer 不参与这次读取。"
        "KV 池路径来自 Mooncake Store。Decode 对 Master 发 GetReplicaList(key)，响应含 replicas、lease_ttl_ms。"
        "MEMORY 副本用 TransferEngine RDMA READ buffer_address。LOCAL_DISK 副本先对持有节点发 BatchGetOffloadObject，"
        "对端从 SSD 读入 ClientBuffer，响应含 batch_id、pointers、transfer_engine_addr，Decode 再 RDMA READ 这块暂存，"
        "然后 release_offload_buffer。SSD 不是 RDMA 的目标地址。"
        "两条路径的 TransferRequest 都只有 opcode、source、target_id、target_offset、length。位置在控制面已经变成地址。"
    )


def slide_write(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "Decode 新 KV 先留在本端，入池才写出去", "2  /  2")
    banner(slide, "本步 attention 的 K、V 写进本地 HBM。MooncakeConnector 的 save_kv_layer 不向外发送。")
    col_w = 6.16
    left = 0.42
    right = 0.42 + col_w + 0.18
    column_head(slide, left, 1.36, col_w, "默认    只留在本端 HBM", NAVY)
    column_head(slide, right, 1.36, col_w, "入池    控制面先分配，再 WRITE", TEAL)
    local = [
        ("数据", "Attention  →  本端 HBM", "新 token 的 K 和 V 写入下一个本地 block。这一步没有 TransferRequest。"),
        ("控制", "调度  →  下一步 Decode", "后续 token 直接读这块本地 KV。不查 Master，也不打开远端 segment。"),
        ("控制", "请求结束  →  仍在本地", "Connector 不把 Decode 新 KV 推回 Prefill。要复用才走右侧入池。"),
    ]
    pool = [
        ("控制", "Decode  →  Master", "PutStart(key, 字节数)。对象先记为 PROCESSING，读者还看不见。"),
        ("控制", "Master  →  Decode", "返回 segment 和 offset。本机 segment 是 Host DDR，远端 segment 是 remote DDR。"),
        ("数据", "Decode  →  这段内存", "本机做拷贝。远端 post RDMA WRITE。盘不在这次请求里。"),
        ("控制", "Decode  →  Master", "PutEnd。副本变为 COMPLETE，其他 Decode 的 GetReplicaList 才能读到。"),
        ("控制", "Master  →  持有副本的客户端", "PutEnd 之后或淘汰时，客户端把对象写入本地 SSD，副本改为 LOCAL_DISK。"),
    ]
    y = 1.80
    for i, (kind, head, body) in enumerate(local):
        step(slide, left, y + i * 1.10, col_w, 1.02, kind, head, body)
    note = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(5.12), Inches(col_w), Inches(1.12)
    )
    solid(note, SOFT, LINE)
    textbox(
        slide,
        left + 0.16,
        5.20,
        col_w - 0.32,
        0.96,
        "这块 HBM 在启动时已经注册。下一步 Decode 按本地 block 地址读取，不打开远端 segment，也不产生完成队列。",
        13,
        False,
        INK,
        anchor=MSO_ANCHOR.TOP,
    )
    for i, (kind, head, body) in enumerate(pool):
        step(slide, right, y + i * 0.90, col_w, 0.84, kind, head, body)
    gold(
        slide,
        6.42,
        0.66,
        "写到哪里在控制面定完。TransferRequest 只剩 opcode、本地地址、segment、offset、length。\n"
        "远端 DDR 是访存语义。SSD 要等卸载控制把对象装进暂存，才再次变成 RDMA READ。",
    )
    slide.notes_slide.notes_text_frame.text = (
        "Decode 写 KV 分默认路径和入池路径。"
        "vLLM MooncakeConnector.save_kv_layer 直接返回，不提交 TransferRequest。"
        "新 token 的 KV 留在 Decode 进程已经 registerLocalMemory 的 HBM 里，下一步 attention 本地读取。"
        "要放进 Mooncake Store 给别的请求复用时，走两阶段元数据。"
        "PutStart 在 Master 上占配额、按 AllocationStrategy 从某个 segment 分配缓冲，副本状态为 PROCESSING，"
        "把 endpoint 和 buffer_address 还给 Decode。segment 属于本机时，写入就是到 Host DDR 的本地拷贝；"
        "segment 属于别的节点时，Decode 作为发起方 post IBV_WR_RDMA_WRITE，对端网卡按 rkey 做 DMA，对端 CPU 不读请求。"
        "PutEnd 校验写者后把副本标为 COMPLETE。失败则 PutRevoke，释放缓冲。"
        "SSD 不是 TransferRequest 的一种 opcode。enable_offload 时，COMPLETE 之后由 Master 控制面排队，"
        "客户端把对象写入 MOONCAKE_OFFLOAD_FILE_STORAGE_PATH，副本类型变为 LOCAL_DISK。"
        "offload_on_evict 则拖到淘汰时才写盘。读这条副本的控制消息是 BatchGetOffloadObject，数据面仍是对暂存缓冲的 RDMA READ。"
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Decode 读写 KV 的控制与数据交互"
    slide_read(prs)
    slide_write(prs)
    out = "/workspace/docs/kv-rdma-rw-flow.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
