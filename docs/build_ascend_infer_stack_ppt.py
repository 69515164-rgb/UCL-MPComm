#!/usr/bin/env python3
"""Editable layer diagram of the Ascend inference software stack."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)
TEAL = (14, 122, 118)
GOLD = (122, 84, 28)
LINE = (214, 220, 226)
SOFT = (244, 247, 250)
PILL = (236, 240, 244)
HDK = (46, 74, 110)
GOLD_BG = (255, 248, 240)


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    strip_shadow(shape)


def write(shape, text, size, bold, color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def box(slide, x, y, w, h, fill, text="", size=12, bold=True, color=WHITE, line=None, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    if text:
        write(shape, text, size, bold, color, align)
    return shape


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def band(slide, x, y, w, h, fill, title, body, title_w=1.70):
    box(slide, x, y, title_w, h, fill, title, 13, True, WHITE)
    box(slide, x + title_w, y, w - title_w, h, fill, "")
    textbox(
        slide,
        x + title_w + 0.12,
        y,
        w - title_w - 0.20,
        h,
        body,
        13,
        False,
        WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def path_rows(slide, x, y, w, h, head_fill, head, rows, control_index=None):
    box(slide, x, y, w, h, WHITE, line=LINE)
    box(slide, x, y, w, 0.34, head_fill, head, 14, True, WHITE)
    top = y + 0.42
    gap = 0.06
    row_h = (h - 0.50 - gap * (len(rows) - 1)) / len(rows)
    name_w = 1.78
    for i, (name, desc) in enumerate(rows):
        ry = top + i * (row_h + gap)
        fill = GOLD if i == control_index else head_fill
        body_fill = GOLD_BG if i == control_index else SOFT
        body_ink = INK
        box(slide, x + 0.08, ry, name_w, row_h, fill, name, 12, True, WHITE)
        box(slide, x + 0.08 + name_w, ry, w - name_w - 0.16, row_h, body_fill, "")
        textbox(
            slide,
            x + name_w + 0.18,
            ry,
            w - name_w - 0.30,
            row_h,
            desc,
            12,
            False,
            body_ink,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def cells(slide, x, y, w, h, items, gap=0.08):
    cell_w = (w - gap * (len(items) - 1)) / len(items)
    for i, (fill, title, body) in enumerate(items):
        cx = x + i * (cell_w + gap)
        box(slide, cx, y, cell_w, h, fill, "")
        textbox(slide, cx + 0.08, y + 0.04, cell_w - 0.16, 0.24, title, 13, True, WHITE)
        textbox(slide, cx + 0.08, y + 0.28, cell_w - 0.16, h - 0.32, body, 11, False, WHITE, anchor=MSO_ANCHOR.TOP)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "昇腾推理软件栈分层"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box(slide, 0, 0, 13.333, 0.08, ACCENT, "")
    textbox(slide, 0.32, 0.16, 12.6, 0.22, "昇腾  ·  推理软件栈  ·  MindIE 3.1", 13, True, ACCENT)
    textbox(slide, 0.32, 0.40, 12.7, 0.40, "两条推理路径在 CANN 汇合，下面共用 HDK 和 NPU", 24, True, NAVY)
    textbox(
        slide,
        0.32,
        0.84,
        12.7,
        0.32,
        "新部署走 MindIE Motor，南向当前是 vLLM-Ascend。MindIE LLM 仍保留原生层次，已进入维护。",
        15,
        False,
        INK,
    )

    label_x, label_w = 0.28, 0.92
    bar_x, bar_w = 1.24, 11.76

    textbox(slide, label_x, 1.22, label_w, 0.36, "集群", 13, True, GOLD, anchor=MSO_ANCHOR.MIDDLE)
    band(
        slide,
        bar_x,
        1.22,
        bar_w,
        0.36,
        GOLD,
        "MindCluster",
        "调用链外侧。Pod、PD 分离、故障恢复。可选 CCAE。",
        title_w=1.70,
    )

    textbox(slide, label_x, 1.66, label_w, 2.48, "服务", 13, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)
    path_w = (bar_w - 0.12) / 2
    path_rows(
        slide,
        bar_x,
        1.66,
        path_w,
        2.48,
        TEAL,
        "路径 A  ·  MindIE LLM 原生引擎",
        (
            ("Server", "Endpoint 兼容 OpenAI、vLLM、Triton、TGI。"),
            ("LLM Manager", "连续组批。Engine 内含 Scheduler。"),
            ("Text Generator", "自回归执行。Executor 负责跨节点下发。"),
            ("Modeling", "ACLGraph 与 ATBGraph。"),
        ),
    )
    path_rows(
        slide,
        bar_x + path_w + 0.12,
        1.66,
        path_w,
        2.48,
        ACCENT,
        "路径 B  ·  MindIE Motor",
        (
            ("Coordinator", "路由、限流、实例隔离。"),
            ("管控面", "Controller 与 NodeManager。"),
            ("EngineServer", "北向接 Coordinator 和 Controller。"),
            ("vLLM-Ascend", "当前唯一南向。前向和实例内组批。"),
        ),
        control_index=1,
    )

    textbox(slide, label_x, 4.22, label_w, 0.56, "KV 池", 12, True, GOLD, anchor=MSO_ANCHOR.MIDDLE)
    cells(
        slide,
        bar_x,
        4.22,
        bar_w,
        0.56,
        (
            (GOLD, "MemCache", "默认后端。元数据与 DRAM 池。"),
            (GOLD, "Mooncake", "mooncake_master。Conductor 管 KV 位置。"),
            (TEAL, "实例内", "Block Manager 与 vLLM 本卡池。另有 YuanRong。"),
        ),
    )

    textbox(slide, label_x, 4.86, label_w, 0.64, "通信", 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    cells(
        slide,
        bar_x,
        4.86,
        bar_w,
        0.64,
        (
            (ACCENT, "HCCL", "集合通信。服务张量并行和专家并行。"),
            (HDK, "HCOMM", "HCCL 的基础库。控制面与数据面。"),
            (TEAL, "HIXL", "单边。LLM-DataDist 拉取和推送 KV。"),
            (NAVY, "SHMEM", "对称内存。卡间单边访问，给定制算子。"),
        ),
    )

    shared = (
        (ACCENT, "运行时", "CANN", "AscendCL、GE、Runtime、算子库、DVPP。通信四库也在这一层。推理包是 NNRT。", 0.40),
        (HDK, "驱动", "HDK / Driver", "驱动与固件，与 CANN 分版本配套。", 0.30),
        (NAVY, "硬件", "昇腾 NPU", "推理任务在昇腾处理器上执行。", 0.30),
    )
    y = 5.60
    for fill, label, title, body, h in shared:
        textbox(slide, label_x, y, label_w, h, label, 13, True, fill, anchor=MSO_ANCHOR.MIDDLE)
        band(slide, bar_x, y, bar_w, h, fill, title, body, title_w=1.70)
        y += h + 0.06

    textbox(
        slide,
        0.32,
        6.78,
        12.7,
        0.58,
        "HCCL 走集合，服务模型并行。HIXL 走单边，服务 P 到 D 的 KV。按层池化目前只在 MemCache 上。\n"
        "HCOMM 是文档里的通信基础库名。MemCache、Mooncake、YuanRong 是池后端，不是三条调用链。",
        13,
        False,
        MUTED,
        anchor=MSO_ANCHOR.TOP,
    )

    slide.notes_slide.notes_text_frame.text = (
        "依据 MindIE 3.1 用户手册 Overview 与 MindIE Motor 3.1 架构说明。"
        "MindIE LLM 进入维护，官方建议新部署使用 MindIE Motor 加 vLLM-Ascend。"
        "MindIE LLM 自上而下为 Server、LLM Manager、Text Generator、Modeling。"
        "LLM Manager 内的 Engine 编排 Scheduler、Executor 和 Worker；Block Manager 管 KV 池。"
        "Executor 把调度结果下发给 Text Generator，覆盖跨节点、跨设备。"
        "Text Generator 含前处理、Generator、Sampler，并支持可插拔的并行解码。"
        "Modeling 提供 ATB Models；开发者架构说明同时支持 ACLGraph 和 ATBGraph。"
        "Motor：Coordinator 为数据面入口；Controller 与 NodeManager 为管控面；EngineServer 北向接二者，南向当前仅 vLLM。"
        "架构页写明南向接口包括 vLLM、SGLang 和 MindIE，当前版本只支持 vLLM。"
        "MindCluster 提供 Kubernetes、PD 分离 CRD 和 Operator，位于调用链外侧。CCAE 可选。"
        "CANN 组成沿用方案概述：AscendCL、GE、Runtime、算子库、HCCL、DVPP。推理运行包为 NNRT。"
        "KV 池：Motor 3.1 默认后端是 MemCache，由 MetaService 管元数据、LocalService 管 DRAM；Mooncake 用 mooncake_master；还有 YuanRong。"
        "vLLM-Ascend 的按层池化目前只支持 MemCache。"
        "Mooncake Conductor 给 Coordinator 提供 KV 位置，用于亲和调度。"
        "通信库都在 CANN 内。HCCL 是集合通信库，服务张量并行和专家并行；HCOMM 是它的基础库，控制面管拓扑和资源，数据面管搬运和同步。"
        "HIXL 是单边通信库，LLM-DataDist 在这里做链路和 KV 的拉取、推送，支持 D2D、D2H、H2D。"
        "SHMEM 是对称内存库，给算子开发提供卡间单边访问。"
        "用户所说的 HCOM 对应文档名 HCOMM。"
    )

    out = "/workspace/docs/ascend-infer-stack.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
