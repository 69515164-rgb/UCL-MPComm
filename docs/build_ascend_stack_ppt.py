#!/usr/bin/env python3
"""Editable layer diagrams of the Ascend training and inference software stacks."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)
TEAL = (14, 122, 118)
GOLD = (122, 84, 28)
GREEN = (27, 110, 72)
LINE = (214, 220, 226)
SOFT = (244, 247, 250)
PILL = (236, 240, 244)
HDK = (46, 74, 110)
CORE = (14, 90, 110)


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    strip_shadow(shape)


def write(shape, text, size, bold, color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def box(slide, x, y, w, h, fill, text="", size=12, bold=True, color=WHITE, line=None, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill, line)
    if text:
        write(shape, text, size, bold, color, align, anchor)
    return shape


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def header(slide, kicker, title, lead):
    box(slide, 0, 0, 13.333, 0.08, ACCENT, "")
    textbox(slide, 0.32, 0.14, 12.6, 0.22, kicker, 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.32, 0.36, 12.7, 0.34, title, 22, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    box(slide, 0.32, 0.78, 0.07, 0.34, ACCENT, "")
    textbox(slide, 0.48, 0.74, 12.5, 0.40, lead, 14, True, INK, anchor=MSO_ANCHOR.MIDDLE)


def layer_bar(slide, x, y, w, h, fill, title, body):
    box(slide, x, y, w, h, fill, "")
    textbox(slide, x + 0.14, y + 0.06, w - 0.28, 0.28, title, 15, True, WHITE)
    textbox(slide, x + 0.14, y + 0.34, w - 0.28, h - 0.40, body, 12, False, WHITE, anchor=MSO_ANCHOR.TOP)


def side_card(slide, x, y, w, h, title, rows):
    box(slide, x, y, w, h, WHITE, line=LINE)
    box(slide, x, y, w, 0.40, TEAL, title, 14, True, WHITE)
    top = y + 0.50
    gap = 0.06
    row_h = (h - 0.60 - gap * (len(rows) - 1)) / len(rows)
    for i, (name, desc) in enumerate(rows):
        ry = top + i * (row_h + gap)
        box(slide, x + 0.10, ry, w - 0.20, row_h, PILL, "")
        textbox(slide, x + 0.18, ry + 0.02, w - 0.36, 0.22, name, 12, True, NAVY)
        textbox(slide, x + 0.18, ry + 0.22, w - 0.36, row_h - 0.26, desc, 11, False, INK, anchor=MSO_ANCHOR.TOP)


def slide_train(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "昇腾软件栈  ·  训练    1 / 2",
        "训练从套件进入框架，再进入 CANN",
        "单步计算走这条调用链。MindCluster 决定任务落在哪些 NPU，不参与单步算子执行。",
    )

    layers = (
        ("任务套件", GOLD, "MindSpeed LLM / MM / RL", "大语言、多模态、强化学习的训练流程。"),
        ("分布式加速", CORE, "MindSpeed Core", "多维并行、内存优化、通算掩盖，对接开源分布式框架。"),
        ("训练框架", TEAL, "MindSpore，以及 PyTorch + FrameworkPTAdapter", "torch_npu 把 PyTorch 算子落到昇腾。MindSpeed 也提供 MindSpore 后端。"),
        ("计算架构", ACCENT, "CANN", "AscendCL、GE、Runtime、算子库、HCCL、DVPP。训练用 Toolkit 或 NNAE。"),
        ("驱动固件", HDK, "HDK / Driver", "驱动与固件。CANN 经 Driver 管理 Device、Context 和 Stream。"),
        ("处理器", NAVY, "昇腾 NPU", "Atlas 训练服务器上的昇腾处理器。"),
    )
    x_label, label_w = 0.28, 1.18
    x_bar, bar_w = 1.52, 7.28
    y0, bar_h, gap = 1.26, 0.78, 0.07
    for i, (label, fill, title, body) in enumerate(layers):
        y = y0 + i * (bar_h + gap)
        textbox(slide, x_label, y, label_w, bar_h, label, 12, True, fill, anchor=MSO_ANCHOR.MIDDLE)
        layer_bar(slide, x_bar, y, bar_w, bar_h, fill, title, body)

    side_card(
        slide,
        8.98,
        1.26,
        4.06,
        5.03,
        "调用链外侧  ·  MindCluster",
        (
            ("Volcano", "按芯片互联做亲和调度，而不只按卡数分配。"),
            ("Ascend Operator", "创建 AscendJob，注入主进程地址和 RankTable。"),
            ("ClusterD", "汇总任务、资源和故障，统一决定处理级别。"),
            ("Device Plugin / NodeD", "向 Kubernetes 注册 NPU，并上报节点状态。"),
            ("断点续训 / FaultDiag", "故障后恢复训练；日志和链路诊断单独成工具。"),
        ),
    )

    textbox(
        slide,
        0.32,
        6.72,
        12.7,
        0.58,
        "自上而下是一次训练步的软件层次。MindCluster 的组件跑在管理节点和节点代理上，负责任务能否跑起来。\n"
        "CANN 与 HDK 分版本配套。NNRT 只覆盖推理，不作为训练运行包。",
        12,
        False,
        MUTED,
        anchor=MSO_ANCHOR.TOP,
    )
    slide.notes_slide.notes_text_frame.text = (
        "训练层次依据昇腾社区 MindSpeed 文档与华为 2025-03 技术文章：MindSpeed 由 Core、LLM、MM、RL 组成。"
        "框架为 MindSpore，以及 PyTorch 经 FrameworkPTAdapter（torch_npu）。"
        "CANN 组成来自华为 CANN 方案概述：AscendCL、GE、Runtime、DVPP、算子库、HCCL、Driver。"
        "软件包 Toolkit、NNAE 覆盖训练，NNRT 仅推理。"
        "MindCluster 提供 Volcano、Ascend Operator、ClusterD、Ascend Device Plugin、NodeD、断点续训和 FaultDiag，位于任务调度外侧。"
    )


def path_card(slide, x, y, w, h, head_fill, head, rows):
    box(slide, x, y, w, h, WHITE, line=LINE)
    box(slide, x, y, w, 0.36, head_fill, head, 14, True, WHITE)
    top = y + 0.46
    gap = 0.06
    row_h = (h - 0.56 - gap * (len(rows) - 1)) / len(rows)
    for i, (name, desc) in enumerate(rows):
        ry = top + i * (row_h + gap)
        box(slide, x + 0.10, ry, 1.70, row_h, head_fill, name, 12, True, WHITE)
        box(slide, x + 1.86, ry, w - 2.06, row_h, SOFT, "")
        textbox(slide, x + 1.96, ry, w - 2.20, row_h, desc, 12, False, INK, anchor=MSO_ANCHOR.MIDDLE)


def slide_infer(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(
        slide,
        "昇腾软件栈  ·  推理    2 / 2",
        "推理在同一套 CANN 上分成两条路径",
        "MindIE LLM 自带服务和实例内调度。MindIE Motor 做跨实例的 PD 调度，当前南向对接 vLLM-Ascend。",
    )

    path_card(
        slide,
        0.28,
        1.26,
        6.32,
        3.28,
        TEAL,
        "路径 A  ·  MindIE LLM 原生引擎",
        (
            ("Server", "服务层。OpenAI、vLLM、Triton 协议由 Endpoint 封装。"),
            ("LLM Manager", "实例内调度。连续组批、KV 池、请求状态。"),
            ("Text Generator", "前处理、自回归执行、后处理。含投机解码和分块 prefill。"),
            ("Modeling", "算子编排与图执行。后端为 ACLGraph 和 ATBGraph。"),
        ),
    )
    path_card(
        slide,
        6.74,
        1.26,
        6.32,
        3.28,
        ACCENT,
        "路径 B  ·  Motor 调度 + vLLM-Ascend",
        (
            ("Coordinator", "请求入口。路由、负载均衡、请求统计、故障实例隔离。"),
            ("Controller", "P/D 身份、故障隔离与恢复，并把状态推给 Coordinator。"),
            ("NodeManager", "节点代理。向 Controller 注册，拉起引擎并上报心跳。"),
            ("vLLM-Ascend", "Motor 当前南向的推理引擎。EngineServer 这一版以 vLLM 为准。"),
        ),
    )

    shared = (
        (TEAL, "MindCluster", "两条路径的 Pod、PD 分离 CRD 和故障恢复都由它承接。可选 CCAE 做算存网运维可视化。"),
        (ACCENT, "CANN", "AscendCL、GE、Runtime、算子库、HCCL、DVPP。推理可用 NNRT；Toolkit 与 NNAE 同时覆盖训练和推理。"),
        (HDK, "HDK / Driver", "驱动与固件，与 CANN 分版本配套。"),
        (NAVY, "昇腾 NPU", "推理任务运行在昇腾处理器上。MindCluster 同时支持训练硬件和推理硬件。"),
    )
    y = 4.66
    heights = (0.48, 0.58, 0.42, 0.42)
    for (fill, title, body), h in zip(shared, heights):
        box(slide, 0.28, y, 1.55, h, fill, title, 12, True, WHITE)
        box(slide, 1.89, y, 11.16, h, fill, body, 13, False, WHITE, align=PP_ALIGN.LEFT)
        y += h + 0.06

    textbox(
        slide,
        0.32,
        6.78,
        12.7,
        0.52,
        "两条路径在 CANN 汇合，互不为对方的下层。Motor 负责请求落到哪一个 Prefill 或 Decode 实例；实例内部的组批在对应引擎里完成。",
        12,
        False,
        MUTED,
        anchor=MSO_ANCHOR.TOP,
    )
    slide.notes_slide.notes_text_frame.text = (
        "推理层次依据 MindIE LLM 架构说明：Server、LLM Manager、Text Generator、Modeling；"
        "Modeling 支持 ACLGraph 和 ATBGraph。加速特性包括 Continuous Batching、PagedAttention、FlashDecoding、SpecDecoding、ChunkPrefill。"
        "MindIE Motor 3.1 架构：Coordinator、Controller、NodeManager、EngineServer、Deployer；向下对接 vLLM-Ascend。"
        "架构说明写明当前 EngineServer 南向仅支持 vLLM。PD 分离设计中 SGLang 会上报 connector 能力。"
        "MindCluster 为 Motor 提供 Kubernetes、PD CRD 和 Operator，训练与推理共用。"
        "CANN 软件包 NNRT 仅推理，Toolkit 与 NNAE 覆盖训练和推理。"
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "昇腾训练与推理软件栈层次"
    slide_train(prs)
    slide_infer(prs)
    out = "/workspace/docs/ascend-train-infer-stack.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
