#!/usr/bin/env python3
"""Editable flowchart: multi-request scheduling and where each SLO challenge sits."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = (14, 42, 71)
INK = (28, 36, 46)
MUTED = (90, 99, 110)
WHITE = (255, 255, 255)
ACCENT = (20, 90, 150)

# Metric colors: tag, card fill, bar.
METRIC = {
    "TTFT": ((20, 90, 150), "E8F1FA", "145A96"),
    "TPOT": ((148, 96, 28), "FBF3EA", "945C14"),
    "两者": ((92, 61, 143), "F3EEF8", "5C3D8F"),
    "产能": ((27, 110, 72), "E8F5EE", "1B6E48"),
}


def rgb(c):
    return RGBColor(*c)


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    sppr = shape._element.spPr
    effect = sppr.find(qn("a:effectLst"))
    if effect is not None:
        sppr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    strip_shadow(shape)


def write_lines(shape, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.06):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.03)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set(
                "anchor",
                {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
            )
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = text
        set_run_font(run, size, bold, color)


def add_text(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        bodyPr = body.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set(
                "anchor",
                {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
            )
    parts = text.split("\n")
    for i, line in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return box


def rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, fill)
    return shape


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "多请求调度：挑战落在哪一步"
    prs.core_properties.subject = "TTFT 与 TPOT 在共享队列、KV、组批和网络上的额外挑战"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rect(slide, 0, 0, 13.333, 0.08, ACCENT)
    add_text(slide, 0.32, 0.14, 12.6, 0.22, "推理调度  ·  相对单请求多出来的难度", 12, True, ACCENT)
    add_text(slide, 0.32, 0.36, 12.6, 0.34, "多请求调度：挑战落在哪一步", 24, True, NAVY)
    rect(slide, 0.32, 0.78, 0.07, 0.40, ACCENT)
    add_text(
        slide,
        0.48,
        0.74,
        12.5,
        0.46,
        "单请求只有一次前向。多请求共用队列、显存、组批和网络，尾部在下面六步被抬过 SLO。色卡是这一步多出来的挑战。",
        14,
        True,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    steps = [
        (
            "1  到达",
            "进入队列",
            [
                ("TTFT", "排队、队头阻塞", "长请求挡住短请求"),
            ],
        ),
        (
            "2  准入",
            "决定接不接",
            [
                ("TPOT", "再接一条", "在飞 TPOT 被推过线"),
                ("产能", "输出长度未知", "花完算力才知超标"),
                ("产能", "目标若是总吞吐", "会接进违约的请求"),
            ],
        ),
        (
            "3  路由",
            "选择副本",
            [
                ("两者", "前缀命中换热点", "TTFT 降，副本 TPOT 升"),
            ],
        ),
        (
            "4  KV 安置",
            "分配显存",
            [
                ("两者", "容量、碎片、换出", "别人的下一步出尖峰"),
            ],
        ),
        (
            "5  组批",
            "合成一次执行",
            [
                ("TPOT", "prefill 插入 decode", "在飞的这一步被拉长"),
                ("TTFT", "分块 prefill", "首 token 被拆成多步"),
            ],
        ),
        (
            "6  通信与结算",
            "搬运、集合、计产能",
            [
                ("TPOT", "多条 KV 同时汇入", "decode 侧形成 incast"),
                ("TPOT", "通信与 KV 抢链路", "步尾部变成 TPOT"),
                ("TPOT", "EP 等最慢的 rank", "整组步长被拉长"),
                ("产能", "两指标同时达标", "超标 token 仍占 GPU"),
            ],
        ),
    ]

    pitch = 2.16
    origin = 0.30
    header_w = 1.90
    header_h = 0.62
    header_y = 1.32
    card_w = 2.06
    card_top = 2.06
    card_h = 1.16
    card_gap = 0.06

    for i, (title, subtitle, cards) in enumerate(steps):
        hx = origin + i * pitch
        header = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(hx), Inches(header_y), Inches(header_w), Inches(header_h)
        )
        solid(header, NAVY)
        try:
            header.adjustments[0] = 0.15
        except Exception:
            pass
        write_lines(
            header,
            [
                (title, 14, True, WHITE),
                (subtitle, 11, False, (214, 224, 234)),
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.04,
        )
        if i < len(steps) - 1:
            ax = hx + header_w + 0.02
            ay = header_y + header_h / 2 - 0.08
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(ay), Inches(0.22), Inches(0.16)
            )
            solid(arrow, ACCENT)

        stem_x = hx + header_w / 2 - 0.01
        stem = rect(slide, stem_x, header_y + header_h, 0.02, 0.12, ACCENT)

        for j, (metric, head, body) in enumerate(cards):
            fg, _fill_hex, _bar = METRIC[metric]
            cy = card_top + j * (card_h + card_gap)
            cx = hx - 0.08
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy), Inches(card_w), Inches(card_h)
            )
            # Fill from hex via RGB tuple.
            fill_rgb = tuple(int(_fill_hex[k : k + 2], 16) for k in (0, 2, 4))
            solid(card, fill_rgb, (220, 226, 232))
            try:
                card.adjustments[0] = 0.08
            except Exception:
                pass
            bar = rect(slide, cx, cy + 0.08, 0.06, card_h - 0.16, fg)
            write_lines(
                card,
                [
                    (metric, 11, True, fg),
                    (head, 12, True, INK),
                    (body, 11, False, MUTED),
                ],
                align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.12,
            )
        del stem

    add_text(
        slide,
        0.30,
        7.12,
        9.3,
        0.26,
        "多请求 TTFT = 排队 + 调度 + prefill    多请求 TPOT = 一步 decode + 被插入的停顿 + 通信排队",
        11,
        False,
        MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    legend = [("TTFT", METRIC["TTFT"][0]), ("TPOT", METRIC["TPOT"][0]), ("两者", METRIC["两者"][0]), ("产能", METRIC["产能"][0])]
    lx = 9.55
    for name, color in legend:
        rect(slide, lx, 7.18, 0.16, 0.14, color)
        add_text(slide, lx + 0.20, 7.12, 0.70, 0.26, name, 11, True, INK, anchor=MSO_ANCHOR.MIDDLE)
        lx += 0.92

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "流程按请求被调度的顺序排列：到达、准入、路由、KV 安置、组批、通信与结算。"
        "每张色卡是相对单请求多出来的挑战，颜色表示先打坏的指标。"
        "蓝 TTFT，琥珀 TPOT，紫表示两个指标一起变坏，绿表示达标结算。"
        "连续批处理的干扰来自 Orca；分块 prefill 来自 Sarathi；"
        "把 prefill 与 decode 拆开的原因来自 DistServe 与 Splitwise。"
        "单请求屋顶测的是一次前向。调度要守的是选定并发下的尾部，而且准入会改变这个尾部。"
    )

    out = "/workspace/docs/multi-request-sched-flow.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
