#!/usr/bin/env python3
"""10T GLM on N+2 @ 2TB @ 480 GB/s DDR, in the 6-slide SLO layout.

GLM-5.3 published weights are FP8, so every GEMM uses half of the N+2
FP4 peak: 7.50 P. Memory and communication bandwidths are at 80%.

H2D is 1.6 Tb/s per GPU: 160 GB/s effective, prefill only.
DDR5 is 1.6 TB/s and 1.3 TB per 8-GPU machine.
D2D for expert traffic is 1.2 TB/s per GPU: 0.96 TB/s effective.
The bus-attached DDR pool used for decode KV swap is 480 GB/s per GPU,
so 384 GB/s effective. Its capacity is unlimited. P2P and the parameter
plane are 800 Gb/s: 80 GB/s effective. Intra-machine fullmesh is 8 GPUs,
so attention TP=8.
"""

from __future__ import annotations

import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
NAVY = "0E2A47"
ACCENT = "145A96"
TEAL = "0E7A76"
INK = "1C242E"
MUTED = "5A636E"
WHITE = "FFFFFF"
LINE = "D6DCE2"
SOFT = "F4F7FA"
HEAD = "0E2A47"
PALE = "E7F3F2"
GOLD_BG = "F8F1E4"
GOLD = "7A541C"

H = 6144
D_MOE = 2048
L = 78
L_MOE = 75
L_DENSE = 3
N_INDEX = 21
TOPK = 8
E0 = 256
E = 3526
N, TP, NODE = 128, 8, 8
DP = N // TP
MLA = 165_019_648
IDX_P = 9_371_648
ONE = 3 * H * D_MOE * L_MOE
DENSE = 3 * H * 12288
ATTN = L * MLA + N_INDEX * IDX_P
VOCAB = 154880
REST = L_DENSE * DENSE + ONE + ATTN + 2 * VOCAB * H
BPT = L * 576 + N_INDEX * 132

PEAK = 7.5e15
HBM_CAP = 432e9
HBM_BW = 14.4e12 * 0.80
D2D = 1.2e12 * 0.80
POOL = 480e9 * 0.80
P2P = 800e9 / 8 * 0.80
PARAM = 800e9 / 8 * 0.80
H2D = 1.6e12 / 8 * 0.80
DDR_NODE_BW = 1.6e12 * 0.80
DDR_SHARE = DDR_NODE_BW / NODE
RESERVE = 6e9
ROOM = N * (HBM_CAP - E * ONE / N - ONE - L_DENSE * DENSE - ATTN / TP - RESERVE)


def rgb(hex_color):
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_run_font(run, size, bold, color, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = OxmlElement(f"a:{tag}")
            rpr.append(el)
        el.set("typeface", name)


def strip_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def solid(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    strip_shadow(shape)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = getattr(tf, "_txBody", None)
    if body is not None:
        body_pr = body.find(qn("a:bodyPr"))
        if body_pr is not None:
            body_pr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)
    return shape


def set_cell_fill(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("a:solidFill"):
            tc_pr.remove(child)
    solid_fill = OxmlElement("a:solidFill")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", hex_color)
    solid_fill.append(srgb)
    tc_pr.append(solid_fill)


def set_cell_borders(cell, hex_color="D6DCE2", width="6350"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for child in list(tc_pr):
            if child.tag == qn(f"a:{edge}"):
                tc_pr.remove(child)
        ln = OxmlElement(f"a:{edge}")
        ln.set("w", width)
        fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_color)
        fill.append(srgb)
        ln.append(fill)
        tc_pr.append(ln)


def write_cell(cell, text, size, bold, color, align):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p_pr = p._p.get_or_add_pPr()
        ln_spc = OxmlElement("a:lnSpc")
        spc = OxmlElement("a:spcPct")
        spc.set("val", "100000")
        ln_spc.append(spc)
        p_pr.append(ln_spc)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, bold, color)


def add_table(slide, x, y, widths, rows, heights, kinds, size=12):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights)))
    table = shape.table
    table.horz_banding = False
    table.vert_banding = False
    tbl = shape._element.graphic.graphicData.tbl
    tbl_pr = tbl.find(qn("a:tblPr"))
    if tbl_pr is None:
        tbl_pr = tbl.tblPr
    for child in list(tbl_pr):
        if child.tag == qn("a:tableStyleId"):
            tbl_pr.remove(child)
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    style = {
        "head": (HEAD, WHITE, True, PP_ALIGN.LEFT),
        "body": (WHITE, INK, False, PP_ALIGN.LEFT),
        "key": (PALE, TEAL, True, PP_ALIGN.LEFT),
        "sum": (SOFT, NAVY, True, PP_ALIGN.LEFT),
        "note": (GOLD_BG, INK, False, PP_ALIGN.LEFT),
    }
    for r, (row, h, kind) in enumerate(zip(rows, heights, kinds)):
        table.rows[r].height = Inches(h)
        fill, color, bold, align = style[kind]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            set_cell_fill(cell, fill if c or kind == "head" else SOFT)
            if kind == "head":
                set_cell_fill(cell, HEAD)
            set_cell_borders(cell)
            this_bold = True if c == 0 and kind != "head" else bold
            this_color = NAVY if c == 0 and kind == "body" else color
            if kind in ("key", "sum", "head", "note"):
                this_color = color
                this_bold = bold or c == 0
            write_cell(cell, text, size, this_bold, this_color, align)
    return shape


def chrome(slide, page, title, kicker, total=6):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    solid(bar, ACCENT)
    textbox(slide, 0.42, 0.16, 10.6, 0.28, kicker, 12, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.42, 0.42, 12.5, 0.42, title, 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 11.4, 0.18, 1.5, 0.26, f"{page}  /  {total}", 12, False, MUTED, PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(7.18), Inches(12.5), Inches(0.01))
    solid(rule, LINE)
    textbox(
        slide, 0.42, 7.20, 12.5, 0.24,
        "10T 为外推    ·    N+2 @ 2TB @ 480 GB/s 内存池、一份 128 卡副本    ·    访存与通信取峰值的 80%    ·    权重 FP8，算力 7.50 P",
        11, False, MUTED, anchor=MSO_ANCHOR.MIDDLE,
    )


def banner(slide, y, text, size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(y), Inches(12.5), Inches(0.46))
    solid(shape, PALE)
    textbox(slide, 0.58, y, 12.18, 0.46, text, size, True, TEAL, anchor=MSO_ANCHOR.MIDDLE)


def chips(slide, items, number_size=28):
    for i, (cap, num, sub) in enumerate(items):
        x = 0.42 + i * 3.20
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.40), Inches(3.05), Inches(1.62))
        solid(card, SOFT, LINE)
        textbox(slide, x + 0.16, 5.48, 2.75, 0.26, cap, 12, True, MUTED)
        textbox(slide, x + 0.16, 5.74, 2.75, 0.50, num, number_size, True, TEAL)
        textbox(slide, x + 0.16, 6.26, 2.75, 0.64, sub, 13, False, INK)


def align(m):
    if m <= 0:
        return 1.0
    return m / (math.ceil(m / 16.0) * 16)


def prefill_one(s):
    sel = min(2048, s)
    idx = N_INDEX * 8192 * s * s
    core = L * 65536 * sel * s
    moe = 2 * (TOPK + 1) * ONE * s
    den = 2 * DENSE * L_DENSE * s
    proj = 2 * ATTN * s
    parts = {
        "idx": idx / (N * PEAK),
        "core": core / (N * PEAK),
        "moe": moe / (N * PEAK),
        "den": den / (N * PEAK),
        "proj": proj / (N * PEAK),
    }
    tc = sum(parts.values())
    tw = (E * ONE / N + ONE + L_DENSE * DENSE + ATTN / TP) / HBM_BW
    ship = s * BPT / (N * PARAM)
    parts.update(tc=tc, tw=tw, ship=ship, t=max(tc, tw) + ship)
    return parts


def max_input(limit=10.0):
    lo, hi = 1.0, 2e8
    for _ in range(80):
        mid = (lo + hi) / 2
        if prefill_one(mid)["t"] > limit:
            hi = mid
        else:
            lo = mid
    return lo


def decode(bs, s):
    frac = 1 - math.exp(-TOPK * bs / E)
    e = E / N
    std = math.sqrt(max(e * frac * (1 - frac), 0.0))
    e_hot = min(e, e * frac + 2 * std)
    tok = (bs * TOPK / E) / frac if frac > 1e-9 else 1.0
    ue, ua = align(tok), align(bs / DP)
    sel = min(2048, s)
    score = N_INDEX * 8192 * s + L * 65536 * sel
    proj = 2 * ATTN
    moe_f = 2 * (TOPK + 1) * ONE
    tw = (e_hot * ONE + ONE) / HBM_BW
    store = bs * BPT * s
    read = N_INDEX * 132 * s + L * sel * 576
    read_gpu = bs * read / N
    spill = max(store - ROOM, 0.0) / store if store else 0.0
    t_kh = read_gpu * (1.0 - spill) / HBM_BW if store else 0.0
    t_pool = read_gpu * spill / POOL
    t_wa = (ATTN / TP) / HBM_BW
    t_ca = bs * (score + proj) / (N * PEAK * ua)
    t_cm = bs * moe_f / (N * PEAK * ue)
    pairs = (e_hot / e) * (bs * TOPK / N)
    t_ep = pairs * 2 * H * L_MOE / D2D
    toks = math.ceil(bs / DP)
    t_ar = 2 * (TP - 1) / TP * (toks * H * 2) * L / P2P
    t_attn = max(t_ca, t_kh + t_wa, t_pool, t_ar)
    t_moe = max(t_cm, tw, t_ep)
    return {
        "t": t_attn + t_moe, "tw": tw, "cm": t_cm, "ar": t_ar, "ca": t_ca,
        "kh": t_kh, "wa": t_wa, "pool": t_pool, "ep": t_ep, "spill": spill,
        "frac": frac, "ehot": e_hot, "ue": ue, "ua": ua, "tok": tok, "toks": toks,
        "bytes": e_hot * ONE + ONE,
    }


def hbm_cap(s):
    return int(ROOM // (BPT * s))


def search_bs(s, limit):
    if decode(1, s)["t"] > limit:
        return 0
    ceiling = 500_000
    lo, hi, best = 1, 1, 0
    while hi < ceiling and decode(hi, s)["t"] <= limit:
        best = hi
        nxt = min(ceiling, hi * 2)
        if nxt == hi:
            break
        hi = nxt
    while lo <= hi:
        mid = (lo + hi) // 2
        if decode(mid, s)["t"] <= limit:
            best = mid
            lo = mid + 1
        else:
            hi = mid - 1
    return best


def prefill_batch(bs, s, hit=0.9):
    q = (1 - hit) * s
    tokens = bs * q
    sel = min(2048, s)
    idx = bs * N_INDEX * 8192 * q * s
    core = bs * L * 65536 * sel * q
    proj = bs * 2 * ATTN * q
    moe = bs * 2 * (TOPK + 1) * ONE * q
    den = bs * 2 * DENSE * L_DENSE * q
    ua = align(max(tokens / DP, 1e-9))
    frac = 1 - math.exp(-TOPK * tokens / E) if tokens > 0 else 0.0
    e = E / N
    std = math.sqrt(max(e * frac * (1 - frac), 0.0))
    e_hot = min(e, e * frac + 2 * std) if frac > 0 else 0.0
    ue = align((tokens * TOPK / E) if frac > 1e-12 else 1.0)
    t_idx = idx / (ua * N * PEAK)
    t_attn = (idx + core + proj) / (ua * N * PEAK)
    t_moe = moe / (ue * N * PEAK)
    t_den = den / (ua * N * PEAK)
    t_we = ((e_hot * ONE + ONE + L_DENSE * DENSE) if e_hot > 0 else (ONE + L_DENSE * DENSE)) / HBM_BW
    t_kv = bs * hit * s * BPT / (N * PARAM)
    t_ar = 2 * (TP - 1) / TP * ((tokens / DP) * H * 2) * L / P2P
    t_ep = ((e_hot / e) * (tokens * TOPK / N) if e else 0.0) * 2 * H * L_MOE / D2D
    t_wa = (ATTN / TP) / HBM_BW
    t_side = max(t_moe + t_den, t_we, t_ep)
    t = t_kv + max(t_attn, t_ar, t_wa) + t_side
    return {
        "t": t, "Q": tokens, "t_ar": t_ar, "t_attn": t_attn, "t_idx": t_idx,
        "t_moe": t_moe, "t_den": t_den, "t_ep": t_ep, "t_we": t_we, "tkv": t_kv,
        "t_side": t_side, "cap": hbm_cap(s),
    }


def max_prefill_bs(s, hit=0.9, limit=10.0):
    cap = hbm_cap(s)
    if cap < 1:
        return 0
    lo, hi, best = 1, 1, 0
    while hi < cap and prefill_batch(hi, s, hit)["t"] <= limit:
        hi = min(cap, hi * 2)
        if hi == cap:
            break
    while lo <= hi:
        mid = (lo + hi) // 2
        if prefill_batch(mid, s, hit)["t"] <= limit:
            best = mid
            lo = mid + 1
        else:
            hi = mid - 1
    return best


def fmt_tok(n):
    if n >= 1e6:
        return f"{n / 1e6:.2f}×10^6"
    if n >= 1e3:
        return f"{n / 1e3:.1f}×10^3"
    return f"{n:.0f}"


def fmt_len(s):
    if s >= 1048576 and s % 1048576 == 0:
        return f"{s // 1048576}M"
    if s % 1024 == 0 and s >= 1024:
        return f"{s // 1024}K"
    return str(s)


def fmt_bytes(n):
    return f"{n / 1e9:.2f} GB" if n >= 1e9 else f"{n / 1e6:.1f} MB"


def fmt_count(n):
    if n >= 10000:
        v = n / 10000
        return f"{v:.1f} 万" if v >= 10 else f"{v:.2f} 万"
    return str(n)


def per_card(bs):
    v = bs / N
    return f"{v:.1f}" if v >= 1 else f"{v:.2f}"


def ms(v):
    x = v * 1e3
    if x >= 100:
        return f"{x:.0f} 毫秒"
    return f"{x:.1f} 毫秒"


def ms2(v):
    return f"{v * 1e3:.2f} 毫秒"


def sec(v):
    if v >= 9.95:
        return f"{v:.1f} 秒"
    return f"{v:.2f} 秒"


def where(spill):
    if spill < 0.005:
        return "全在 HBM"
    return f"{spill * 100:.1f}% 在总线内存池"


def hbm_span(rows, batch_key, dec_key, bs):
    labels = []
    for row in rows:
        if row[batch_key] == bs and row[dec_key]["spill"] < 0.005:
            labels.append(fmt_len(row["S"]))
        else:
            break
    if not labels:
        return "短输入"
    if len(labels) == 1:
        return labels[0]
    return f"{labels[0]} – {labels[-1]}"


def main():
    s_max = max_input()
    s_show = int(s_max)
    p = prefill_one(s_show)
    weight_gb = (E * ONE / N + ONE + L_DENSE * DENSE + ATTN / TP) / 1e9
    kv_gb = s_show * BPT / 1e9
    bs10 = search_bs(4096, 0.010)
    bs30 = search_bs(4096, 0.030)
    d10 = decode(bs10, 4096)
    d30 = decode(bs30, 4096)
    full_w = (E / N * ONE + ONE) / HBM_BW
    linear = p["core"] + p["moe"] + p["den"] + p["proj"]

    lengths = (4096, 8192, 32768, 131072, 262144, 524288, 1048576)
    sweep = []
    for s in lengths:
        b10 = search_bs(s, 0.010)
        b30 = search_bs(s, 0.030)
        sweep.append({
            "S": s, "store": BPT * s, "cap": hbm_cap(s),
            "b10": b10, "b30": b30, "d10": decode(b10, s), "d30": decode(b30, s),
        })
    extra = {}
    for s in (65_536, 2_000_000, s_show):
        b10 = search_bs(s, 0.010)
        b30 = search_bs(s, 0.030)
        extra[s] = {"b10": b10, "b30": b30, "d10": decode(b10, s), "d30": decode(b30, s), "cap": hbm_cap(s)}
    by_len = {row["S"]: row for row in sweep}

    pb = []
    for s in lengths:
        bs = max_prefill_bs(s)
        row = prefill_batch(bs, s)
        row["S"] = s
        row["bs"] = bs
        pb.append(row)
    by_s = {row["S"]: row for row in pb}
    p4, p1m = by_s[4096], by_s[1048576]
    one_1m = prefill_batch(1, 1048576)
    span10 = hbm_span(sweep, "b10", "d10", bs10)
    span30 = hbm_span(sweep, "b30", "d30", bs30)
    prefill_rps = p4["bs"] / 10.0
    dec10_rps = bs10 / 0.010 / 512
    dec30_rps = bs30 / 0.030 / 512
    pd10 = prefill_rps / dec10_rps
    pd30 = prefill_rps / dec30_rps
    s_label = f"{s_show / 1e6:.2f}×10^6"
    total10 = E * ONE + REST
    total53 = E0 * ONE + REST
    active = TOPK * ONE + ONE + L_DENSE * DENSE + ATTN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "10T GLM：N+2 @ 2TB @ 480 GB/s 内存池上的 TTFT / TPOT"
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    chrome(s1, 1, "同一骨架，只加路由专家", "GLM-5.3 发布权重    ·    外推到 10T")
    textbox(
        s1, 0.42, 0.90, 12.5, 0.30,
        "最新发布权重是 GLM-5.3，架构与 GLM-5.2 相同。官方口径 744B / 40B。10T 冻结层数、宽度和 top-8，只加路由专家。",
        14, False, INK, anchor=MSO_ANCHOR.MIDDLE,
    )
    rows = [
        ["项目", "GLM-5.3 发布权重", "10T 外推", "备注"],
        ["回推总参数", "7.433×10^11", "1.000×10^13", "不含 MTP。官方 744B，差额是 MTP 与舍入"],
        ["路由专家", "256", "3526", "×13.8。增量只记在专家个数上"],
        ["激活参数", "3.92×10^10", "3.92×10^10", "每 token 仍是 8 个路由专家 + 1 个共享"],
        ["激活比", "5.28%", "0.39%", "专家表变大，单 token 激活不变"],
        ["FP8 专家表", "0.725 TB", "9.983 TB", "单专家全层 2.831 GB。发布精度是 FP8"],
        ["层 / 隐藏 / 中间维", "78 / 6144 / 2048", "同左", "前 3 层稠密，后 75 层 MoE"],
        ["注意力", "MLA-256，DSA top-2048", "同左", "21 层有 indexer，每 4 层共享一次"],
        ["单 token 专家通信", "1.106×10^7 字节", "同左", "8 × 3 × 6144 × 75，与专家总数无关"],
        ["单 token KV", "47700 字节", "同左", "78×576 + 21×132，按 FP8 布局"],
    ]
    add_table(
        s1, 0.42, 1.22, [2.35, 2.85, 2.45, 4.85], rows,
        [0.30] + [0.40] * 9, ["head"] + ["body"] * 9, size=12,
    )
    note = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(5.28), Inches(12.5), Inches(1.58))
    solid(note, GOLD_BG)
    textbox(
        s1, 0.58, 5.34, 12.18, 1.46,
        "外推式：总参数 = 专家数 × 2.831×10^9 + 1.848×10^10。单个专家 = 3 × 6144 × 2048 × 75。\n"
        "N+2 @ 2TB @ 480 GB/s 内存池：发布权重是 FP8，GEMM 用 7.50 P，即 15 P FP4 的一半。HBM 432 GB，14.4 TB/s 取 80% 为 11.52 TB/s。"
        "H2D 每卡 1.6 Tb/s 取 80% 为 160 GB/s，只在 prefill。机内 DDR5 每机 1.6 TB/s、1.3 TB。\n"
        "D2D 每卡 1.2 TB/s 取 80% 为 0.96 TB/s，专家通信走这里。内存池每卡 480 GB/s 取 80% 为 384 GB/s，容量不设上限，供 decode 换入。"
        "8 卡 P2P 与参数面都是 800 Gb/s，取 80% 为 80 GB/s。超节点 128 卡，机内 8 卡 fullmesh。",
        13, False, GOLD, anchor=MSO_ANCHOR.MIDDLE,
    )
    s1.notes_slide.notes_text_frame.text = (
        f"GLM-5.3 config：78 层，隐藏 6144，MoE 中间维 2048，256 个路由专家，1 个共享，top-8，前 3 层稠密，中间维 12288。"
        f"MLA：kv_lora 512 + rope 64 = 576，64 头。DSA index_topk 2048。indexer 在 21 层计算，其余层复用下标。"
        f"主干回推 {total53/1e9:.1f}B，激活 {active/1e9:.2f}B。官方口径 744B-A40B。"
        f"10T：E={E}，总参数 {total10:.6e}。专家 FP8 字节等于参数个数。"
        f"每卡路由专家 {E*ONE/N/1e9:.2f} GB，共享 {ONE/1e9:.2f} GB，稠密 {L_DENSE*DENSE/1e9:.2f} GB，注意力 {ATTN/TP/1e9:.2f} GB。"
        f"扣掉 6 GB 预留后，每卡 KV 余量 {(HBM_CAP - E*ONE/N - ONE - L_DENSE*DENSE - ATTN/TP - RESERVE)/1e9:.2f} GB，128 卡 {ROOM/1e12:.2f} TB。"
        "词表 154880 和 MTP 留在参数回推里，时延页的显存扣减沿用原先的 GLM 页，不另扣这两项。"
        "后面几页是 N+2 @ 2TB @ 480 GB/s 内存池。FP4 峰值是 15 P。GLM-5.3 发布精度是 FP8，GEMM 取一半，7.50 P。注意力和专家用同一个峰值。"
        "H2D 1.6 Tb/s = 200 GB/s 峰值，80% 后 160 GB/s，只在 prefill。DDR5 的 1.6 TB/s 和 1.3 TB 按一台 8 卡机器记，每卡 160 GB/s、162.5 GB。"
        "D2D 每卡 1.2 TB/s，80% 后 0.96 TB/s，专家通信走这里。总线内存池每卡 480 GB/s，80% 后 384 GB/s，容量不设上限，专供 decode 换入。"
        "若把 480 GB/s 理解成 128 卡共享，短输入的两档全局批不变，更长上下文的批会下降。本页按每卡。"
        "机内 8 卡 fullmesh，P2P 800 Gb/s，80% 后 80 GB/s。参数面同样。注意力 TP=8，16 台机器做 DP=16。MoE EP=128 走 D2D。"
        "整卡专家加共享读完是 "
        f"{full_w*1e3:.2f} ms，短于 10 ms，所以 10 毫秒这一档也能把每卡 27.5 个专家读满。"
    )

    s2 = prs.slides.add_slide(blank)
    chrome(s2, 2, f"10 秒对应的最大输入是 {s_label}", "Prefill    ·    TTFT ≤ 10 秒")
    banner(
        s2, 0.92,
        f"一条请求放在一个 128 卡副本上。计算 {p['tc']:.2f} 秒，读权重 {ms2(p['tw'])}，送 KV {ms2(p['ship'])}。S² 占 {p['idx']:.2f} 秒，线性项占 {linear:.2f} 秒。",
        size=13,
    )
    textbox(s2, 0.42, 1.46, 6.2, 0.24, "并行策略", 13, True, ACCENT)
    add_table(
        s2, 0.42, 1.74, [1.55, 5.40, 5.55],
        [
            ["", "Attention", "MoE"],
            ["并行", "机内 TP=8\n16 台机器做 DP=16", "超节点 EP=128\n一份模型铺满 128 卡"],
            ["通信域", "8 卡 fullmesh\n有效 80 GB/s", "128 卡 D2D，有效 0.96 TB/s\nKV 走参数面，有效 80 GB/s"],
            ["每卡权重", f"注意力 {ATTN/1e9:.2f} GB / 8\n= {ATTN/TP/1e9:.2f} GB FP8",
             f"路由专家 {E/N:.1f} 个，{E*ONE/N/1e9:.1f} GB FP8\n共享专家 {ONE/1e9:.2f} GB，每卡一份"],
            ["本阶段角色", "S² 只来自 21 层 indexer\n决定输入能写多长", f"专家 GEMM 随 token 线性\n到这个长度约占 {sec(p['moe'])}"],
        ],
        [0.28, 0.50, 0.56, 0.52, 0.50],
        ["head", "body", "body", "body", "body"],
    )
    textbox(s2, 0.42, 4.26, 12.5, 0.22, "10 秒怎么落到这个输入长度", 13, True, ACCENT)
    add_table(
        s2, 0.42, 4.60, [1.70, 7.55, 3.25],
        [
            ["这一笔", "算法", f"S = {s_show/1e6:.2f}×10^6 时"],
            ["indexer", "21 层 × 8192 × S²，除以 128 × 7.50×10^15", sec(p["idx"])],
            ["稀疏 MLA", "78 层，每个 token 看 2048 条 latent", sec(p["core"])],
            ["专家 GEMM", "9 个专家 × 2.831×10^9 × 2 × S，除以同一批卡的 FP8 算力", sec(p["moe"])],
            ["投影 + 稠密", "注意力投影，加前 3 层稠密 MLP", sec(p["proj"] + p["den"])],
            ["读权重 + 送 KV", f"每卡 {weight_gb:.2f} GB / 11.52 TB/s；KV {kv_gb:.1f} GB / (128×80 GB/s)", f"{ms2(p['tw'])} + {ms2(p['ship'])}"],
            ["合计", "计算和读权重取长的一项，再加送 KV。机内 allreduce 记在下一页", "10.0 秒"],
        ],
        [0.28, 0.28, 0.28, 0.28, 0.28, 0.30, 0.30],
        ["head", "body", "body", "body", "body", "body", "sum"],
        size=12,
    )
    s2.notes_slide.notes_text_frame.text = (
        "单请求不跨副本。增加 Prefill 节点是多放几份 128 卡副本，单条 TTFT 不变。"
        "时间 = max(计算, 读权重) + 把 KV 送到 Decode。计算全部除以 128×7.50×10^15。"
        "indexer 点积：32 头 × 128 维 × 2 = 8192 FLOP/对，只有 21 层计算。"
        "稀疏 MLA：64 头 × 256 维，QK 和 AV 合计每条 65536 FLOP，78 层都做，每 token 只看 2048 条。"
        f"S={s_show}：indexer {p['idx']:.4f}s，稀疏 {p['core']:.4f}s，专家 {p['moe']:.4f}s，投影 {p['proj']:.4f}s，稠密 {p['den']:.4f}s，合计计算 {p['tc']:.4f}s。"
        f"读权重 {p['tw']*1e3:.2f} ms，KV 搬运 {p['ship']*1e3:.2f} ms。S² {p['idx']:.2f}s，线性项 {linear:.2f}s。"
        f"4K 单条 TTFT 是 {prefill_one(4096)['t']*1e3:.2f} 毫秒，停在读权重。1M 单条、无命中是 {prefill_one(1048576)['t']:.3f} 秒。"
        f"该长度的 KV 约 {kv_gb:.1f} GB，经 128×80 GB/s 的参数面送出。HBM 余量 {ROOM/1e12:.2f} TB，这一条放得下，H2D 不进这条 10 秒。"
        "冷启动单请求的机内 allreduce 记在下一页的批上。"
    )

    s3 = prs.slides.add_slide(blank)
    chrome(s3, 3, f"4K 的 prefill 全局批是 {p4['bs']}，1M 是 {p1m['bs']}", "Prefill    ·    TTFT ≤ 10 秒    ·    KV 命中 90%")
    banner(
        s3, 0.92,
        "多条同时进同一份 128 卡副本。命中 90% 的 KV 走参数面，只现算 10%。注意力侧取 allreduce，再加 D2D 上的专家通信。",
        size=14,
    )
    textbox(s3, 0.42, 1.44, 12.5, 0.22, "10 秒在两个端点上怎么加完", 13, True, ACCENT)
    walk = [
        ["笔", "算法", f"4K，全局批 {p4['bs']}", f"1M，全局批 {p1m['bs']}"],
        ["现算 token", "批 × 输入 × 0.1", fmt_tok(p4["Q"]), fmt_tok(p1m["Q"])],
        ["机内 allreduce", "每组现算 token × 6144 × 2 字节\n78 层 × 2×7/8 / 80 GB/s", sec(p4["t_ar"]), sec(p1m["t_ar"])],
        ["注意力计算", "indexer、2048 条稀疏 MLA、投影\n除以 FP8 的 7.50 P", sec(p4["t_attn"]), sec(p1m["t_attn"])],
        ["专家通信", "最忙卡 BF16 combine\n除以每卡 0.96 TB/s 的 D2D", sec(p4["t_ep"]), sec(p1m["t_ep"])],
        ["相加", "注意力侧取 allreduce，专家侧取 EP，再加参数面读 KV", sec(p4["t"]), sec(p1m["t"])],
    ]
    add_table(
        s3, 0.42, 1.68, [1.70, 5.15, 2.85, 2.80], walk,
        [0.26, 0.30, 0.50, 0.48, 0.50, 0.32],
        ["head", "body", "key", "body", "body", "sum"],
    )
    textbox(s3, 0.42, 4.18, 12.5, 0.22, "同一条 10 秒，输入从 4K 到 1M", 13, True, ACCENT)
    sweep_head = ["输入", "10 秒全局批", "每卡", "现算 token", "allreduce", "专家通信", "注意力计算", "读 KV"]
    sweep_body, sweep_kinds = [], ["head"]
    for row in pb:
        sweep_kinds.append("key" if row["S"] in (4096, 1048576) else "body")
        sweep_body.append([
            fmt_len(row["S"]), f"{row['bs']}", per_card(row["bs"]), fmt_tok(row["Q"]),
            sec(row["t_ar"]), sec(row["t_ep"]), sec(row["t_attn"]), ms(row["tkv"]),
        ])
    add_table(
        s3, 0.42, 4.50, [1.05, 1.55, 0.95, 1.70, 1.55, 1.55, 1.85, 2.30],
        [sweep_head] + sweep_body, [0.24] + [0.24] * len(sweep_body), sweep_kinds, size=11,
    )
    textbox(
        s3, 0.42, 6.66, 12.5, 0.28,
        f"这些批的 KV 都放得进 HBM：4K 能放 {fmt_count(p4['cap'])} 条，1M 能放 {p1m['cap']} 条。参数面读命中 KV 约 {ms(p4['tkv'])}。",
        13, False, MUTED, anchor=MSO_ANCHOR.MIDDLE,
    )
    s3.notes_slide.notes_text_frame.text = (
        "全局批是 128 卡副本上的条数。注意力 TP=8、DP=16，MoE EP=128。"
        "KV 命中 90%。每条只把 10% 的 token 送进计算，这 10% 的 query 仍看全长上下文。"
        "时间 = 参数面读 KV + 注意力侧较长的一项 + 专家侧较长的一项。"
        "注意力侧在注意力计算、机内 allreduce、注意力权重之间取最大。专家侧在专家 GEMM 加稠密 MLP、读专家权重、EP 之间取最大。这一页专家侧取 D2D combine。"
        "机内 allreduce：每组 token 数 = 现算 token / 16，每 token 6144×2 字节，环上 2×7/8，78 层，除以 80 GB/s。"
        f"4K 全局批 {p4['bs']}：现算 {p4['Q']:.4e}，allreduce {p4['t_ar']:.3f}s，注意力 {p4['t_attn']:.3f}s，"
        f"专家 GEMM {p4['t_moe']:.3f}s，稠密 {p4['t_den']:.4f}s，EP {p4['t_ep']:.3f}s，读权重 {p4['t_we']*1e3:.2f} ms，"
        f"读 KV {p4['tkv']*1e3:.1f} ms，合计 {p4['t']:.3f}s。再加一条是 {prefill_batch(p4['bs']+1, 4096)['t']:.3f}s。"
        f"1M 全局批 {p1m['bs']}：allreduce {p1m['t_ar']:.3f}s，注意力 {p1m['t_attn']:.3f}s（indexer {p1m['t_idx']:.3f}s），"
        f"专家侧 {p1m['t_side']:.3f}s，EP {p1m['t_ep']:.3f}s，合计 {p1m['t']:.3f}s。"
        f"HBM 余量 {ROOM/1e12:.2f} TB。4K 能放 {p4['cap']} 条，1M 能放 {p1m['cap']} 条，都高于这页的时间批。"
        f"单条 1,048,576 token、90% 命中的 TTFT 是 {one_1m['t']:.4f} 秒，其中 allreduce {one_1m['t_ar']:.4f} 秒。"
    )

    s4 = prs.slides.add_slide(blank)
    chrome(s4, 4, f"10 毫秒、短输入，全局批 {bs10}", "Decode    ·    TPOT ≤ 10 毫秒")
    banner(
        s4, 0.92,
        f"整卡专家 {ms2(d10['tw'])}，机内 allreduce {ms2(d10['ar'])}。专家 GEMM {ms2(d10['cm'])}，D2D 上的 EP {ms2(d10['ep'])}。",
        size=14,
    )
    textbox(s4, 0.42, 1.46, 6.0, 0.24, "并行策略", 13, True, ACCENT)
    textbox(s4, 7.15, 1.46, 5.7, 0.24, "短输入（4K）这一步怎么加到 10 毫秒", 13, True, ACCENT)
    add_table(
        s4, 0.42, 1.78, [1.15, 2.55, 2.70],
        [
            ["", "Attention", "MoE"],
            ["并行", "TP=8，DP=16\n每组 token 数 = ⌈BS / 16⌉", f"EP=128\n每卡 {E/N:.1f} 个专家"],
            ["通信", f"每层机内 allreduce\n这一步 {ms2(d10['ar'])}", f"最忙卡 BF16 combine\n这一步 {ms2(d10['ep'])}"],
            ["访存", f"注意力权重 {ATTN/TP/1e9:.2f} GB\n这一步 {ms2(d10['wa'])}", f"{E/N:.1f} 个专家这一步 {ms2(d10['tw'])}\n整张本地表都读到"],
            ["计算", "行数 = ⌈BS / 16⌉\n按 16 行砖对齐", f"每个命中专家约 {d10['tok']:.1f} 个 token\n16 行砖，效率 {d10['ue']*100:.0f}%"],
        ],
        [0.28, 0.58, 0.58, 0.62, 0.58],
        ["head", "body", "body", "key", "body"], size=12,
    )
    add_table(
        s4, 7.05, 1.78, [1.20, 3.50, 1.15],
        [
            ["笔", "白话", "时间"],
            ["命中", f"1 − exp(−8×{bs10} / {E}) = {d10['frac']*100:.0f}%\n最忙卡按均值 + 2σ，顶到 {d10['ehot']:.1f} 个", "—"],
            ["读专家", f"{d10['bytes']/1e9:.2f} GB（路由 + 共享）\n除以 11.52 TB/s", ms2(d10["tw"])],
            ["机内通信", f"每组 {d10['toks']} 条 × 6144 × 2 字节\n78 层 × 2×7/8 / 80 GB/s", ms2(d10["ar"])],
            ["专家计算", f"每个命中专家约 {d10['tok']:.1f} 个 token\n16 行砖，效率 {d10['ue']*100:.0f}%", ms2(d10["cm"])],
            ["专家通信", "最忙卡送出的 BF16 combine / 0.96 TB/s", ms2(d10["ep"])],
            ["相加", f"注意力侧 {d10['ar']*1e3:.2f}，专家侧 {d10['tw']*1e3:.2f}", "10.0 毫秒"],
        ],
        [0.28, 0.54, 0.50, 0.50, 0.48, 0.40, 0.38],
        ["head", "body", "key", "key", "body", "body", "sum"], size=12,
    )
    chips(
        s4,
        [
            (span10, str(bs10), f"全局批，每卡 {per_card(bs10)} 条\n全在 HBM"),
            ("512K", str(by_len[524288]["b10"]), f"全局批，每卡 {per_card(by_len[524288]['b10'])} 条\n{where(by_len[524288]['d10']['spill'])}"),
            ("1M", str(by_len[1048576]["b10"]), f"全局批，每卡 {per_card(by_len[1048576]['b10'])} 条\n{where(by_len[1048576]['d10']['spill'])}"),
            (s_label, str(extra[s_show]["b10"]), f"全局批，每卡 {per_card(extra[s_show]['b10'])} 条\n{where(extra[s_show]['d10']['spill'])}"),
        ],
    )
    s4.notes_slide.notes_text_frame.text = (
        "Decode 与 Prefill 同一份切分：TP=8，DP=16，EP=128。一步 = 注意力侧最长的一笔 + 专家侧最长的一笔。"
        f"BS={bs10}，命中率 {d10['frac']:.4f}，最忙卡 {d10['ehot']:.3f} 个，已经顶到每卡全部本地专家。"
        f"读权重 {d10['bytes']/1e9:.3f} GB / 11.52 TB/s = {d10['tw']*1e3:.3f} ms。整卡读完是 {full_w*1e3:.3f} ms。"
        f"机内 allreduce：ceil({bs10}/16) = {d10['toks']} 条，每条 6144×2 字节，环上 2×7/8，78 层，除以 80 GB/s，得 {d10['ar']*1e3:.3f} ms。"
        f"两项相加 {d10['t']*1e3:.3f} ms，表上记 10.0 毫秒。{bs10+1} 条是 {decode(bs10+1, 4096)['t']*1e3:.3f} ms。"
        f"专家计算 {d10['cm']*1e3:.3f} ms，效率 {d10['ue']:.3f}。每个命中专家只有约 {d10['tok']:.2f} 个 token，16 行砖的效率是 {d10['ue']*100:.0f}%。"
        f"EP {d10['ep']*1e3:.3f} ms。注意力计算 {d10['ca']*1e3:.3f} ms。注意力权重 {d10['wa']*1e3:.3f} ms。"
        f"{span10} 的全局批都是 {bs10}，KV 全在 HBM。"
        f"512K 降到 {by_len[524288]['b10']}，存量 {by_len[524288]['d10']['spill']*100:.1f}% 在总线内存池，池上稀疏读 {by_len[524288]['d10']['pool']*1e3:.2f} ms，和读专家之后剩下的预算对齐。"
        f"1M 降到 {by_len[1048576]['b10']}，2M 降到 {extra[2_000_000]['b10']}，{s_label} 降到 {extra[s_show]['b10']}。"
        "换入走每卡 384 GB/s 的总线内存池，即 480 GB/s 的 80%，容量不设上限。H2D 只在 prefill。"
        "前 3 层稠密 MLP 的 GEMM 在这个批上远短于读专家，沿用原先的 GLM 解码式，不单列。"
    )

    s5 = prs.slides.add_slide(blank)
    chrome(s5, 5, f"30 毫秒、短输入，全局批 {bs30}", "Decode    ·    TPOT ≤ 30 毫秒")
    banner(
        s5, 0.92,
        f"专家整卡读完 {ms2(d30['tw'])}，机内 allreduce {ms2(d30['ar'])}。D2D 上的 EP {ms2(d30['ep'])}。",
        size=14,
    )
    textbox(s5, 0.42, 1.46, 6.2, 0.24, "并行策略    ·    和 10 毫秒同一份切分", 13, True, ACCENT)
    textbox(s5, 7.15, 1.46, 5.7, 0.24, f"短输入、全局批 {bs30}，这一步怎么加", 13, True, ACCENT)
    add_table(
        s5, 0.42, 1.78, [1.15, 2.55, 2.70],
        [
            ["", "Attention", "MoE"],
            ["并行", f"仍是 TP=8，DP=16\n{bs30} / 16 = {d30['toks']} 条 / 组", f"仍是 EP=128\n每卡 {E/N:.1f} 个专家全醒"],
            ["通信", "allreduce 随批线性涨\n这一档是注意力侧最长", f"combine 走 0.96 TB/s\n{ms2(d30['ep'])}，短于读权重"],
            ["访存", "短输入全在 HBM\n放不下走每卡内存池 384 GB/s", f"{E*ONE/N/1e9:.1f} GB 路由 + {ONE/1e9:.2f} GB 共享\n/ 11.52 TB/s = {ms2(d30['tw'])}"],
            ["计算", f"注意力计算 {ms2(d30['ca'])}\n短于 allreduce", f"每专家约 {d30['tok']:.1f} token\n16 行砖，效率 {d30['ue']*100:.0f}%，{ms2(d30['cm'])}"],
        ],
        [0.28, 0.58, 0.58, 0.62, 0.58],
        ["head", "body", "key", "body", "body"], size=12,
    )
    add_table(
        s5, 7.05, 1.78, [1.25, 3.45, 1.15],
        [
            ["笔", "白话", "时间"],
            ["读权重", f"命中率 {d30['frac']*100:.0f}%，{d30['ehot']:.1f} 个专家全读\n{d30['bytes']/1e9:.2f} GB / 11.52 TB/s", ms2(d30["tw"])],
            ["机内通信", f"每组 {d30['toks']} 条 × 6144 × 2 字节\n78 层 × 2×7/8 / 80 GB/s", ms2(d30["ar"])],
            ["专家计算", f"{bs30} × 8 / {E} ≈ {d30['tok']:.1f} token / 专家\n对齐效率 {d30['ue']*100:.0f}%", ms2(d30["cm"])],
            ["专家通信", "最忙卡 combine / 0.96 TB/s", ms2(d30["ep"])],
            ["注意力计算", "短输入的 indexer、稀疏 MLA 和投影", ms2(d30["ca"])],
            ["相加", f"注意力侧 {d30['ar']*1e3:.2f}，专家侧 {d30['tw']*1e3:.2f}", "30.0 毫秒"],
        ],
        [0.28, 0.52, 0.52, 0.48, 0.40, 0.40, 0.42],
        ["head", "key", "key", "body", "body", "body", "sum"], size=12,
    )
    chips(
        s5,
        [
            (span30, str(bs30), f"全局批，每卡 {per_card(bs30)} 条\n全在 HBM"),
            ("128K", str(by_len[131072]["b30"]), f"全局批，每卡 {per_card(by_len[131072]['b30'])} 条\n{where(by_len[131072]['d30']['spill'])}"),
            ("1M", str(by_len[1048576]["b30"]), f"全局批，每卡 {per_card(by_len[1048576]['b30'])} 条\n{where(by_len[1048576]['d30']['spill'])}"),
            (s_label, str(extra[s_show]["b30"]), f"全局批，每卡 {per_card(extra[s_show]['b30'])} 条\n{where(extra[s_show]['d30']['spill'])}"),
        ],
        number_size=26,
    )
    s5.notes_slide.notes_text_frame.text = (
        f"BS={bs30}，命中率 {d30['frac']:.4f}，每卡 {d30['ehot']:.3f} 个专家全醒。"
        f"读权重 {d30['tw']*1e3:.3f} ms。allreduce ceil({bs30}/16)={d30['toks']} 条，{d30['ar']*1e3:.3f} ms。"
        f"两项相加 {d30['t']*1e3:.3f} ms，表上记 30.0 毫秒。{bs30+1} 条是 {decode(bs30+1, 4096)['t']*1e3:.3f} ms。"
        f"专家计算 {d30['cm']*1e3:.3f} ms，效率 {d30['ue']:.3f}。EP {d30['ep']*1e3:.3f} ms。注意力计算 {d30['ca']*1e3:.3f} ms。"
        f"显存余量 {ROOM/1e12:.2f} TB。溢出进卡的口是每卡总线内存池 {POOL/1e9:.0f} GB/s，即 480 GB/s 的 80%，容量不设上限。D2D 上的专家通信是每卡 1.2 TB/s 取 80%，即 0.96 TB/s。"
        f"64K 仍是 {extra[65536]['b30']}，存量 {extra[65536]['d30']['spill']*100:.1f}% 在池上，稀疏读 {extra[65536]['d30']['pool']*1e3:.2f} ms，短于 allreduce {extra[65536]['d30']['ar']*1e3:.2f} ms，所以批保持。"
        f"128K 降到 {by_len[131072]['b30']}，存量 {by_len[131072]['d30']['spill']*100:.1f}% 在池上，稀疏读 {by_len[131072]['d30']['pool']*1e3:.2f} ms，已经绑住这一步。"
        f"256K 降到 {by_len[262144]['b30']}，1M 降到 {by_len[1048576]['b30']}，2M 降到 {extra[2_000_000]['b30']}，{s_label} 降到 {extra[s_show]['b30']}。"
        f"4K、输出 512、命中 90%：Prefill {prefill_rps:.1f} 请求/秒。10 毫秒 Decode {dec10_rps:.1f} 请求/秒，P:D 约 1:{pd10:.1f}。"
        f"30 毫秒 Decode {dec30_rps:.1f} 请求/秒，P:D 约 1:{pd30:.1f}。"
    )

    s6 = prs.slides.add_slide(blank)
    chrome(s6, 6, f"128K 的 10 毫秒仍是 {bs10}，30 毫秒降到 {by_len[131072]['b30']}", "Decode    ·    不同输入下的最大全局批")
    banner(
        s6, 0.92,
        "全局批是 128 卡上的条数。放不下的 KV 进每卡总线内存池，480 GB/s 取 80% 为 384 GB/s，容量不设上限。",
        size=13,
    )
    head = ["输入", "每条 KV", "HBM 能放", "10 毫秒全局批", "每卡", "30 毫秒全局批", "每卡", "30 毫秒的 KV 放哪"]
    body, kinds = [], ["head"]
    for row in sweep:
        kinds.append("key" if row["d30"]["spill"] >= 0.005 else "body")
        body.append([
            fmt_len(row["S"]), fmt_bytes(row["store"]), fmt_count(row["cap"]),
            f"{row['b10']}", per_card(row["b10"]), f"{row['b30']}", per_card(row["b30"]),
            where(row["d30"]["spill"]),
        ])
    add_table(
        s6, 0.42, 1.50, [1.15, 1.35, 1.35, 1.70, 1.00, 1.75, 1.00, 3.20],
        [head] + body, [0.36] + [0.42] * len(body), kinds, size=12,
    )
    d128 = by_len[131072]["d30"]
    note = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(5.02), Inches(12.5), Inches(1.84))
    solid(note, GOLD_BG)
    textbox(
        s6, 0.58, 5.08, 12.18, 1.72,
        f"10 毫秒在 {span10} 都是 {bs10}，全在 HBM。一步是机内 allreduce {d10['ar']*1e3:.2f} 毫秒加读专家 {d10['tw']*1e3:.2f} 毫秒。"
        f"512K 降到 {by_len[524288]['b10']}（{where(by_len[524288]['d10']['spill'])}）。"
        f"1M 降到 {by_len[1048576]['b10']}，{s_label} 降到 {extra[s_show]['b10']}。\n"
        f"30 毫秒在 {span30} 仍是 {bs30}，全在 HBM。128K 降到 {by_len[131072]['b30']}，{where(d128['spill'])}，池上稀疏读绑住这一步。"
        f"1M 为 {by_len[1048576]['b30']}，{s_label} 为 {extra[s_show]['b30']}。\n"
        f"Prefill 的 10 秒全局批在第 3 页：4K 为 {p4['bs']}，1M 为 {p1m['bs']}。"
        f"4K 进、512 出、90% 命中时，P:D 在 10 毫秒约 1:{pd10:.1f}，在 30 毫秒约 1:{pd30:.1f}。",
        14, False, GOLD, anchor=MSO_ANCHOR.MIDDLE,
    )
    s6.notes_slide.notes_text_frame.text = (
        "存量 = S×47700 字节。一步真正读的是 21 层 indexer 的 132 字节/token，加上 78 层里最多 2048 条 latent，每条 576 字节。"
        f"HBM 余量 {ROOM/1e12:.2f} TB。溢出进 GPU 的口是每卡总线内存池，有效 {POOL/1e9:.0f} GB/s，容量不设上限。"
        "这个带宽是每卡 480 GB/s 取 80%。D2D 上的专家通信是每卡 1.2 TB/s 取 80%，即 0.96 TB/s。H2D 1.6 Tb/s 只在 prefill，decode 换入走内存池。"
        + " ".join(
            f"S={row['S']} b10={row['b10']} spill10={row['d10']['spill']*100:.1f}% pool10={row['d10']['pool']*1e3:.2f} "
            f"b30={row['b30']} spill30={row['d30']['spill']*100:.1f}% pool30={row['d30']['pool']*1e3:.2f}"
            for row in sweep
        )
        + f" 2M 的 10 毫秒批 {extra[2_000_000]['b10']}，30 毫秒批 {extra[2_000_000]['b30']}。"
        f"单条冷 1M 的 TTFT 是 {prefill_one(1048576)['t']:.3f} 秒。单条 1M、90% 命中是 {one_1m['t']:.4f} 秒。"
        f"P:D 按 4K 进、512 出、90% 命中：Prefill {prefill_rps:.1f} 请求/秒，"
        f"10 毫秒 Decode {dec10_rps:.1f} 请求/秒（1:{pd10:.2f}），30 毫秒 Decode {dec30_rps:.1f} 请求/秒（1:{pd30:.2f}）。"
    )

    out = "/workspace/docs/glm-n2-2tb-480-slo.pptx"
    prs.save(out)
    print(out)
    print("ROOM_TB", round(ROOM / 1e12, 3), "S", s_show, "bs10", bs10, "bs30", bs30)
    print("pf", p4["bs"], p1m["bs"], "PD", round(pd10, 2), round(pd30, 2))
    print("span", span10, span30)
    print("tc", round(p["tc"], 4), "idx", round(p["idx"], 4), "ship", round(p["ship"] * 1e3, 2))
    if not (p["t"] <= 10.0 + 1e-6 and d10["t"] <= 0.010 + 1e-9 and d30["t"] <= 0.030 + 1e-9):
        raise SystemExit("SLO miss")
    if any(row["t_ep"] + 1e-9 < max(row["t_moe"] + row["t_den"], row["t_we"]) for row in pb):
        raise SystemExit("prefill expert side is no longer the D2D combine")
    if not (d10["ep"] < d10["tw"] and d30["ep"] < d30["tw"]):
        raise SystemExit("decode EP is no longer under the weight read")


if __name__ == "__main__":
    main()
