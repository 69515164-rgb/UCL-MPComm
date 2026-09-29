#!/usr/bin/env python3
"""Decode KV pool: when SSD is enough, and when DDR is required."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

W, H = 1920, 1080
OUT = Path("/workspace/docs")
SLIDE_DIR = OUT / "decode-kv-pool"
FONT_PATH = "/workspace/docs/fonts/NotoSansSC.ttf"

BG = (255, 255, 255)
INK = (22, 27, 34)
MUTED = (88, 98, 110)
ACCENT = (14, 74, 138)
SOFT = (241, 245, 249)
LINE = (214, 221, 229)
CHIP = (232, 239, 246)


def font(size: int, weight: int = 450) -> ImageFont.FreeTypeFont:
    fnt = ImageFont.truetype(FONT_PATH, size)
    fnt.set_variation_by_axes([weight])
    return fnt


def text_w(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> int:
    return int(draw.textlength(text, font=fnt))


def wrap(draw, text: str, fnt, max_w: int) -> list[str]:
    """Wrap on CJK characters. Keep a number or Latin word on one line."""
    lines: list[str] = []
    for para in text.split("\n"):
        tokens: list[str] = []
        buf = ""
        for ch in para:
            if ch.isascii() and not ch.isspace():
                buf += ch
            else:
                if buf:
                    tokens.append(buf)
                    buf = ""
                tokens.append(ch)
        if buf:
            tokens.append(buf)
        line = ""
        for tok in tokens:
            trial = line + tok
            if draw.textlength(trial, font=fnt) <= max_w:
                line = trial
                continue
            if line.strip():
                lines.append(line.rstrip())
            if draw.textlength(tok, font=fnt) <= max_w:
                line = tok.lstrip()
                continue
            piece = ""
            for ch in tok:
                if draw.textlength(piece + ch, font=fnt) <= max_w:
                    piece += ch
                else:
                    if piece:
                        lines.append(piece)
                    piece = ch
            line = piece
        if line.strip():
            lines.append(line.rstrip())
    return lines


def draw_text(draw, xy, text, fnt, fill=INK, anchor="lt"):
    draw.text(xy, text, font=fnt, fill=fill, anchor=anchor)


def wrap_draw(draw, xy, text, fnt, max_w, fill=INK, leading=1.35):
    x, y = xy
    lines = wrap(draw, text, fnt, max_w)
    h = int(fnt.size * leading)
    for line in lines:
        draw_text(draw, (x, y), line, fnt, fill)
        y += h
    return y


def canvas() -> Image.Image:
    return Image.new("RGB", (W, H), BG)


def top_bar(img: Image.Image) -> None:
    ImageDraw.Draw(img).rectangle((0, 0, W, 8), fill=ACCENT)


def footer(img: Image.Image, page: str) -> None:
    draw = ImageDraw.Draw(img)
    draw.line((64, 1028, W - 64, 1028), fill=LINE, width=1)
    draw_text(draw, (64, 1044), "Decode KV 池  ·  SSD 与 DDR", font(16, 450), MUTED)
    draw_text(draw, (W - 64, 1044), page, font(16, 550), ACCENT, anchor="rt")


def kicker_title(img: Image.Image, kicker: str, title: str, claim: str | None = None) -> int:
    top_bar(img)
    draw = ImageDraw.Draw(img)
    draw_text(draw, (64, 28), kicker, font(18, 650), ACCENT)
    draw_text(draw, (64, 58), title, font(36, 700), INK)
    y = 112
    if claim:
        y = wrap_draw(draw, (64, 108), claim, font(20, 450), W - 128, MUTED, 1.4)
        y += 8
    draw.line((64, y + 6, W - 64, y + 6), fill=LINE, width=1)
    return y + 22


def card(draw, box, fill=BG, outline=LINE, width=1, radius=16):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def slide_01() -> Image.Image:
    img = canvas()
    y0 = kicker_title(
        img,
        "结论",
        "Decode 何时用 SSD 池，何时必须用 DDR 池",
        "历史 KV 的数值不变，每步用新的 Q 重读一遍。池子只承接 HBM 放不下的缺页；这笔读取从第二个 token 起进入 TPOT。",
    )
    draw = ImageDraw.Draw(img)
    cols = [
        (
            "不进池",
            "HBM 装得下",
            "这条请求要常驻的历史 KV 能留在 decode 侧 HBM。每步只读 HBM，并把新 token 的 KV 写回 HBM。SSD 和 DDR 都不在 TPOT 上。",
            [
                "Prefill 交接在第一个 token 之前，记在 TTFT。",
                "冷前缀提前整段装进 HBM，装载也记在 TTFT。",
                "之后每步重读同一份历史，只追加一个 token。",
            ],
            "这一步不访问池子。",
        ),
        (
            "可以用 SSD",
            "大页把接口喂满",
            "缺页在 TPOT 上，但同一层按 token 连续存放，按层做一次大读。平均传输时间由接口带宽决定。",
            [
                "单次 IO 大到把介质时延摊薄，IOPS 不成瓶颈。",
                "合同看平均 TPOT。垃圾回收的长尾不写进预算。",
                "没有 TPOT 合同的离线吞吐，同样可以用 SSD。",
            ],
            "平均 TPOT 与 DDR 相同。",
        ),
        (
            "必须用 DDR",
            "喂不满，或要保 p99",
            "缺页在 TPOT 上，而且 SSD 的实际带宽或长尾进了本步预算。接口标称带宽相同也不够。",
            [
                "小页随机读：IO 次数随上下文线性增加，SSD 喂不满接口。",
                "合同写 p99：垃圾回收的停顿不随大流量摊薄。",
                "单次 IO 太小：介质时延本身大于剩余的 TPOT 预算。",
            ],
            "缺页读改走 DDR。",
        ),
    ]
    gap = 22
    left = 64
    width = (W - 128 - gap * 2) // 3
    top = y0 + 16
    bottom = 900
    for i, (tag, head, lead, bullets, foot) in enumerate(cols):
        x = left + i * (width + gap)
        fill = SOFT if i == 2 else BG
        card(draw, (x, top, x + width, bottom), fill=fill)
        draw.rectangle((x, top, x + width, top + 8), fill=ACCENT)
        draw_text(draw, (x + 28, top + 28), tag, font(18, 700), ACCENT)
        draw_text(draw, (x + 28, top + 62), head, font(28, 700), INK)
        yy = wrap_draw(draw, (x + 28, top + 114), lead, font(20, 450), width - 56, INK, 1.45)
        yy += 22
        for b in bullets:
            draw.ellipse((x + 32, yy + 10, x + 42, yy + 20), fill=ACCENT)
            yy = wrap_draw(draw, (x + 56, yy), b, font(18, 450), width - 88, MUTED, 1.4)
            yy += 16
        draw.line((x + 28, bottom - 78, x + width - 28, bottom - 78), fill=LINE, width=1)
        draw_text(draw, (x + 28, bottom - 52), foot, font(20, 700), ACCENT)
    footer(img, "01  /  05")
    return img


def slide_02() -> Image.Image:
    img = canvas()
    y0 = kicker_title(
        img,
        "计时",
        "只有缺页的重读进入 TPOT",
        "第一个 token 之前的搬运，包括交接和第一次把缺页拉进 HBM，都记在 TTFT。",
    )
    draw = ImageDraw.Draw(img)
    steps = [
        ("1", "交接", "TTFT", "Prefill 把 prompt 的历史 KV 交给 decode。数值到这里已经定稿。"),
        ("2", "第一个 token", "TTFT", "读 HBM 里的历史 KV，算出这个 token，把增量 KV 写回 HBM。这时若还要从池里取缺页，仍算在 TTFT。"),
        ("3", "后续每个 token", "TPOT", "用新的 Q 重读全部历史 KV。留在 HBM 的读 HBM，不在 HBM 的读池。被驱逐的页下一步还要再读。"),
        ("4", "逐层重叠", "TPOT", "下一层的取数可以盖在当前层的计算下面。取数不超过计算，池子只占带宽；超出的部分加进这一步。"),
    ]
    top = y0 + 18
    row_h = 168
    for i, (n, title, clock, body) in enumerate(steps):
        y = top + i * (row_h + 16)
        card(draw, (64, y, 1856, y + row_h), fill=BG)
        draw.rounded_rectangle((92, y + 52, 160, y + 120), radius=12, fill=ACCENT)
        draw_text(draw, (126, y + 86), n, font(28, 700), (255, 255, 255), anchor="mm")
        draw_text(draw, (188, y + 28), title, font(26, 700), INK)
        pill_w = text_w(draw, clock, font(16, 700)) + 28
        title_w = text_w(draw, title, font(26, 700))
        px = 188 + title_w + 16
        draw.rounded_rectangle((px, y + 28, px + pill_w, y + 64), radius=14, fill=CHIP)
        draw_text(draw, (px + 14, y + 34), clock, font(16, 700), ACCENT)
        wrap_draw(draw, (188, y + 80), body, font(20, 450), 1620, INK, 1.4)
    footer(img, "02  /  05")
    return img


def slide_03() -> Image.Image:
    img = canvas()
    y0 = kicker_title(
        img,
        "判据",
        "接口带宽相同，差别只在单次 IO 有多大",
        "平均时间对齐的条件，是池子真的把这条接口喂满。长上下文只提供字节数，IO 大小决定这些字节怎么走。",
    )
    draw = ImageDraw.Draw(img)
    card(draw, (64, y0 + 12, 940, 960), fill=BG)
    card(draw, (980, y0 + 12, 1856, 960), fill=SOFT)
    draw_text(draw, (96, y0 + 36), "每步访问量", font(22, 700), ACCENT)
    formulas = [
        "每 token 的 KV 字节 = 2 × 层数 × KV 头数 × 头维 × 每元素字节",
        "单步读取 = 当前缓存长度 × 每 token 的 KV 字节",
        "单步写入 = 每 token 的 KV 字节",
        "本步池读取 = 本步不在 HBM 中的历史 KV 字节",
    ]
    yy = y0 + 84
    for line in formulas:
        draw.rounded_rectangle((96, yy, 908, yy + 58), radius=10, fill=SOFT)
        draw_text(draw, (112, yy + 14), line, font(18, 550), INK)
        yy += 72
    note = "70B、GQA、BF16：80 层、8 个 KV 头、头维 128。每 token 320KB。缓存长度 32768 时，单步读取 10GB，其中每层 128MB。历史内容与上一步相同，只是多了一个 token。"
    wrap_draw(draw, (96, yy + 8), note, font(18, 450), 812, MUTED, 1.4)

    draw_text(draw, (1012, y0 + 36), "单次读取时间", font(22, 700), ACCENT)
    draw.rounded_rectangle((1012, y0 + 84, 1824, y0 + 150), radius=10, fill=BG)
    draw_text(draw, (1028, y0 + 102), "单次读取时间 = IO 字节 / 接口带宽 + 介质时延", font(18, 600), INK)
    rules = [
        ("时延可以忽略", "单次 IO 字节远大于接口带宽乘两种介质的时延差。按 50GB/s、时延差 100 微秒估算，门槛大约是 5MB。一层 128MB 的大读里，时延只占百分之几，SSD 和 DDR 的平均时间相同。"),
        ("IOPS 不成瓶颈", "单步 IO 次数 = 本步池读取 / 单次 IO 字节。一层一次读，一步只有几十个 IO。"),
        ("小页必须改用 DDR", "每层每个 token 约 4KB。32K 上下文一步约 260 万次读。要把 50GB/s 喂满，需要约每秒 1200 万次 4KB 读。SSD 的随机 IOPS 先到顶，实际带宽低于接口。"),
    ]
    yy = y0 + 174
    for title, body in rules:
        draw_text(draw, (1012, yy), title, font(20, 700), INK)
        yy = wrap_draw(draw, (1012, yy + 34), body, font(18, 450), 812, MUTED, 1.38)
        yy += 18
    footer(img, "03  /  05")
    return img


def slide_04() -> Image.Image:
    img = canvas()
    y0 = kicker_title(
        img,
        "页大小",
        "写用小页，读用大页",
        "读写页不必一样大。同一层的 KV 按 token 顺序连续存放，小页是这段区间上的刻度，大读是一次区间读。",
    )
    draw = ImageDraw.Draw(img)
    items = [
        ("写", "一个 token 在一层里的 KV，只有几 KB，追加在末尾。不重写已经写下的历史。"),
        ("读", "这一层里还不在 HBM 的历史，收成一次大读。全量注意力本步本来就要把这段全部读走。"),
        ("长度", "大读不超过 HBM 里的逐层缓冲，通常就是一层。整份历史 KV 不能当成一页读回来。"),
        ("扣除", "attention sink 和最近的 token 若已在 HBM，从这次区间里去掉，避免重复占用接口。"),
        ("下发", "逻辑上的一页，在池子上再拆成数 MB 的条带。单块介质的停顿不会卡住整层。"),
    ]
    top = y0 + 28
    row_h = 118
    for i, (tag, body) in enumerate(items):
        y = top + i * (row_h + 14)
        card(draw, (64, y, 1856, y + row_h), fill=BG if i % 2 == 0 else SOFT)
        draw.rounded_rectangle((96, y + 34, 224, y + 84), radius=12, fill=ACCENT)
        draw_text(draw, (160, y + 59), tag, font(22, 700), (255, 255, 255), anchor="mm")
        body_f = font(22, 500)
        lines = wrap(draw, body, body_f, 1560)
        block = int(len(lines) * body_f.size * 1.35)
        yy = y + (row_h - block) // 2
        wrap_draw(draw, (256, yy), body, body_f, 1560, INK, 1.35)
    footer(img, "04  /  05")
    return img


def slide_05() -> Image.Image:
    img = canvas()
    y0 = kicker_title(
        img,
        "总表",
        "按缺页是否在 TPOT 上，以及这次读能不能喂满接口",
        "先定计时落在 TTFT 还是 TPOT，再定介质。",
    )
    draw = ImageDraw.Draw(img)
    headers = ["场景", "选择", "原因"]
    rows = [
        ("历史 KV 能整段留在 decode 侧 HBM", "不进池", "每步只读 HBM，池子不在 TPOT 上"),
        ("冷前缀在 decode 开始前整段装入 HBM", "SSD 可以", "只搬一次，时间记在 TTFT"),
        ("长上下文，按层连续大读，合同看平均 TPOT", "SSD 可以", "接口被喂满，平均时间与 DDR 相同"),
        ("离线吞吐，没有 TPOT 合同", "SSD 可以", "时延和垃圾回收长尾不进预算"),
        ("小页随机读，上下文一长 IO 次数就上去", "必须 DDR", "SSD 喂不满同一条接口"),
        ("合同写的是 p99 TPOT", "必须 DDR", "垃圾回收的停顿不随大流量摊薄"),
        ("单次 IO 小，时延差大于剩余预算", "必须 DDR", "介质时延直接加进每一个 token"),
    ]
    xs = [64, 980, 1280]
    xe = [960, 1260, 1856]
    y = y0 + 28
    h = 52
    draw.rectangle((64, y, 1856, y + h), fill=ACCENT)
    for x, label in zip(xs, headers):
        draw_text(draw, (x + 24, y + 14), label, font(18, 700), (255, 255, 255))
    y += h
    body_f = font(20, 500)
    choice_f = font(20, 700)
    for i, (a, b, c) in enumerate(rows):
        rh = 96
        fill = SOFT if i % 2 == 0 else BG
        draw.rectangle((64, y, 1856, y + rh), fill=fill)
        draw.line((64, y + rh, 1856, y + rh), fill=LINE, width=1)
        a_lines = wrap(draw, a, body_f, xe[0] - xs[0] - 48)
        c_lines = wrap(draw, c, body_f, xe[2] - xs[2] - 48)
        a_block = int(len(a_lines) * body_f.size * 1.35)
        c_block = int(len(c_lines) * body_f.size * 1.35)
        wrap_draw(draw, (xs[0] + 24, y + (rh - a_block) // 2), a, body_f, xe[0] - xs[0] - 48, INK, 1.35)
        draw_text(draw, (xs[1] + 24, y + rh // 2), b, choice_f, ACCENT, anchor="lm")
        wrap_draw(draw, (xs[2] + 24, y + (rh - c_block) // 2), c, body_f, xe[2] - xs[2] - 48, MUTED, 1.35)
        y += rh
    footer(img, "05  /  05")
    return img


def build() -> None:
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    slides = [slide_01(), slide_02(), slide_03(), slide_04(), slide_05()]
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, img in enumerate(slides, start=1):
        path = SLIDE_DIR / f"{i:02d}.png"
        img.save(path, "PNG")
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    out = OUT / "decode-kv-pool.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
